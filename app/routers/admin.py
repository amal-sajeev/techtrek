import csv
import io
import json
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, File, Query, Request, UploadFile
from fastapi.responses import JSONResponse, RedirectResponse, Response, StreamingResponse
from sqlalchemy import Date, cast, func, or_
from sqlalchemy.orm import Session, joinedload

from app.csrf import csrf_protection
from app.config import settings
from app.dependencies import flash, get_db, now_ist, template_ctx, templates
from app.services.admin_metrics_bundle import build_admin_metrics_bundle
from app.services.metrics_report_ai import finalize_metrics_narrative, run_metrics_report_ai
from app.services.metrics_report_pdf import generate_platform_metrics_report_pdf
from app.services.activity_log import log_activity
from app.models.activity_log import ActivityLog
from app.models.auditorium import Auditorium
from app.services.invoice import generate_invoice_pdf
from app.models.booking import Booking
from app.models.city import City
from app.models.college import College
from app.models.seat import Seat
from app.models.seat_type import SeatType
from app.models.session import Session as SessionModel
from app.models.session_speaker import SessionSpeaker, SPEAKER_ROLES
from app.models.speaker import Speaker
from app.models.agenda import AgendaItem
from app.models.session_recording import SessionRecording
from app.models.event import Event
from app.models.certificate_template import CertificateTemplate
from app.services.certificate import merge_cert_scalar_from_template, normalize_cert_scalar_for_storage
from app.models.coupon import Coupon
from app.models.feedback import Feedback
from app.models.testimonial import Testimonial
from app.models.user import User
from app.services.razorpay import process_refund as rz_process_refund
from app.models.waitlist import Waitlist
from app.models.site_setting import SiteSetting
from app.models.newsletter import Newsletter
from app.models.testimonial import NewsletterSubscriber
from app.models.gallery_image import GalleryImage
from app.models.uploaded_image import UploadedImage
from app.models.event_break import EventBreak
from app.models.event_addon import EventAddOn, BookingAddOn
from app.models.event_session import EventSession
from app.models.feedback_template import FeedbackTemplate, TemplateQuestion, FeedbackResponse, QuestionResponse
from app.models.poll import Poll, PollOption, PollVote
from app.models.event_alert import EventAlert
from app.models.session_feedback import SessionFeedback
from app.models.webhook_log import WebhookLog
from app.models.ticket_share import TicketShare


RECORDING_ALLOWED_HOSTS = {
    "youtube.com", "www.youtube.com", "youtu.be",
    "vimeo.com", "player.vimeo.com",
    "dailymotion.com", "www.dailymotion.com", "dai.ly",
    "twitch.tv", "www.twitch.tv", "clips.twitch.tv",
    "facebook.com", "www.facebook.com", "fb.watch",
    "streamable.com",
    "wistia.com", "fast.wistia.com",
    "loom.com", "www.loom.com",
    "drive.google.com",
}

ADMIN_PAGE_SIZE = 50


def _validate_recording_url(url: str | None) -> str | None:
    """Return an error message if the URL is invalid, or None if it's acceptable."""
    if not url:
        return None
    parsed = urlparse(url)
    if parsed.scheme != "https":
        return "Recording URL must use HTTPS."
    host = parsed.hostname or ""
    if host not in RECORDING_ALLOWED_HOSTS:
        return (
            f"Unsupported recording host '{host}'. "
            "Supported: YouTube, Vimeo, Dailymotion, Twitch, Facebook, "
            "Streamable, Wistia, Loom, Google Drive."
        )
    return None


router = APIRouter(prefix="/admin", tags=["admin"], dependencies=[Depends(csrf_protection)])


def _require_admin(request: Request, db: Session) -> User | None:
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    user = db.query(User).filter(User.id == user_id).first()
    if not user or not user.is_admin:
        return None
    return user


def _admin_ctx(request: Request, active_page: str = "", **kwargs):
    ctx = template_ctx(request, active_page=active_page)
    ctx.update(kwargs)
    return ctx


async def _form(request: Request):
    return await request.form()


def _extract_custom_prices(form) -> dict | None:
    """Build {seat_type_key: price} dict from csp_* form fields."""
    keys = form.getlist("csp_key")
    if not keys:
        return None
    result = {}
    for key in keys:
        raw = (form.get(f"csp_price_{key}", "") or "").strip()
        if raw:
            try:
                result[key] = float(raw)
            except ValueError:
                pass
    return result or None


def _auditorium_seat_types(db) -> dict:
    """Return {aud_id: [sorted list of distinct seat_type strings]} for all auditoriums."""
    rows = (
        db.query(Seat.auditorium_id, Seat.seat_type)
        .filter(Seat.seat_type != "aisle", Seat.is_active == True)
        .distinct()
        .all()
    )
    mapping = defaultdict(set)
    for aud_id, stype in rows:
        mapping[aud_id].add(stype)
    return {k: sorted(v) for k, v in mapping.items()}


def _custom_types_map(db) -> dict:
    """Return {\"custom_N\": {id, name, colour}} for all custom seat types."""
    cts = db.query(SeatType).order_by(SeatType.name).all()
    return {f"custom_{ct.id}": {"id": ct.id, "name": ct.name, "colour": ct.colour} for ct in cts}


def _seat_type_display_name(seat_type_key: str | None, custom_map: dict) -> str:
    """Human label for Seat.seat_type; custom_* keys use SeatType.name from custom_map."""
    k = (seat_type_key or "standard").strip() or "standard"
    if k in custom_map:
        return custom_map[k]["name"]
    builtin = {
        "standard": "Standard",
        "vip": "VIP",
        "accessible": "Accessible",
        "aisle": "Aisle",
        "reserved": "Reserved",
    }
    if k in builtin:
        return builtin[k]
    return k.replace("_", " ").strip().title()


# ─── Dashboard ───

@router.get("/")
def dashboard(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login?next=/admin/", status_code=303)

    total_users = db.query(func.count(User.id)).scalar()
    _paid_not_shared = [Booking.payment_status == "paid", Booking.is_shared_ticket == False]
    total_bookings = db.query(func.count(Booking.id)).filter(*_paid_not_shared).scalar()
    total_revenue = db.query(func.sum(Booking.amount_paid)).filter(*_paid_not_shared).scalar() or 0

    now = now_ist()
    today = now.date()
    upcoming_count = db.query(func.count(Event.id)).filter(
        Event.status == "published", Event.start_date >= today
    ).scalar()

    total_checked_in = db.query(func.count(Booking.id)).filter(Booking.checked_in == True).scalar()
    total_refunded = db.query(func.count(Booking.id)).filter(Booking.payment_status == "refunded").scalar()

    # ── Week-over-week trends ─────────────────────────────────────────────
    week_ago = now - timedelta(days=7)
    two_weeks_ago = now - timedelta(days=14)
    bookings_this_week = db.query(func.count(Booking.id)).filter(
        *_paid_not_shared, Booking.booked_at >= week_ago
    ).scalar() or 0
    bookings_last_week = db.query(func.count(Booking.id)).filter(
        *_paid_not_shared, Booking.booked_at >= two_weeks_ago, Booking.booked_at < week_ago
    ).scalar() or 0
    revenue_this_week = float(db.query(func.sum(Booking.amount_paid)).filter(
        *_paid_not_shared, Booking.booked_at >= week_ago
    ).scalar() or 0)
    revenue_last_week = float(db.query(func.sum(Booking.amount_paid)).filter(
        *_paid_not_shared, Booking.booked_at >= two_weeks_ago, Booking.booked_at < week_ago
    ).scalar() or 0)
    users_this_week = db.query(func.count(User.id)).filter(User.created_at >= week_ago).scalar() or 0
    users_last_week = db.query(func.count(User.id)).filter(
        User.created_at >= two_weeks_ago, User.created_at < week_ago
    ).scalar() or 0

    def _trend(current, previous):
        """Return (delta, direction) where direction is 'up', 'down', or 'flat'."""
        delta = current - previous
        if delta > 0:
            return delta, "up"
        elif delta < 0:
            return abs(delta), "down"
        return 0, "flat"

    bookings_trend = _trend(bookings_this_week, bookings_last_week)
    revenue_trend  = _trend(revenue_this_week,  revenue_last_week)
    users_trend    = _trend(users_this_week,    users_last_week)

    active_events_raw = (
        db.query(Event)
        .filter(Event.status.notin_(["completed", "cancelled"]))
        .order_by(Event.start_date.desc().nullslast(), Event.created_at.desc())
        .all()
    )
    active_events = []
    for ev in active_events_raw:
        session_count = len(ev.event_sessions) if ev.event_sessions else 0
        booking_count = db.query(func.count(Booking.id)).filter(
            Booking.event_id == ev.id, Booking.payment_status == "paid",
            Booking.is_shared_ticket == False,
        ).scalar() or 0
        checked_in = db.query(func.count(Booking.id)).filter(
            Booking.event_id == ev.id, Booking.payment_status == "paid", Booking.checked_in == True
        ).scalar() or 0
        waitlist_count = db.query(func.count(Waitlist.id)).filter(Waitlist.event_id == ev.id).scalar() or 0
        active_events.append({
            "event": ev, "session_count": session_count, "bookings": booking_count,
            "checked_in": checked_in, "waitlist": waitlist_count,
        })

    return templates.TemplateResponse(
        "admin/dashboard.html",
        _admin_ctx(
            request,
            active_page="dashboard",
            total_users=total_users,
            total_bookings=total_bookings,
            total_revenue=float(total_revenue),
            upcoming_count=upcoming_count,
            total_checked_in=total_checked_in,
            total_refunded=total_refunded,
            active_events=active_events,
            bookings_trend=bookings_trend,
            revenue_trend=revenue_trend,
            users_trend=users_trend,
        ),
    )


# ─── Metrics ───

_METRICS_TAB_IDS = frozenset(
    {"overview", "users", "events", "revenue", "feedback", "sessions", "system"}
)


@router.get("/metrics")
def metrics_page(
    request: Request,
    db: Session = Depends(get_db),
    date_from: str = Query("", alias="date_from"),
    date_to: str = Query("", alias="date_to"),
    event_id: str = Query("", alias="event_id"),
    college_id: str = Query("", alias="college_id"),
    tab: str = Query("", alias="tab"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login?next=/admin/metrics", status_code=303)

    f_tab = tab if tab in _METRICS_TAB_IDS else ""

    bundle = build_admin_metrics_bundle(
        db,
        date_from=date_from,
        date_to=date_to,
        event_id=event_id,
        college_id=college_id,
    )
    all_events = db.query(Event).order_by(Event.start_date.desc().nullslast()).all()
    all_colleges = db.query(College).filter(College.is_active == True).order_by(College.name).all()
    metrics_ctx = {k: v for k, v in bundle.items() if k not in ("filter_params", "single_event_selected")}

    return templates.TemplateResponse(
        "admin/metrics.html",
        _admin_ctx(
            request,
            active_page="metrics",
            all_events=all_events, all_colleges=all_colleges,
            f_date_from=date_from, f_date_to=date_to,
            f_event_id=event_id, f_college_id=college_id, f_tab=f_tab,
            **metrics_ctx,
        ),
    )


@router.get("/metrics/report.pdf")
def metrics_report_pdf(
    request: Request,
    db: Session = Depends(get_db),
    date_from: str = Query("", alias="date_from"),
    date_to: str = Query("", alias="date_to"),
    event_id: str = Query("", alias="event_id"),
    college_id: str = Query("", alias="college_id"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login?next=/admin/metrics", status_code=303)

    bundle = build_admin_metrics_bundle(
        db,
        date_from=date_from,
        date_to=date_to,
        event_id=event_id,
        college_id=college_id,
    )
    ai_run = run_metrics_report_ai(settings, bundle)
    narrative = finalize_metrics_narrative(
        bundle, ai_run.narrative, ai_failure_reason=ai_run.failure_reason
    )
    pdf_bytes = generate_platform_metrics_report_pdf(bundle, narrative)
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="techtrek-platform-metrics-report.pdf"'},
    )


# ─── Cities ───

@router.get("/cities")
def cities_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    cities = db.query(City).order_by(City.name).all()
    return templates.TemplateResponse(
        "admin/cities.html",
        _admin_ctx(request, active_page="cities", cities=cities),
    )


@router.get("/cities/new")
def city_new(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    return templates.TemplateResponse(
        "admin/city_form.html",
        _admin_ctx(request, active_page="cities", city=None),
    )


@router.post("/cities/new")
async def city_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    city = City(
        name=form.get("name", "").strip(),
        state=form.get("state", "").strip(),
        is_active="is_active" in form,
    )
    db.add(city)
    db.flush()
    log_activity(db, category="admin", action="create", description=f"Created city '{city.name}'", request=request, user_id=admin.id, target_type="city", target_id=city.id)
    db.commit()
    flash(request, f"City '{city.name}' created.", "success")
    return RedirectResponse("/admin/cities", status_code=303)


@router.get("/cities/{city_id}/edit")
def city_edit(request: Request, city_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    city = db.query(City).get(city_id)
    if not city:
        flash(request, "City not found.", "danger")
        return RedirectResponse("/admin/cities", status_code=303)
    return templates.TemplateResponse(
        "admin/city_form.html",
        _admin_ctx(request, active_page="cities", city=city),
    )


@router.post("/cities/{city_id}/edit")
async def city_update(request: Request, city_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    city = db.query(City).get(city_id)
    if not city:
        return RedirectResponse("/admin/cities", status_code=303)
    form = await _form(request)
    city.name = form.get("name", city.name).strip()
    city.state = form.get("state", city.state).strip()
    city.is_active = "is_active" in form
    log_activity(db, category="admin", action="update", description=f"Updated city '{city.name}'", request=request, user_id=admin.id, target_type="city", target_id=city.id)
    db.commit()
    flash(request, f"City '{city.name}' updated.", "success")
    return RedirectResponse("/admin/cities", status_code=303)


@router.post("/cities/{city_id}/delete")
def city_delete(request: Request, city_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    city = db.query(City).get(city_id)
    if city:
        log_activity(db, category="admin", action="delete", description=f"Deleted city '{city.name}'", request=request, user_id=admin.id, target_type="city", target_id=city_id)
        db.delete(city)
        db.commit()
        flash(request, f"City '{city.name}' deleted.", "success")
    return RedirectResponse("/admin/cities", status_code=303)


@router.post("/cities/{city_id}/toggle")
def city_toggle(request: Request, city_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    city = db.query(City).get(city_id)
    if city:
        city.is_active = not city.is_active
        status = "active" if city.is_active else "inactive"
        log_activity(db, category="admin", action="update", description=f"Toggled city '{city.name}' to {status}", request=request, user_id=admin.id, target_type="city", target_id=city_id)
        db.commit()
        flash(request, f"City '{city.name}' is now {status}.", "success")
    return RedirectResponse("/admin/cities", status_code=303)


# ─── Colleges ───

@router.get("/colleges")
def colleges_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    colleges = db.query(College).order_by(College.name).all()
    enriched = []
    for col in colleges:
        city = db.query(City).get(col.city_id) if col.city_id else None
        aud_names = " ".join(a.name for a in col.auditoriums)
        enriched.append({"college": col, "city": city, "aud_count": len(col.auditoriums), "aud_names": aud_names})
    return templates.TemplateResponse(
        "admin/colleges.html",
        _admin_ctx(request, active_page="colleges", colleges=enriched),
    )


@router.get("/colleges/new")
def college_new(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    cities = db.query(City).filter(City.is_active == True).order_by(City.name).all()
    return templates.TemplateResponse(
        "admin/college_form.html",
        _admin_ctx(request, active_page="colleges", college=None, cities=cities),
    )


@router.post("/colleges/new")
async def college_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    col = College(
        name=form.get("name", "").strip(),
        city_id=int(form.get("city_id")),
        address=form.get("address", "").strip() or None,
        logo_url=form.get("logo_url", "").strip() or None,
        is_active="is_active" in form,
    )
    db.add(col)
    db.flush()
    log_activity(db, category="admin", action="create", description=f"Created college '{col.name}'", request=request, user_id=admin.id, target_type="college", target_id=col.id)
    db.commit()
    flash(request, f"College '{col.name}' created.", "success")
    return RedirectResponse("/admin/colleges", status_code=303)


@router.get("/colleges/{college_id}/edit")
def college_edit(request: Request, college_id: int, tab: str = "details", db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    col = db.query(College).get(college_id)
    if not col:
        flash(request, "College not found.", "danger")
        return RedirectResponse("/admin/colleges", status_code=303)
    cities = db.query(City).filter(City.is_active == True).order_by(City.name).all()
    auditoriums = db.query(Auditorium).filter(Auditorium.college_id == college_id).order_by(Auditorium.name).all()
    return templates.TemplateResponse(
        "admin/college_form.html",
        _admin_ctx(request, active_page="colleges", college=col, cities=cities, auditoriums=auditoriums, active_tab=tab),
    )


@router.post("/colleges/{college_id}/edit")
async def college_update(request: Request, college_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    col = db.query(College).get(college_id)
    if not col:
        return RedirectResponse("/admin/colleges", status_code=303)
    form = await _form(request)
    col.name = form.get("name", col.name).strip()
    col.city_id = int(form.get("city_id", col.city_id))
    col.address = form.get("address", "").strip() or None
    col.logo_url = form.get("logo_url", "").strip() or None
    col.is_active = "is_active" in form
    log_activity(db, category="admin", action="update", description=f"Updated college '{col.name}'", request=request, user_id=admin.id, target_type="college", target_id=college_id)
    db.commit()
    flash(request, f"College '{col.name}' updated.", "success")
    return RedirectResponse("/admin/colleges", status_code=303)


@router.post("/colleges/{college_id}/delete")
def college_delete(request: Request, college_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    col = db.query(College).get(college_id)
    if col:
        log_activity(db, category="admin", action="delete", description=f"Deleted college '{col.name}'", request=request, user_id=admin.id, target_type="college", target_id=college_id)
        db.delete(col)
        db.commit()
        flash(request, f"College '{col.name}' deleted.", "success")
    return RedirectResponse("/admin/colleges", status_code=303)


# ─── Auditoriums (nested under colleges) ───

@router.get("/auditoriums")
def auditoriums_list_redirect(request: Request):
    return RedirectResponse("/admin/colleges", status_code=303)


@router.post("/colleges/{college_id}/auditoriums/new")
async def auditorium_create(request: Request, college_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    col = db.query(College).get(college_id)
    if not col:
        return RedirectResponse("/admin/colleges", status_code=303)

    form = await _form(request)
    aud = Auditorium(
        name=form.get("name", "").strip(),
        college_id=college_id,
        location=form.get("location", "").strip(),
        description=form.get("description", "").strip(),
        total_rows=int(form.get("total_rows", 10)),
        total_cols=int(form.get("total_cols", 15)),
    )
    db.add(aud)
    db.commit()
    db.refresh(aud)
    log_activity(db, category="admin", action="create", description=f"Created auditorium '{aud.name}'", request=request, user_id=admin.id, target_type="auditorium", target_id=aud.id)
    db.commit()
    flash(request, f"Auditorium '{aud.name}' created.", "success")
    return RedirectResponse(f"/admin/colleges/{college_id}/edit?tab=auditoriums", status_code=303)


@router.post("/colleges/{college_id}/auditoriums/{aud_id}/edit")
async def auditorium_update(request: Request, college_id: int, aud_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    aud = db.query(Auditorium).get(aud_id)
    if not aud:
        flash(request, "Auditorium not found.", "danger")
        return RedirectResponse(f"/admin/colleges/{college_id}/edit?tab=auditoriums", status_code=303)

    form = await _form(request)
    aud.name = form.get("name", aud.name).strip()
    aud.location = form.get("location", aud.location).strip()
    aud.description = form.get("description", "").strip()
    aud.total_rows = int(form.get("total_rows", aud.total_rows))
    aud.total_cols = int(form.get("total_cols", aud.total_cols))
    log_activity(db, category="admin", action="update", description=f"Updated auditorium '{aud.name}'", request=request, user_id=admin.id, target_type="auditorium", target_id=aud_id)
    db.commit()
    flash(request, f"Auditorium '{aud.name}' updated.", "success")
    return RedirectResponse(f"/admin/colleges/{college_id}/edit?tab=auditoriums", status_code=303)


@router.post("/colleges/{college_id}/auditoriums/{aud_id}/delete")
def auditorium_delete(request: Request, college_id: int, aud_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    aud = db.query(Auditorium).get(aud_id)
    if aud:
        log_activity(db, category="admin", action="delete", description=f"Deleted auditorium '{aud.name}'", request=request, user_id=admin.id, target_type="auditorium", target_id=aud_id)
        db.delete(aud)
        db.commit()
        flash(request, f"Auditorium '{aud.name}' deleted.", "success")
    return RedirectResponse(f"/admin/colleges/{college_id}/edit?tab=auditoriums", status_code=303)


# ─── Seat Types ───

@router.get("/seat-types")
def seat_types_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    seat_types = db.query(SeatType).order_by(SeatType.name).all()
    return templates.TemplateResponse(
        "admin/seat_types.html",
        _admin_ctx(request, active_page="seat_types", seat_types=seat_types),
    )


@router.get("/seat-types/new")
def seat_type_new(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    return templates.TemplateResponse(
        "admin/seat_type_form.html",
        _admin_ctx(request, active_page="seat_types", seat_type=None),
    )


@router.post("/seat-types/new")
async def seat_type_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    st = SeatType(
        name=form.get("name", "").strip(),
        colour=form.get("colour", "#6366f1").strip(),
        icon=form.get("icon", "").strip() or None,
        price=float(form["price"]) if form.get("price", "").strip() else None,
        is_custom=True,
    )
    db.add(st)
    db.flush()
    log_activity(db, category="admin", action="create", description=f"Created seat type '{st.name}'", request=request, user_id=admin.id, target_type="seat_type", target_id=st.id)
    db.commit()
    flash(request, f"Seat type '{st.name}' created.", "success")
    return RedirectResponse("/admin/seat-types", status_code=303)


@router.get("/seat-types/{st_id}/edit")
def seat_type_edit(request: Request, st_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    st = db.query(SeatType).get(st_id)
    if not st:
        flash(request, "Seat type not found.", "danger")
        return RedirectResponse("/admin/seat-types", status_code=303)
    return templates.TemplateResponse(
        "admin/seat_type_form.html",
        _admin_ctx(request, active_page="seat_types", seat_type=st),
    )


@router.post("/seat-types/{st_id}/edit")
async def seat_type_update(request: Request, st_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    st = db.query(SeatType).get(st_id)
    if not st:
        return RedirectResponse("/admin/seat-types", status_code=303)
    form = await _form(request)
    st.name = form.get("name", st.name).strip()
    st.colour = form.get("colour", st.colour).strip()
    st.icon = form.get("icon", "").strip() or None
    st.price = float(form["price"]) if form.get("price", "").strip() else None
    log_activity(db, category="admin", action="update", description=f"Updated seat type '{st.name}'", request=request, user_id=admin.id, target_type="seat_type", target_id=st_id)
    db.commit()
    flash(request, f"Seat type '{st.name}' updated.", "success")
    return RedirectResponse("/admin/seat-types", status_code=303)


@router.post("/seat-types/{st_id}/delete")
def seat_type_delete(request: Request, st_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    st = db.query(SeatType).get(st_id)
    if st:
        in_use = db.query(Seat).filter(Seat.seat_type == f"custom_{st.id}").count()
        if in_use:
            flash(request, f"Cannot delete '{st.name}' — it is used by {in_use} seat(s).", "danger")
            return RedirectResponse("/admin/seat-types", status_code=303)
        log_activity(db, category="admin", action="delete", description=f"Deleted seat type '{st.name}'", request=request, user_id=admin.id, target_type="seat_type", target_id=st_id)
        db.delete(st)
        db.commit()
        flash(request, f"Seat type '{st.name}' deleted.", "success")
    return RedirectResponse("/admin/seat-types", status_code=303)


# ─── Seat Layout Designer ───

@router.get("/auditoriums/{aud_id}/layout")
def seat_layout(request: Request, aud_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    aud = db.query(Auditorium).get(aud_id)
    if not aud:
        flash(request, "Auditorium not found.", "danger")
        return RedirectResponse("/admin/colleges", status_code=303)

    seats = db.query(Seat).filter(Seat.auditorium_id == aud_id).order_by(Seat.row_num, Seat.col_num).all()
    seat_data = [
        {"id": s.id, "row": s.row_num, "col": s.col_num, "label": s.label,
         "type": s.seat_type, "active": s.is_active}
        for s in seats
    ]

    custom_types = db.query(SeatType).filter(SeatType.is_custom == True).order_by(SeatType.name).all()
    custom_types_data = [
        {"id": st.id, "name": st.name, "colour": st.colour, "icon": st.icon}
        for st in custom_types
    ]

    return templates.TemplateResponse(
        "admin/seat_layout.html",
        _admin_ctx(
            request,
            active_page="auditoriums",
            auditorium=aud,
            seat_data=seat_data,
            custom_types=custom_types_data,
            row_gaps=json.loads(aud.row_gaps) if aud.row_gaps else [],
            col_gaps=json.loads(aud.col_gaps) if aud.col_gaps else [],
            entry_exit_config=aud.entry_exit_config or [],
        ),
    )


@router.post("/auditoriums/{aud_id}/layout")
async def seat_layout_save(request: Request, aud_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    aud = db.query(Auditorium).get(aud_id)
    if not aud:
        return RedirectResponse("/admin/colleges", status_code=303)

    form = await _form(request)
    layout_json = form.get("layout_data", "[]")
    total_rows = form.get("total_rows")
    total_cols = form.get("total_cols")
    stage_cols_raw = form.get("stage_cols")
    try:
        layout = json.loads(layout_json)
    except json.JSONDecodeError:
        flash(request, "Invalid layout data.", "danger")
        return RedirectResponse(f"/admin/auditoriums/{aud_id}/layout", status_code=303)

    if total_rows is not None and total_cols is not None:
        try:
            r, c = int(total_rows), int(total_cols)
            if 1 <= r <= 50 and 1 <= c <= 50:
                aud.total_rows = r
                aud.total_cols = c
        except ValueError:
            pass

    if stage_cols_raw is not None:
        try:
            sc = int(stage_cols_raw)
            if 1 <= sc <= aud.total_cols:
                aud.stage_cols = sc
            else:
                aud.stage_cols = aud.total_cols
        except ValueError:
            aud.stage_cols = None
    else:
        aud.stage_cols = None

    stage_offset_raw = form.get("stage_offset", "0")
    try:
        aud.stage_offset = max(0, int(stage_offset_raw))
    except (ValueError, TypeError):
        aud.stage_offset = 0

    stage_label_raw = form.get("stage_label", "Stage")
    aud.stage_label = (stage_label_raw or "Stage").strip()[:100]

    entry_exit_raw = form.get("entry_exit_config", "[]")
    try:
        aud.entry_exit_config = json.loads(entry_exit_raw)
    except (json.JSONDecodeError, ValueError):
        aud.entry_exit_config = None

    row_gaps_raw = form.get("row_gaps", "[]")
    col_gaps_raw = form.get("col_gaps", "[]")
    try:
        rg = json.loads(row_gaps_raw)
        aud.row_gaps = json.dumps([int(x) for x in rg if 1 <= int(x) < aud.total_rows]) if rg else None
    except (json.JSONDecodeError, ValueError):
        aud.row_gaps = None
    try:
        cg = json.loads(col_gaps_raw)
        aud.col_gaps = json.dumps([int(x) for x in cg if 1 <= int(x) < aud.total_cols]) if cg else None
    except (json.JSONDecodeError, ValueError):
        aud.col_gaps = None

    seat_ids = [s.id for s in db.query(Seat.id).filter(Seat.auditorium_id == aud_id).all()]
    if seat_ids:
        db.query(Booking).filter(Booking.seat_id.in_(seat_ids)).delete(synchronize_session="fetch")
        db.query(Seat).filter(Seat.id.in_(seat_ids)).delete(synchronize_session="fetch")

    for item in layout:
        seat = Seat(
            auditorium_id=aud_id,
            row_num=item["row"],
            col_num=item["col"],
            label=item.get("label", f"{chr(64 + item['row'])}{item['col']}"),
            seat_type=item.get("type", "standard"),
            is_active=item.get("active", True),
        )
        db.add(seat)

    aud.layout_config = layout
    seat_count = sum(1 for item in layout if item.get("type") != "aisle")
    log_activity(db, category="admin", action="update", description=f"Saved seat layout for '{aud.name}' ({seat_count} seats)", request=request, user_id=admin.id, target_type="auditorium", target_id=aud_id)
    db.commit()
    flash(request, "Seat layout saved.", "success")
    return RedirectResponse(f"/admin/auditoriums/{aud_id}/layout", status_code=303)


# ─── Speakers ───

@router.get("/speakers")
def speakers_list(
    request: Request,
    db: Session = Depends(get_db),
    role: str = Query("", alias="role"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    query = db.query(Speaker)
    if role:
        query = query.join(SessionSpeaker).filter(SessionSpeaker.role == role)
    speakers = query.order_by(Speaker.name).all()
    return templates.TemplateResponse(
        "admin/speakers.html",
        _admin_ctx(request, active_page="speakers", speakers=speakers,
                   speaker_roles=SPEAKER_ROLES, role_filter=role),
    )


@router.get("/speakers/new")
def speaker_new(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    return templates.TemplateResponse(
        "admin/speaker_form.html",
        _admin_ctx(request, active_page="speakers", speaker=None),
    )


@router.post("/speakers/new")
async def speaker_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    sp = Speaker(
        name=form.get("name", "").strip(),
        title=form.get("title", "").strip() or None,
        bio=form.get("bio", "").strip() or None,
        photo_url=form.get("photo_url", "").strip() or None,
        email=form.get("email", "").strip() or None,
    )
    db.add(sp)
    db.flush()
    log_activity(db, category="admin", action="create", description=f"Created speaker '{sp.name}'", request=request, user_id=admin.id, target_type="speaker", target_id=sp.id)
    db.commit()
    flash(request, f"Speaker '{sp.name}' created.", "success")
    return RedirectResponse("/admin/speakers", status_code=303)


@router.get("/speakers/{speaker_id}/edit")
def speaker_edit(request: Request, speaker_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    sp = db.query(Speaker).get(speaker_id)
    if not sp:
        flash(request, "Speaker not found.", "danger")
        return RedirectResponse("/admin/speakers", status_code=303)
    return templates.TemplateResponse(
        "admin/speaker_form.html",
        _admin_ctx(request, active_page="speakers", speaker=sp),
    )


@router.post("/speakers/{speaker_id}/edit")
async def speaker_update(request: Request, speaker_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    sp = db.query(Speaker).get(speaker_id)
    if not sp:
        return RedirectResponse("/admin/speakers", status_code=303)
    form = await _form(request)
    sp.name = form.get("name", sp.name).strip()
    sp.title = form.get("title", "").strip() or None
    sp.bio = form.get("bio", "").strip() or None
    sp.photo_url = form.get("photo_url", "").strip() or None
    sp.email = form.get("email", "").strip() or None
    log_activity(db, category="admin", action="update", description=f"Updated speaker '{sp.name}'", request=request, user_id=admin.id, target_type="speaker", target_id=speaker_id)
    db.commit()
    flash(request, f"Speaker '{sp.name}' updated.", "success")
    return RedirectResponse("/admin/speakers", status_code=303)


@router.get("/speakers/{speaker_id}/delete-check")
def speaker_delete_check(request: Request, speaker_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "unauthorized"}, status_code=403)
    sp = db.query(Speaker).get(speaker_id)
    if not sp:
        return JSONResponse({"error": "not_found"}, status_code=404)
    from app.models.event_session import EventSession
    now = now_ist()
    primary_sessions = (
        db.query(SessionModel, Event)
        .join(EventSession, EventSession.session_id == SessionModel.id)
        .join(Event, EventSession.event_id == Event.id)
        .filter(SessionModel.speaker_id == speaker_id, Event.start_date >= now.date())
        .all()
    )
    agenda_sessions = (
        db.query(SessionModel, Event)
        .join(AgendaItem, AgendaItem.session_id == SessionModel.id)
        .join(EventSession, EventSession.session_id == SessionModel.id)
        .join(Event, EventSession.event_id == Event.id)
        .filter(AgendaItem.speaker_id == speaker_id, Event.start_date >= now.date())
        .all()
    )
    seen = set()
    sessions_list = []
    for s, ev in primary_sessions:
        if s.id not in seen:
            seen.add(s.id)
            date_str = ev.start_date.strftime("%b %d, %Y") if ev and ev.start_date else ""
            sessions_list.append({"id": s.id, "title": s.title, "date": date_str, "role": "Primary Speaker"})
    for s, ev in agenda_sessions:
        if s.id not in seen:
            seen.add(s.id)
            date_str = ev.start_date.strftime("%b %d, %Y") if ev and ev.start_date else ""
            sessions_list.append({"id": s.id, "title": s.title, "date": date_str, "role": "Agenda Item"})
    return JSONResponse({"speaker_name": sp.name, "has_sessions": len(sessions_list) > 0, "sessions": sessions_list})


@router.post("/speakers/{speaker_id}/delete")
async def speaker_delete(request: Request, speaker_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    sp = db.query(Speaker).get(speaker_id)
    if not sp:
        return RedirectResponse("/admin/speakers", status_code=303)
    form = await _form(request)
    session_action = form.get("session_action", "draft")
    linked_primary = db.query(SessionModel).filter(SessionModel.speaker_id == speaker_id).all()
    linked_agenda = db.query(AgendaItem).filter(AgendaItem.speaker_id == speaker_id).all()
    if session_action == "delete":
        for s in linked_primary:
            db.delete(s)
    else:
        for s in linked_primary:
            s.speaker_id = None
    for ai in linked_agenda:
        ai.speaker_id = None
        ai.speaker_name = None
    db.flush()
    log_activity(db, category="admin", action="delete", description=f"Deleted speaker '{sp.name}' (sessions: {session_action})", request=request, user_id=admin.id, target_type="speaker", target_id=speaker_id)
    db.delete(sp)
    db.commit()
    flash(request, f"Speaker '{sp.name}' deleted.", "success")
    return RedirectResponse("/admin/speakers", status_code=303)


@router.post("/speakers/{speaker_id}/invite")
def speaker_invite(request: Request, speaker_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    sp = db.query(Speaker).get(speaker_id)
    if not sp:
        flash(request, "Speaker not found.", "danger")
        return RedirectResponse("/admin/speakers", status_code=303)
    if not sp.email:
        flash(request, "Speaker has no email address. Add an email first.", "danger")
        return RedirectResponse("/admin/speakers", status_code=303)
    if sp.user_id:
        flash(request, f"Speaker '{sp.name}' is already linked to a user account.", "warning")
        return RedirectResponse("/admin/speakers", status_code=303)

    import secrets
    sp.invite_token = secrets.token_hex(32)
    sp.invite_token_expires = now_ist() + timedelta(days=7)
    db.commit()

    base_url = str(request.base_url).rstrip("/")
    invite_url = f"{base_url}/auth/speaker-invite/{sp.invite_token}"

    from app.services.email import send_speaker_invite
    sent = send_speaker_invite(sp.email, sp.name, invite_url)

    if sent:
        log_activity(db, category="admin", action="invite_sent", description=f"Sent speaker invite to '{sp.name}' ({sp.email})", request=request, user_id=admin.id, target_type="speaker", target_id=speaker_id)
        db.commit()
        flash(request, f"Invite sent to {sp.email}.", "success")
    else:
        log_activity(db, category="system", action="email_failed", description=f"Speaker invite email to '{sp.name}' ({sp.email}) failed", request=request, user_id=admin.id, target_type="speaker", target_id=speaker_id)
        db.commit()
        flash(request, f"Invite token created but email to {sp.email} failed. Check SMTP settings or share the link manually: {invite_url}", "warning")
    return RedirectResponse("/admin/speakers", status_code=303)


# ─── Sessions ───

@router.get("/sessions")
def sessions_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    from app.models.event_session import EventSession
    all_sessions = db.query(SessionModel).order_by(SessionModel.created_at.desc()).all()
    enriched = []
    for s in all_sessions:
        linked_count = db.query(func.count(EventSession.id)).filter(EventSession.session_id == s.id).scalar() or 0
        enriched.append({"session": s, "linked_events": linked_count})
    return templates.TemplateResponse(
        "admin/sessions.html",
        _admin_ctx(request, active_page="sessions", sessions=enriched),
    )


@router.get("/sessions/new")
def session_new(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    speakers = db.query(Speaker).order_by(Speaker.name).all()
    return templates.TemplateResponse(
        "admin/session_form.html",
        _admin_ctx(request, active_page="sessions", lecture=None,
                   speakers=speakers,
                   agenda_items=[], session_speakers=[], speaker_roles=SPEAKER_ROLES),
    )


@router.post("/sessions/new")
async def session_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    form = await _form(request)

    speaker_id_raw = form.get("speaker_id")
    speaker_id = int(speaker_id_raw) if speaker_id_raw and speaker_id_raw != "" else None

    session_obj = SessionModel(
        speaker_id=speaker_id,
        title=form.get("title", "").strip(),
        speaker_name=form.get("speaker_name", "").strip() or form.get("speaker", "").strip(),
        abstract=form.get("abstract", "").strip() or form.get("description", "").strip(),
        key_learning_outcomes=form.get("key_learning_outcomes", "").strip(),
        description=form.get("abstract", "").strip() or form.get("description", "").strip(),
        banner_url=form.get("banner_url", "").strip() or None,
        duration_minutes=int(form.get("duration_minutes", 30)),
    )

    db.add(session_obj)
    db.flush()

    _save_agenda_items(db, form, session_obj.id)
    _save_session_speakers(db, form, session_obj.id)
    _save_gallery_images(db, form, "session", session_obj.id)

    log_activity(db, category="admin", action="create", description=f"Created session '{session_obj.title}'", request=request, user_id=admin.id, target_type="session", target_id=session_obj.id)
    db.commit()
    flash(request, f"Session '{session_obj.title}' created.", "success")
    return RedirectResponse("/admin/sessions", status_code=303)


@router.post("/api/sessions/create-quick")
async def session_create_quick(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"error": "Invalid request"}, status_code=400)
    title = (body.get("title") or "").strip()
    if not title:
        return JSONResponse({"error": "Title is required"}, status_code=400)
    speaker_name = (body.get("speaker_name") or "").strip()
    speaker_id_raw = body.get("speaker_id")
    speaker_id = int(speaker_id_raw) if speaker_id_raw else None
    description = (body.get("description") or "").strip()
    session_obj = SessionModel(
        speaker_id=speaker_id,
        title=title,
        speaker_name=speaker_name,
        description=description,
        abstract=(body.get("abstract") or "").strip() or description,
        key_learning_outcomes=(body.get("key_learning_outcomes") or "").strip(),
        banner_url=(body.get("banner_url") or "").strip() or None,
        duration_minutes=int(body.get("duration_minutes") or 30),
    )
    db.add(session_obj)
    db.flush()
    log_activity(db, category="admin", action="create",
                 description=f"Quick-created session '{session_obj.title}'",
                 request=request, user_id=admin.id,
                 target_type="session", target_id=session_obj.id)
    db.commit()
    return JSONResponse({
        "id": session_obj.id,
        "title": session_obj.title,
        "speaker_name": session_obj.speaker_name,
        "speaker_id": session_obj.speaker_id,
        "duration": session_obj.duration_minutes,
        "description": session_obj.description or "",
        "abstract": session_obj.abstract or "",
        "key_learning_outcomes": session_obj.key_learning_outcomes or "",
        "banner_url": session_obj.banner_url or "",
    })


@router.post("/api/speakers/create-quick")
async def speaker_create_quick(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"error": "Invalid request"}, status_code=400)
    name = (body.get("name") or "").strip()
    if not name:
        return JSONResponse({"error": "Speaker name is required"}, status_code=400)
    email = (body.get("email") or "").strip()
    if not email:
        return JSONResponse({"error": "Email is required"}, status_code=400)
    sp = Speaker(name=name, email=email)
    db.add(sp)
    db.flush()
    log_activity(db, category="admin", action="create",
                 description=f"Quick-created speaker '{sp.name}'",
                 request=request, user_id=admin.id,
                 target_type="speaker", target_id=sp.id)
    db.commit()
    return JSONResponse({"id": sp.id, "name": sp.name, "title": sp.title or "", "email": sp.email})


@router.get("/sessions/{sess_id}/edit")
def session_edit(request: Request, sess_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    lecture = db.query(SessionModel).get(sess_id)
    if not lecture:
        flash(request, "Session not found.", "danger")
        return RedirectResponse("/admin/sessions", status_code=303)
    speakers = db.query(Speaker).order_by(Speaker.name).all()
    agenda_items = db.query(AgendaItem).filter(AgendaItem.session_id == sess_id).order_by(AgendaItem.order).all()
    session_speakers = db.query(SessionSpeaker).filter(SessionSpeaker.session_id == sess_id).all()
    gallery = db.query(GalleryImage).filter(
        GalleryImage.owner_type == "session", GalleryImage.owner_id == sess_id
    ).order_by(GalleryImage.position).all()
    return templates.TemplateResponse(
        "admin/session_form.html",
        _admin_ctx(request, active_page="sessions", lecture=lecture,
                   speakers=speakers,
                   agenda_items=agenda_items, session_speakers=session_speakers,
                   speaker_roles=SPEAKER_ROLES, gallery_images=gallery),
    )


@router.post("/sessions/{sess_id}/edit")
async def session_update(request: Request, sess_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    lecture = db.query(SessionModel).get(sess_id)
    if not lecture:
        return RedirectResponse("/admin/sessions", status_code=303)

    form = await _form(request)

    speaker_id_raw = form.get("speaker_id")
    lecture.speaker_id = int(speaker_id_raw) if speaker_id_raw and speaker_id_raw != "" else None
    lecture.title = form.get("title", lecture.title).strip()
    lecture.speaker_name = form.get("speaker_name", "").strip() or form.get("speaker", lecture.speaker_name).strip()
    lecture.abstract = form.get("abstract", "").strip() or form.get("description", "").strip()
    lecture.key_learning_outcomes = form.get("key_learning_outcomes", "").strip()
    lecture.description = lecture.abstract
    lecture.banner_url = form.get("banner_url", "").strip() or None
    lecture.duration_minutes = int(form.get("duration_minutes", 30))

    _save_agenda_items(db, form, sess_id)
    _save_session_speakers(db, form, sess_id)
    _save_gallery_images(db, form, "session", sess_id)

    log_activity(db, category="admin", action="update", description=f"Updated session '{lecture.title}'", request=request, user_id=admin.id, target_type="session", target_id=sess_id)
    db.commit()
    flash(request, f"Session '{lecture.title}' updated.", "success")
    return RedirectResponse(f"/admin/sessions/{sess_id}/edit", status_code=303)


def _collect_speaker_ids_from_form(form) -> list[int]:
    ids = []
    idx = 0
    while True:
        raw = form.get(f"session_speaker_id_{idx}")
        if raw is None:
            break
        raw = raw.strip()
        if raw:
            ids.append(int(raw))
        idx += 1
    return ids


def _save_gallery_images(db: Session, form, owner_type: str, owner_id: int):
    """Sync gallery images from form data. Expects gallery_image_ids as a
    comma-separated list of UploadedImage IDs in display order."""
    db.query(GalleryImage).filter(
        GalleryImage.owner_type == owner_type,
        GalleryImage.owner_id == owner_id,
    ).delete()
    raw = form.get("gallery_image_ids", "").strip()
    if not raw:
        return
    for pos, img_id_str in enumerate(raw.split(",")):
        img_id_str = img_id_str.strip()
        if not img_id_str or not img_id_str.isdigit():
            continue
        db.add(GalleryImage(
            owner_type=owner_type,
            owner_id=owner_id,
            image_id=int(img_id_str),
            position=pos,
        ))


def _save_agenda_items(db: Session, form, session_id: int):
    db.query(AgendaItem).filter(AgendaItem.session_id == session_id).delete()
    idx = 0
    while True:
        title = form.get(f"agenda_title_{idx}")
        if title is None:
            break
        title = title.strip()
        if title:
            sp_id_raw = form.get(f"agenda_speaker_id_{idx}", "").strip()
            sp_id = int(sp_id_raw) if sp_id_raw else None
            sp_name = None
            if sp_id:
                sp_obj = db.query(Speaker).get(sp_id)
                sp_name = sp_obj.name if sp_obj else None
            item = AgendaItem(
                session_id=session_id,
                order=idx,
                title=title,
                speaker_id=sp_id,
                speaker_name=sp_name,
                duration_minutes=int(form.get(f"agenda_duration_{idx}", 20) or 20),
                description=form.get(f"agenda_desc_{idx}", "").strip() or None,
            )
            db.add(item)
        idx += 1
    db.commit()


def _save_event_sessions(db: Session, form, event_id: int) -> int:
    """Sync EventSession rows from the event form. Returns count linked."""
    from app.models.event_session import EventSession
    db.query(EventSession).filter(EventSession.event_id == event_id).delete()
    db.flush()
    sess_indices = form.getlist("sess_idx")
    linked = 0
    for idx in sess_indices:
        sess_id = form.get(f"sess_id_{idx}", "").strip()
        if not sess_id or not sess_id.isdigit():
            continue
        sess = db.query(SessionModel).get(int(sess_id))
        if not sess:
            continue
        start_time = None
        start_str = form.get(f"sess_start_{idx}", "")
        if start_str:
            try:
                start_time = datetime.fromisoformat(start_str)
            except ValueError:
                pass
        speaker_id = None
        speaker_name = None
        speaker_id_raw = form.get(f"sess_speaker_id_{idx}", "").strip()
        if speaker_id_raw and speaker_id_raw.isdigit():
            speaker = db.query(Speaker).get(int(speaker_id_raw))
            if speaker:
                speaker_id = speaker.id
                speaker_name = speaker.name
        order = int(form.get(f"sess_order_{idx}", 0) or 0)

        custom_title = form.get(f"sess_custom_title_{idx}", "").strip() or None
        custom_description = form.get(f"sess_custom_description_{idx}", "").strip() or None
        custom_abstract = form.get(f"sess_custom_abstract_{idx}", "").strip() or None
        custom_klo = form.get(f"sess_custom_key_learning_outcomes_{idx}", "").strip() or None
        custom_banner_url = form.get(f"sess_custom_banner_url_{idx}", "").strip() or None
        custom_dur_raw = form.get(f"sess_custom_duration_minutes_{idx}", "").strip()
        custom_duration_minutes = int(custom_dur_raw) if custom_dur_raw and custom_dur_raw.isdigit() else None
        custom_recording_url = form.get(f"sess_custom_recording_url_{idx}", "").strip() or None
        custom_is_rec_raw = form.get(f"sess_custom_is_recording_public_{idx}", "")
        custom_is_recording_public = True if custom_is_rec_raw == "1" else (False if custom_is_rec_raw == "0" else None)

        es = EventSession(
            event_id=event_id,
            session_id=sess.id,
            order=order,
            start_time=start_time,
            speaker_id=speaker_id,
            speaker_name=speaker_name,
            custom_title=custom_title,
            custom_description=custom_description,
            custom_abstract=custom_abstract,
            custom_key_learning_outcomes=custom_klo,
            custom_banner_url=custom_banner_url,
            custom_duration_minutes=custom_duration_minutes,
            custom_recording_url=custom_recording_url,
            custom_is_recording_public=custom_is_recording_public,
        )
        db.add(es)
        linked += 1
    return linked


def _save_event_breaks(db: Session, form, event_id: int):
    db.query(EventBreak).filter(EventBreak.event_id == event_id).delete()
    idx = 0
    while True:
        title = form.get(f"break_title_{idx}")
        if title is None:
            break
        title = title.strip()
        if title:
            start_raw = form.get(f"break_start_time_{idx}", "").strip()
            start_time = None
            if start_raw:
                try:
                    start_time = datetime.fromisoformat(start_raw)
                except ValueError:
                    pass
            brk = EventBreak(
                event_id=event_id,
                order=int(form.get(f"break_order_{idx}", idx) or idx),
                title=title,
                description=form.get(f"break_desc_{idx}", "").strip() or None,
                duration_minutes=int(form.get(f"break_duration_{idx}", 15) or 15),
                start_time=start_time,
            )
            db.add(brk)
        idx += 1
    db.commit()


def _save_event_addons(db: Session, form, event_id: int):
    db.query(EventAddOn).filter(EventAddOn.event_id == event_id).delete()
    idx = 0
    while True:
        title = form.get(f"addon_title_{idx}")
        if title is None:
            break
        title = title.strip()
        if title:
            in_agenda = form.get(f"addon_in_agenda_{idx}") == "1"
            raw_start = (form.get(f"addon_start_time_{idx}") or "").strip()
            start_time = None
            if raw_start:
                try:
                    from datetime import datetime
                    start_time = datetime.fromisoformat(raw_start)
                except (ValueError, TypeError):
                    pass
            addon = EventAddOn(
                event_id=event_id,
                title=title,
                description=form.get(f"addon_desc_{idx}", "").strip() or None,
                price=float(form.get(f"addon_price_{idx}", 0) or 0),
                max_quantity=int(form.get(f"addon_max_qty_{idx}") or 0) or None,
                is_active=True,
                in_agenda=in_agenda,
                order=int(form.get(f"addon_order_{idx}") or 0),
                start_time=start_time,
            )
            db.add(addon)
        idx += 1
    db.commit()


def _save_session_speakers(db: Session, form, session_id: int):
    db.query(SessionSpeaker).filter(SessionSpeaker.session_id == session_id).delete()
    idx = 0
    while True:
        sp_id_raw = form.get(f"session_speaker_id_{idx}")
        if sp_id_raw is None:
            break
        sp_id_raw = sp_id_raw.strip()
        if sp_id_raw:
            role = form.get(f"session_speaker_role_{idx}", "Guest").strip()
            if role not in SPEAKER_ROLES:
                role = "Guest"
            ss = SessionSpeaker(
                session_id=session_id,
                speaker_id=int(sp_id_raw),
                role=role,
            )
            db.add(ss)
        idx += 1
    db.flush()



@router.post("/sessions/{sess_id}/delete")
def session_delete(request: Request, sess_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    lecture = db.query(SessionModel).get(sess_id)
    if lecture:
        log_activity(db, category="admin", action="delete", description=f"Deleted session '{lecture.title}'", request=request, user_id=admin.id, target_type="session", target_id=sess_id)
        db.delete(lecture)
        db.commit()
        flash(request, f"Session '{lecture.title}' deleted.", "success")
    return RedirectResponse("/admin/sessions", status_code=303)




# ─── Certificate Preview (Event-level) ───

@router.get("/events/{event_id}/certificate/preview")
def event_certificate_preview(
    request: Request, event_id: int, db: Session = Depends(get_db)
):
    import io as _io
    from types import SimpleNamespace
    from app.services.certificate import generate_certificate_pdf

    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    event = db.query(Event).get(event_id)
    if not event:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    auditorium = db.query(Auditorium).get(event.auditorium_id) if event.auditorium_id else None

    dummy_booking = SimpleNamespace(
        booking_ref="PREVIEW",
        qr_code_data="CERT-PREVIEW-SAMPLE",
    )
    dummy_user = SimpleNamespace(
        full_name="Sample Attendee",
        username="sample_attendee",
    )

    pdf_bytes = generate_certificate_pdf(dummy_booking, dummy_user, event, event, auditorium)

    return StreamingResponse(
        _io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={
            "Content-Disposition": 'inline; filename="certificate-preview.pdf"',
            "Cache-Control": "no-store",
        },
    )


@router.post("/events/{event_id}/certificate/save")
async def event_certificate_save(
    request: Request, event_id: int, db: Session = Depends(get_db)
):
    """Save only certificate template fields on the event."""
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    event = db.query(Event).get(event_id)
    if not event:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    form = await request.form()
    event.cert_title = normalize_cert_scalar_for_storage(form.get("cert_title"))
    event.cert_subtitle = normalize_cert_scalar_for_storage(form.get("cert_subtitle"))
    event.cert_footer = normalize_cert_scalar_for_storage(form.get("cert_footer"))
    event.cert_signer_name = normalize_cert_scalar_for_storage(form.get("cert_signer_name"))
    event.cert_signer_designation = normalize_cert_scalar_for_storage(form.get("cert_signer_designation"))
    event.cert_signature_url = normalize_cert_scalar_for_storage(form.get("cert_signature_url"))
    event.cert_logo_url = normalize_cert_scalar_for_storage(form.get("cert_logo_url"))
    event.cert_bg_url = normalize_cert_scalar_for_storage(form.get("cert_bg_url"))
    event.cert_color_scheme = normalize_cert_scalar_for_storage(form.get("cert_color_scheme"))
    event.cert_style = normalize_cert_scalar_for_storage(form.get("cert_style"))
    db.commit()

    flash(request, "Certificate template saved.", "success")
    return RedirectResponse(f"/admin/events/{event_id}/edit?step=5", status_code=303)


def _apply_certificate_template_to_event(event: Event, tpl: CertificateTemplate) -> None:
    event.cert_template_id = tpl.id
    event.cert_style = tpl.cert_style
    event.cert_title = merge_cert_scalar_from_template(tpl.cert_title, event.cert_title)
    event.cert_subtitle = merge_cert_scalar_from_template(tpl.cert_subtitle, event.cert_subtitle)
    event.cert_footer = merge_cert_scalar_from_template(tpl.cert_footer, event.cert_footer)
    event.cert_signer_name = merge_cert_scalar_from_template(tpl.cert_signer_name, event.cert_signer_name)
    event.cert_signer_designation = merge_cert_scalar_from_template(
        tpl.cert_signer_designation, event.cert_signer_designation
    )
    event.cert_logo_url = merge_cert_scalar_from_template(tpl.cert_logo_url, event.cert_logo_url)
    event.cert_bg_url = merge_cert_scalar_from_template(tpl.cert_bg_url, event.cert_bg_url)
    event.cert_signature_url = merge_cert_scalar_from_template(
        tpl.cert_signature_url, event.cert_signature_url
    )
    event.cert_color_scheme = merge_cert_scalar_from_template(
        tpl.cert_color_scheme, event.cert_color_scheme
    )


@router.post("/events/{event_id}/certificate/apply-template")
async def event_certificate_apply_template(
    request: Request, event_id: int, db: Session = Depends(get_db)
):
    """Copy a saved certificate template onto the event and record cert_template_id."""
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    event = db.query(Event).get(event_id)
    if not event:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    form = await request.form()
    raw_id = (form.get("template_id") or "").strip()
    if not raw_id.isdigit():
        flash(request, "Select a certificate template.", "warning")
        return RedirectResponse(f"/admin/events/{event_id}/edit?step=5", status_code=303)

    tpl = db.query(CertificateTemplate).get(int(raw_id))
    if not tpl:
        flash(request, "Certificate template not found.", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/edit?step=5", status_code=303)

    _apply_certificate_template_to_event(event, tpl)
    log_activity(
        db,
        category="admin",
        action="update",
        description=f"Applied certificate template '{tpl.name}' to event '{event.name}'",
        request=request,
        user_id=admin.id,
        target_type="event",
        target_id=event.id,
    )
    db.commit()
    flash(request, f"Applied template “{tpl.name}”.", "success")
    return RedirectResponse(f"/admin/events/{event_id}/edit?step=5", status_code=303)


@router.get("/events/{event_id}/certificate/designer")
def certificate_designer_page(
    request: Request, event_id: int, db: Session = Depends(get_db)
):
    from app.services.certificate import (
        _raw_cert_style_dict,
        default_freeform_cert_style_dict,
        should_render_certificate_as_freeform,
        _parse_cert_style,
    )

    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    raw = _raw_cert_style_dict(event)
    if should_render_certificate_as_freeform(raw):
        initial_style = raw
    else:
        initial_style = default_freeform_cert_style_dict()
        sty = _parse_cert_style(event)
        initial_style["border_style"] = sty.get("border_style", initial_style["border_style"])
        initial_style["border_width"] = float(sty.get("border_width", initial_style["border_width"]))
        initial_style["bg_size"] = sty.get("bg_size", initial_style["bg_size"])
        initial_style["bg_offset_x"] = float(sty.get("bg_offset_x", 0))
        initial_style["bg_offset_y"] = float(sty.get("bg_offset_y", 0))
        for k, sk in (
            ("border_color_primary", "border_color_primary"),
            ("border_color_secondary", "border_color_secondary"),
            ("border_color_tertiary", "border_color_tertiary"),
        ):
            if sty.get(sk):
                initial_style[k] = sty[sk]

    if not str(initial_style.get("background_image_url") or "").strip():
        bg = getattr(event, "cert_bg_url", None) or None
        if bg:
            initial_style = dict(initial_style)
            initial_style["background_image_url"] = bg

    eid = event.id
    _save = f"/admin/events/{eid}/certificate/designer"
    _leg = f"/admin/events/{eid}/certificate/legacy-to-freeform"
    designer_bootstrap = {
        "eventId": eid,
        "templateId": None,
        "eventName": event.name or "",
        "isPersistedFreeform": should_render_certificate_as_freeform(raw),
        "initialStyle": initial_style,
        "saveUrl": _save,
        "legacyConvertUrl": _leg,
        "designerSaveUrl": _save,
        "designerLegacyUrl": _leg,
        "aiGenerateUrl": "/admin/certificate-ai/generate-template",
    }
    return templates.TemplateResponse(
        "admin/certificate_designer.html",
        _admin_ctx(
            request,
            active_page="events",
            event=event,
            cert_template=None,
            designer_bootstrap=designer_bootstrap,
        ),
    )


@router.post("/events/{event_id}/certificate/designer")
async def certificate_designer_save(
    request: Request, event_id: int, db: Session = Depends(get_db)
):
    ct = (request.headers.get("content-type") or "").lower()
    admin = _require_admin(request, db)
    if not admin:
        if "application/json" in ct:
            return JSONResponse({"error": "Unauthorized"}, status_code=401)
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        if "application/json" in ct:
            return JSONResponse({"error": "Not found"}, status_code=404)
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    if "application/json" in ct:
        body = await request.json()
        payload = body.get("cert_style")
        if isinstance(payload, dict):
            event.cert_style = json.dumps(payload)
        elif payload is None or payload == "":
            event.cert_style = None
        else:
            event.cert_style = str(payload).strip() or None
        log_activity(
            db,
            category="admin",
            action="update",
            description=f"Updated certificate visual layout for event '{event.name}'",
            request=request,
            user_id=admin.id,
            target_type="event",
            target_id=event.id,
        )
        db.commit()
        return JSONResponse({"ok": True})
    form = await request.form()
    event.cert_style = (form.get("cert_style") or "").strip() or None
    log_activity(
        db,
        category="admin",
        action="update",
        description=f"Updated certificate visual layout for event '{event.name}'",
        request=request,
        user_id=admin.id,
        target_type="event",
        target_id=event.id,
    )
    db.commit()
    flash(request, "Certificate layout saved.", "success")
    return RedirectResponse(f"/admin/events/{event_id}/certificate/designer", status_code=303)


@router.get("/events/{event_id}/certificate/legacy-to-freeform")
def certificate_legacy_to_freeform(
    request: Request, event_id: int, db: Session = Depends(get_db)
):
    from app.services.certificate import legacy_cert_style_to_freeform_dict

    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)
    event = db.query(Event).get(event_id)
    if not event:
        return JSONResponse({"error": "Not found"}, status_code=404)
    return JSONResponse(legacy_cert_style_to_freeform_dict(event))


@router.post("/events/certificate/preview-image")
async def event_certificate_preview_image(
    request: Request, db: Session = Depends(get_db)
):
    """Generate a PNG thumbnail of the certificate from live form values."""
    import io as _io
    from types import SimpleNamespace
    import pypdfium2
    from app.services.certificate import (
        generate_certificate_pdf,
        should_render_certificate_as_freeform,
        _raw_cert_style_dict,
    )
    from fastapi.responses import Response

    admin = _require_admin(request, db)
    if not admin:
        return Response(status_code=403)

    form = await request.form()

    aud_id = form.get("auditorium_id", "")
    auditorium = None
    if aud_id and aud_id.strip().isdigit():
        auditorium = db.query(Auditorium).get(int(aud_id))

    form_cert_style = normalize_cert_scalar_for_storage(form.get("cert_style"))
    cert_style_for_pdf = form_cert_style
    raw_from_form = _raw_cert_style_dict(SimpleNamespace(cert_style=form_cert_style or ""))
    if not should_render_certificate_as_freeform(raw_from_form):
        peid = (form.get("preview_event_id") or "").strip()
        if peid.isdigit():
            ev_row = db.query(Event).get(int(peid))
            if ev_row and ev_row.cert_style:
                raw_db = _raw_cert_style_dict(ev_row)
                if should_render_certificate_as_freeform(raw_db):
                    cert_style_for_pdf = ev_row.cert_style

    draft_event = SimpleNamespace(
        name=form.get("name", "").strip() or "Event Title",
        start_date=date.today(),
        cert_title=normalize_cert_scalar_for_storage(form.get("cert_title")),
        cert_subtitle=normalize_cert_scalar_for_storage(form.get("cert_subtitle")),
        cert_footer=normalize_cert_scalar_for_storage(form.get("cert_footer")),
        cert_signer_name=normalize_cert_scalar_for_storage(form.get("cert_signer_name")),
        cert_signer_designation=normalize_cert_scalar_for_storage(form.get("cert_signer_designation")),
        cert_signature_url=normalize_cert_scalar_for_storage(form.get("cert_signature_url")),
        cert_logo_url=normalize_cert_scalar_for_storage(form.get("cert_logo_url")),
        cert_bg_url=normalize_cert_scalar_for_storage(form.get("cert_bg_url")),
        cert_color_scheme=normalize_cert_scalar_for_storage(form.get("cert_color_scheme")),
        cert_style=cert_style_for_pdf,
    )
    dummy_booking = SimpleNamespace(
        booking_ref="PREVIEW",
        qr_code_data="CERT-PREVIEW-SAMPLE",
    )
    dummy_user = SimpleNamespace(
        full_name="Sample Attendee",
        username="sample_attendee",
    )

    pdf_bytes = generate_certificate_pdf(dummy_booking, dummy_user, draft_event, draft_event, auditorium)

    pdf_doc = pypdfium2.PdfDocument(pdf_bytes)
    page = pdf_doc[0]
    scale = 1.5
    bitmap = page.render(scale=scale)
    pil_image = bitmap.to_pil()
    png_buf = _io.BytesIO()
    pil_image.save(png_buf, format="PNG", optimize=True)
    page.close()
    pdf_doc.close()

    return Response(
        content=png_buf.getvalue(),
        media_type="image/png",
        headers={"Cache-Control": "no-store"},
    )



# Session-level recordings routes removed — managed via Event Management Hub


# ─── Bookings ───

@router.get("/bookings")
def bookings_list(
    request: Request,
    db: Session = Depends(get_db),
    q: str = Query("", alias="q"),
    status_filter: str = Query("", alias="status"),
    event_filter: str = Query("", alias="event_id"),
    page: int = Query(1, ge=1),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    query = (
        db.query(Booking)
        .outerjoin(User, Booking.user_id == User.id)
        .outerjoin(Event, Booking.event_id == Event.id)
    )
    if status_filter:
        query = query.filter(Booking.payment_status == status_filter)
    else:
        query = query.filter(Booking.payment_status.in_(["paid", "hold", "refunded"]))

    if event_filter:
        try:
            query = query.filter(Booking.event_id == int(event_filter))
        except ValueError:
            pass

    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Event.name.ilike(like),
                Booking.booking_ref.ilike(like),
                Booking.ticket_id.ilike(like),
            )
        )

    total_count = query.count()
    total_pages = max(1, (total_count + ADMIN_PAGE_SIZE - 1) // ADMIN_PAGE_SIZE)
    page = min(page, total_pages)

    bookings = (
        query.order_by(Booking.booked_at.desc())
        .offset((page - 1) * ADMIN_PAGE_SIZE)
        .limit(ADMIN_PAGE_SIZE)
        .all()
    )

    enriched = []
    for b in bookings:
        u = db.query(User).get(b.user_id)
        event = db.query(Event).get(b.event_id) if b.event_id else None
        seat = db.query(Seat).get(b.seat_id)
        enriched.append({"booking": b, "user": u, "event": event, "seat": seat})

    all_events = db.query(Event).order_by(Event.name).all()

    return templates.TemplateResponse(
        "admin/bookings.html",
        _admin_ctx(
            request,
            active_page="bookings",
            bookings=enriched,
            q=q,
            status_filter=status_filter,
            session_filter=event_filter,
            all_events=all_events,
            page=page,
            total_pages=total_pages,
            total_count=total_count,
        ),
    )


@router.post("/bookings/backfill-qr-urls")
async def bookings_backfill_qr_urls(request: Request, db: Session = Depends(get_db)):
    """One-time style fix: set qr_code_data to public verification URL for bookings still using legacy payloads."""
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    base = settings.base_url.rstrip("/") if getattr(settings, "base_url", None) else ""
    if not base:
        flash(request, "base_url is not configured in settings.", "danger")
        return RedirectResponse("/admin/bookings", status_code=303)
    updated = 0
    for b in db.query(Booking).all():
        if not b.ticket_id:
            continue
        want = f"{base}/certificate/verify/{b.ticket_id}"
        cur = (b.qr_code_data or "").strip()
        if cur.startswith("http"):
            continue
        if cur == want:
            continue
        b.qr_code_data = want
        updated += 1
    db.commit()
    flash(request, f"Updated {updated} booking(s) with certificate verification URL for QR.", "success")
    return RedirectResponse("/admin/bookings", status_code=303)


def _csv_safe(val, default=""):
    if val is None:
        return default
    s = str(val).replace("\r", " ").replace("\n", " ").replace('"', '""')
    return s


@router.get("/bookings/export")
def bookings_csv(
    request: Request,
    db: Session = Depends(get_db),
    q: str = Query("", alias="q"),
    status_filter: str = Query("", alias="status"),
    event_filter: str = Query("", alias="event_id"),
    use_filters: str = Query("0", alias="use_filters"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    query = db.query(Booking)
    apply_filters = use_filters.strip().lower() in ("1", "true", "yes")
    if apply_filters:
        if status_filter:
            query = query.filter(Booking.payment_status == status_filter)
        else:
            query = query.filter(Booking.payment_status.in_(["paid", "hold", "refunded"]))
        if event_filter:
            try:
                query = query.filter(Booking.event_id == int(event_filter))
            except ValueError:
                pass
    else:
        query = query.filter(Booking.payment_status.in_(["paid", "hold", "refunded"]))

    bookings = query.order_by(Booking.booked_at.desc()).all()

    header = [
        "Booking Ref", "Ticket ID", "Event", "Seat", "Status", "Amount Paid", "Refund Amount",
        "Booked At", "Checked In", "Refund Status",
        "user_id", "user_email", "user_username", "user_full_name", "user_phone",
        "user_college", "user_discipline", "user_domain", "user_year_of_study",
        "user_is_admin", "user_is_supervisor", "user_supervisor_college", "user_created_at",
        "user_oauth_provider", "user_oauth_id",
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(header)
    for b in bookings:
        u = db.query(User).get(b.user_id)
        event = db.query(Event).get(b.event_id) if b.event_id else None
        seat = db.query(Seat).get(b.seat_id)
        if apply_filters and q:
            search = q.lower()
            match = (
                (u and (search in (u.username or "").lower() or search in (u.email or "").lower() or (u.full_name and search in u.full_name.lower())))
                or (event and search in event.name.lower())
                or (b.booking_ref and search in b.booking_ref.lower())
                or (b.ticket_id and search in b.ticket_id.lower())
            )
            if not match:
                continue
        sup_college = (u.supervised_college.name if u and u.supervised_college else "") or ""
        writer.writerow([
            b.booking_ref or "",
            b.ticket_id or "",
            event.name if event else "",
            seat.label if seat else "",
            b.payment_status or "",
            b.amount_paid if b.amount_paid is not None else "",
            b.refund_amount if b.refund_amount is not None else "",
            b.booked_at.strftime("%Y-%m-%d %H:%M") if b.booked_at else "",
            "Yes" if b.checked_in else "No",
            b.refund_status or "",
            u.id if u else "",
            _csv_safe(u.email if u else None),
            _csv_safe(u.username if u else None),
            _csv_safe(u.full_name if u else None),
            _csv_safe(u.phone if u else None),
            _csv_safe(u.college if u else None),
            _csv_safe(u.discipline if u else None),
            _csv_safe(u.domain if u else None),
            u.year_of_study if u and u.year_of_study is not None else "",
            "1" if u and u.is_admin else "0",
            "1" if u and u.is_supervisor else "0",
            sup_college,
            u.created_at.strftime("%Y-%m-%d %H:%M:%S") if u and u.created_at else "",
            _csv_safe(u.oauth_provider if u else None),
            _csv_safe(u.oauth_id if u else None),
        ])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=techtrek_bookings.csv"},
    )


@router.post("/bookings/{booking_id}/cancel")
def booking_cancel(request: Request, booking_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    b = db.query(Booking).get(booking_id)
    if not b:
        flash(request, "Booking not found.", "danger")
        return RedirectResponse("/admin/bookings", status_code=303)
    if b.payment_status not in ("paid", "hold"):
        flash(request, f"Booking {b.booking_ref} is '{b.payment_status}' — cannot cancel.", "danger")
        return RedirectResponse("/admin/bookings", status_code=303)
    if b.checked_in:
        flash(request, f"Booking {b.booking_ref} is checked in — cancel not allowed.", "danger")
        return RedirectResponse("/admin/bookings", status_code=303)
    b.payment_status = "cancelled"
    log_activity(db, category="admin", action="cancel", description=f"Admin cancelled booking {b.booking_ref}", request=request, user_id=admin.id, target_type="booking", target_id=booking_id)
    db.commit()
    flash(request, f"Booking {b.booking_ref} cancelled.", "success")
    return RedirectResponse("/admin/bookings", status_code=303)


@router.get("/bookings/{booking_id}/invoice")
def admin_booking_invoice(request: Request, booking_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    booking = db.query(Booking).get(booking_id)
    if not booking:
        flash(request, "Booking not found.", "danger")
        return RedirectResponse("/admin/bookings", status_code=303)

    group_bookings = (
        db.query(Booking)
        .filter(
            Booking.booking_group == booking.booking_group,
            Booking.payment_status.in_(["paid", "refunded"]),
        )
        .all()
    ) if booking.booking_group else [booking]
    if not group_bookings:
        group_bookings = [booking]

    user = db.query(User).get(booking.user_id)
    event = db.query(Event).get(booking.event_id) if booking.event_id else None
    auditorium = db.query(Auditorium).get(event.auditorium_id) if event and event.auditorium_id else None
    if not user or not event or not auditorium:
        flash(request, "Related data not found.", "danger")
        return RedirectResponse("/admin/bookings", status_code=303)

    seats = [db.query(Seat).get(b.seat_id) for b in group_bookings]
    custom_types_map = {f"custom_{st.id}": st for st in db.query(SeatType).filter(SeatType.is_custom == True).all()}
    pdf_bytes = generate_invoice_pdf(group_bookings, user, event, auditorium, seats, custom_types_map, db=db)
    ref = booking.booking_ref or "invoice"
    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="invoice-{ref}.pdf"'},
    )


@router.post("/bookings/{booking_id}/refund")
def booking_refund(request: Request, booking_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    b = db.query(Booking).get(booking_id)
    if b and b.checked_in:
        flash(request, f"Booking {b.booking_ref} is checked in — refund not allowed.", "danger")
        return RedirectResponse("/admin/bookings", status_code=303)
    can_refund = b and b.payment_status in ("paid", "refunded") and (
        b.payment_status == "paid" or b.refund_status == "failed"
    )
    if can_refund:
        price = b.amount_paid or 0
        refund_amount = float(price)

        rz_ok = True
        rz_result = None
        if b.razorpay_payment_id and refund_amount > 0:
            rz_result = rz_process_refund(b.razorpay_payment_id, int(refund_amount * 100))
            rz_ok = rz_result is not None

        b.payment_status = "refunded"
        b.refund_amount = refund_amount
        if rz_result and isinstance(rz_result, dict):
            b.refund_id = rz_result.get("id")
            b.refund_status = "initiated"
        elif rz_ok:
            b.refund_status = "completed"
            b.refund_processed_at = now_ist()
        else:
            b.refund_status = "failed"

        log_activity(db, category="admin", action="refund", description=f"Admin refunded booking {b.booking_ref} (₹{refund_amount:.0f})", request=request, user_id=admin.id, target_type="booking", target_id=booking_id)
        db.commit()

        if rz_ok:
            flash(request, f"Booking {b.booking_ref} refunded (₹{refund_amount:.0f}).", "success")
        else:
            flash(request, f"Booking {b.booking_ref} marked refunded but Razorpay API call failed — process the ₹{refund_amount:.0f} refund manually.", "warning")
    elif b and b.payment_status != "paid":
        flash(request, f"Booking {b.booking_ref} is not in 'paid' status — cannot refund.", "danger")
    return RedirectResponse("/admin/bookings", status_code=303)


# ─── Check-in ───

@router.get("/checkin")
def checkin_page(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    events_list = (
        db.query(Event)
        .filter(Event.status.in_(["published", "completed"]))
        .order_by(Event.start_date.desc())
        .all()
    )
    return templates.TemplateResponse(
        "admin/checkin.html",
        _admin_ctx(request, active_page="checkin", events_list=events_list, result=None),
    )


@router.post("/checkin")
async def checkin_verify(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    form = await _form(request)
    ticket_id = form.get("ticket_id", "").strip()
    event_id_raw = form.get("event_id", "")

    events_list = (
        db.query(Event)
        .filter(Event.status.in_(["published", "completed"]))
        .order_by(Event.start_date.desc())
        .all()
    )

    if not ticket_id:
        return templates.TemplateResponse(
            "admin/checkin.html",
            _admin_ctx(request, active_page="checkin", events_list=events_list, result={"status": "error", "msg": "Please enter a ticket ID."}),
        )

    is_group = ticket_id.startswith("GROUP-")

    if is_group:
        group_id = ticket_id[6:]
        all_group_any_status = db.query(Booking).filter(
            Booking.booking_group == group_id,
        ).all()
        all_group = [b for b in all_group_any_status if b.payment_status == "paid"]
        refunded_count = sum(1 for b in all_group_any_status if b.payment_status in ("refunded", "cancelled"))

        result = None
        group_bookings = []

        if not all_group_any_status:
            result = {"status": "error", "msg": f"Group '{group_id}' not found."}
        elif not all_group:
            result = {"status": "error", "msg": f"No valid (paid) tickets in this group — {refunded_count} ticket(s) are refunded/cancelled."}
        elif event_id_raw:
            try:
                group_bookings = [b for b in all_group if b.event_id == int(event_id_raw)]
            except ValueError:
                group_bookings = all_group
            if not group_bookings:
                result = {"status": "error", "msg": f"No tickets in this group match the selected event."}
        else:
            group_bookings = all_group

        if group_bookings:
            now = now_ist()
            newly_checked = []
            already_checked = []
            for gb in group_bookings:
                if gb.checked_in:
                    seat = db.query(Seat).get(gb.seat_id)
                    already_checked.append(seat.label if seat else gb.ticket_id)
                else:
                    gb.checked_in = True
                    gb.checked_in_at = now
                    seat = db.query(Seat).get(gb.seat_id)
                    newly_checked.append(seat.label if seat else gb.ticket_id)
            db.commit()

            user = db.query(User).get(group_bookings[0].user_id)
            event = db.query(Event).get(group_bookings[0].event_id) if group_bookings[0].event_id else None
            event_name = event.name if event else "unknown"
            refunded_note = f" ({refunded_count} ticket(s) in this group are refunded/cancelled.)" if refunded_count else ""

            if newly_checked and not already_checked:
                msg = f"Check-in successful! {len(newly_checked)} ticket(s) for '{event_name}'.{refunded_note}"
                status = "success"
                log_activity(db, category="admin", action="checkin", description=f"Group check-in: {len(newly_checked)} ticket(s) for '{event_name}'", request=request, user_id=admin.id, target_type="booking", target_id=group_bookings[0].id)
            elif newly_checked and already_checked:
                msg = f"Checked in {len(newly_checked)} ticket(s). {len(already_checked)} already checked in.{refunded_note}"
                status = "success"
                log_activity(db, category="admin", action="checkin", description=f"Partial group check-in: {len(newly_checked)} new for '{event_name}'", request=request, user_id=admin.id, target_type="booking", target_id=group_bookings[0].id)
            else:
                msg = f"Re-entry — all {len(already_checked)} ticket(s) already checked in. Ticket is valid.{refunded_note}"
                status = "reentry"

            result = {
                "status": status,
                "msg": msg,
                "is_group": True,
                "user_name": user.full_name or user.username if user else "Unknown",
                "user_email": user.email if user else "",
                "event_name": event_name,
                "newly_checked": newly_checked,
                "already_checked": already_checked,
                "refunded_count": refunded_count,
            }
    else:
        query = db.query(Booking).filter(Booking.ticket_id == ticket_id, Booking.payment_status == "paid")
        if event_id_raw:
            try:
                query = query.filter(Booking.event_id == int(event_id_raw))
            except ValueError:
                pass
        booking = query.first()

        if not booking:
            result = {"status": "error", "msg": f"Ticket '{ticket_id}' not found or not valid."}
        elif booking.checked_in:
            user = db.query(User).get(booking.user_id)
            seat = db.query(Seat).get(booking.seat_id)
            event = db.query(Event).get(booking.event_id) if booking.event_id else None
            time_str = booking.checked_in_at.strftime('%I:%M %p') if booking.checked_in_at else 'earlier'
            result = {
                "status": "reentry",
                "msg": f"Re-entry — ticket valid. Originally checked in at {time_str}.",
                "user_name": user.full_name or user.username if user else "Unknown",
                "user_email": user.email if user else "",
                "seat_label": seat.label if seat else "",
                "event_name": event.name if event else "",
                "ticket_id": ticket_id,
            }
        else:
            booking.checked_in = True
            booking.checked_in_at = now_ist()
            user = db.query(User).get(booking.user_id)
            seat = db.query(Seat).get(booking.seat_id)
            event = db.query(Event).get(booking.event_id) if booking.event_id else None
            event_name = event.name if event else "unknown"
            log_activity(db, category="admin", action="checkin", description=f"Checked in ticket '{ticket_id}' (seat {seat.label if seat else '?'}) for '{event_name}'", request=request, user_id=admin.id, target_type="booking", target_id=booking.id)
            db.commit()
            result = {
                "status": "success",
                "msg": "Check-in successful!",
                "user_name": user.full_name or user.username if user else "Unknown",
                "user_email": user.email if user else "",
                "seat_label": seat.label if seat else "",
                "event_name": event_name,
                "ticket_id": ticket_id,
            }

    stats = None
    if event_id_raw:
        try:
            eid = int(event_id_raw)
            total_booked = db.query(func.count(Booking.id)).filter(Booking.event_id == eid, Booking.payment_status == "paid").scalar()
            checked_in_count = db.query(func.count(Booking.id)).filter(Booking.event_id == eid, Booking.payment_status == "paid", Booking.checked_in == True).scalar()
            stats = {"total": total_booked, "checked_in": checked_in_count}
        except ValueError:
            pass

    return templates.TemplateResponse(
        "admin/checkin.html",
        _admin_ctx(request, active_page="checkin", events_list=events_list, result=result, stats=stats, selected_event=event_id_raw),
    )


# ─── Waitlist ───

@router.get("/waitlist")
def waitlist_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    entries = db.query(Waitlist).order_by(Waitlist.joined_at.desc()).all()
    enriched = []
    events_map = {}
    for w in entries:
        u = db.query(User).get(w.user_id)
        event = db.query(Event).get(w.event_id) if w.event_id else None
        enriched.append({"entry": w, "user": u, "event": event})
        if event:
            if event.id not in events_map:
                events_map[event.id] = {"event": event, "count": 0}
            events_map[event.id]["count"] += 1

    waitlist_events = sorted(events_map.values(), key=lambda x: x["count"], reverse=True)
    all_events = db.query(Event).filter(Event.status == "published").order_by(Event.name).all()

    return templates.TemplateResponse(
        "admin/waitlist.html",
        _admin_ctx(
            request,
            active_page="waitlist",
            entries=enriched,
            waitlist_events=waitlist_events,
            all_events=all_events,
        ),
    )


@router.post("/waitlist/notify-about-event")
async def waitlist_notify_about_event(
    request: Request,
    db: Session = Depends(get_db),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    form = await _form(request)
    source_event_id = int(form.get("source_event_id", 0))
    target_event_id = int(form.get("target_event_id", 0))

    source_event = db.query(Event).get(source_event_id) if source_event_id else None
    target_event = db.query(Event).get(target_event_id) if target_event_id else None

    if not source_event or not target_event:
        flash(request, "Invalid event selection.", "danger")
        return RedirectResponse("/admin/waitlist", status_code=303)

    entries = db.query(Waitlist).filter(Waitlist.event_id == source_event_id).all()
    if not entries:
        flash(request, "No waitlist entries for that event.", "warning")
        return RedirectResponse("/admin/waitlist", status_code=303)

    from app.services.email import send_waitlist_notification

    base = settings.base_url.rstrip("/") if hasattr(settings, "base_url") and settings.base_url else ""
    event_url = f"{base}/events/{target_event_id}"
    sent = 0
    for w in entries:
        user = db.query(User).get(w.user_id)
        if user and user.email:
            send_waitlist_notification(
                user.email, user.username,
                source_event.name, target_event.name, event_url,
            )
            sent += 1

    log_activity(
        db, category="admin", action="waitlist_notify",
        description=f"Notified {sent} waitlisted user(s) from '{source_event.name}' about '{target_event.name}'",
        request=request, user_id=admin.id,
        target_type="event", target_id=target_event_id,
    )
    db.commit()

    flash(request, f"Notified {sent} waitlisted user(s) about '{target_event.name}'.", "success")
    return RedirectResponse("/admin/waitlist", status_code=303)


@router.post("/waitlist/{entry_id}/delete")
def waitlist_delete(request: Request, entry_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    entry = db.query(Waitlist).get(entry_id)
    if entry:
        log_activity(db, category="admin", action="delete", description=f"Removed waitlist entry #{entry_id}", request=request, user_id=admin.id, target_type="waitlist", target_id=entry_id)
        db.delete(entry)
        db.commit()
        flash(request, "Waitlist entry removed.", "success")
    return RedirectResponse("/admin/waitlist", status_code=303)


# ─── Users ───

@router.get("/users")
def users_list(
    request: Request,
    db: Session = Depends(get_db),
    q: str = Query("", alias="q"),
    page: int = Query(1, ge=1),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    query = db.query(User).filter(User.deleted_at.is_(None))
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(User.college.ilike(like))
    total_count = query.count()
    total_pages = max(1, (total_count + ADMIN_PAGE_SIZE - 1) // ADMIN_PAGE_SIZE)
    page = min(page, total_pages)
    users = (
        query.order_by(User.created_at.desc())
        .offset((page - 1) * ADMIN_PAGE_SIZE)
        .limit(ADMIN_PAGE_SIZE)
        .all()
    )
    colleges = db.query(College).filter(College.is_active == True).order_by(College.name).all()
    return templates.TemplateResponse(
        "admin/users.html",
        _admin_ctx(
            request,
            active_page="users",
            users=users,
            q=q,
            colleges=colleges,
            page=page,
            total_pages=total_pages,
            total_count=total_count,
        ),
    )


@router.get("/users/export")
def users_csv(
    request: Request,
    db: Session = Depends(get_db),
    q: str = Query("", alias="q"),
    use_filters: str = Query("0", alias="use_filters"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    query = db.query(User)
    apply_filters = use_filters.strip().lower() in ("1", "true", "yes")
    search_term = q.strip().lower() if apply_filters and q.strip() else ""
    if search_term:
        like = f"%{search_term}%"
        query = query.filter(User.college.ilike(like))
    users = query.order_by(User.created_at.desc()).all()
    if search_term:
        filtered = []
        for u in users:
            if any(
                search_term in (getattr(u, f, "") or "").lower()
                for f in ("username", "email", "full_name", "college")
            ):
                filtered.append(u)
        users = filtered

    def _safe(u, attr, default=""):
        v = getattr(u, attr, None)
        if v is None:
            return default
        return str(v).replace("\r", " ").replace("\n", " ").replace('"', '""')

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "id", "email", "username", "full_name", "phone", "college", "discipline", "domain",
        "year_of_study", "is_admin", "is_supervisor", "supervisor_college", "created_at",
        "oauth_provider", "oauth_id",
    ])
    for u in users:
        sup_college = ""
        if u.supervised_college:
            sup_college = u.supervised_college.name or ""
        writer.writerow([
            u.id,
            _safe(u, "email"),
            _safe(u, "username"),
            _safe(u, "full_name"),
            _safe(u, "phone"),
            _safe(u, "college"),
            _safe(u, "discipline"),
            _safe(u, "domain"),
            u.year_of_study if u.year_of_study is not None else "",
            "1" if u.is_admin else "0",
            "1" if u.is_supervisor else "0",
            sup_college,
            u.created_at.strftime("%Y-%m-%d %H:%M:%S") if u.created_at else "",
            _safe(u, "oauth_provider"),
            _safe(u, "oauth_id"),
        ])
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=techtrek_users.csv"},
    )


@router.post("/users/{user_id}/toggle-admin")
def toggle_admin(request: Request, user_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    u = db.query(User).get(user_id)
    if u and u.id != admin.id:
        u.is_admin = not u.is_admin
        status = "admin" if u.is_admin else "regular user"
        log_activity(db, category="admin", action="role_change", description=f"Changed {u.username} role to {status}", request=request, user_id=admin.id, target_type="user", target_id=user_id)
        db.commit()
        flash(request, f"{u.username} is now a {status}.", "success")
    return RedirectResponse("/admin/users", status_code=303)


@router.post("/users/{user_id}/toggle-supervisor")
async def toggle_supervisor(request: Request, user_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    u = db.query(User).get(user_id)
    if not u or u.id == admin.id:
        return RedirectResponse("/admin/users", status_code=303)

    if u.is_supervisor:
        u.is_supervisor = False
        u.supervisor_college_id = None
        log_activity(db, category="admin", action="role_change", description=f"Removed supervisor role from {u.username}", request=request, user_id=admin.id, target_type="user", target_id=user_id)
        db.commit()
        flash(request, f"{u.username} is no longer a supervisor.", "success")
    else:
        form = await request.form()
        college_id_raw = form.get("college_id", "")
        if not college_id_raw:
            flash(request, "Please select a college to assign the supervisor to.", "danger")
            return RedirectResponse("/admin/users", status_code=303)
        try:
            cid = int(college_id_raw)
        except ValueError:
            flash(request, "Invalid college selection.", "danger")
            return RedirectResponse("/admin/users", status_code=303)
        college = db.query(College).get(cid)
        if not college:
            flash(request, "Selected college not found.", "danger")
            return RedirectResponse("/admin/users", status_code=303)
        u.is_supervisor = True
        u.supervisor_college_id = cid
        log_activity(db, category="admin", action="role_change", description=f"Made {u.username} supervisor at {college.name}", request=request, user_id=admin.id, target_type="user", target_id=user_id)
        db.commit()
        flash(request, f"{u.username} is now a supervisor at {college.name}.", "success")
    return RedirectResponse("/admin/users", status_code=303)


@router.post("/users/{user_id}/delete")
async def user_delete(request: Request, user_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    u = db.query(User).get(user_id)
    if not u or u.id == admin.id:
        flash(request, "Cannot delete this user.", "danger")
        return RedirectResponse("/admin/users", status_code=303)

    form = await request.form()
    delete_type = form.get("delete_type", "soft")

    if delete_type == "hard":
        db.query(Booking).filter(Booking.user_id == user_id).delete()
        db.query(Feedback).filter(Feedback.user_id == user_id).delete()
        log_activity(db, category="admin", action="delete",
                     description=f"Hard-deleted user '{u.username}' (id={user_id})",
                     request=request, user_id=admin.id, target_type="user", target_id=user_id)
        db.delete(u)
        db.commit()
        flash(request, f"User '{u.username}' permanently deleted.", "success")
    else:
        u.deleted_at = now_ist()
        log_activity(db, category="admin", action="deactivate",
                     description=f"Soft-deleted user '{u.username}'",
                     request=request, user_id=admin.id, target_type="user", target_id=user_id)
        db.commit()
        flash(request, f"User '{u.username}' deactivated.", "success")

    return RedirectResponse("/admin/users", status_code=303)


# ─── Schedule (Admin) ───

@router.get("/schedule")
def admin_schedule(
    request: Request,
    db: Session = Depends(get_db),
    college_id: str = Query("", alias="college_id"),
    auditorium_id: str = Query("", alias="auditorium_id"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    query = db.query(Event).filter(
        Event.status.in_(["published", "completed"])
    )
    if auditorium_id:
        try:
            query = query.filter(Event.auditorium_id == int(auditorium_id))
        except ValueError:
            pass
    elif college_id:
        try:
            aud_ids = [a.id for a in db.query(Auditorium.id).filter(Auditorium.college_id == int(college_id)).all()]
            if aud_ids:
                query = query.filter(Event.auditorium_id.in_(aud_ids))
            else:
                query = query.filter(False)
        except ValueError:
            pass

    events = query.order_by(Event.start_date).all()

    grouped = defaultdict(list)
    for ev in events:
        aud = db.query(Auditorium).get(ev.auditorium_id) if ev.auditorium_id else None
        date_key = ev.start_date.strftime("%Y-%m-%d") if ev.start_date else "TBD"
        grouped[date_key].append({"event": ev, "auditorium": aud})

    colleges = db.query(College).order_by(College.name).all()
    auditoriums = db.query(Auditorium).order_by(Auditorium.name).all()

    return templates.TemplateResponse(
        "admin/schedule_admin.html",
        _admin_ctx(
            request, active_page="schedule",
            grouped=dict(sorted(grouped.items())),
            colleges=colleges, auditoriums=auditoriums,
            college_id=college_id, auditorium_id=auditorium_id,
        ),
    )


# ─── Activity Log ───

@router.get("/activity-log")
def activity_log_page(
    request: Request,
    db: Session = Depends(get_db),
    category: str = Query("", alias="category"),
    q: str = Query("", alias="q"),
    date_from: str = Query("", alias="date_from"),
    date_to: str = Query("", alias="date_to"),
    page: int = Query(1, alias="page", ge=1),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    query = db.query(ActivityLog)

    if category:
        query = query.filter(ActivityLog.category == category)
    if q:
        search = f"%{q}%"
        query = query.filter(
            ActivityLog.description.ilike(search)
            | ActivityLog.action.ilike(search)
        )
    if date_from:
        try:
            dt_from = datetime.fromisoformat(date_from)
            query = query.filter(ActivityLog.timestamp >= dt_from)
        except ValueError:
            pass
    if date_to:
        try:
            dt_to = datetime.fromisoformat(date_to)
            dt_to = dt_to.replace(hour=23, minute=59, second=59)
            query = query.filter(ActivityLog.timestamp <= dt_to)
        except ValueError:
            pass

    total = query.count()
    total_pages = max(1, (total + ADMIN_PAGE_SIZE - 1) // ADMIN_PAGE_SIZE)
    page = min(page, total_pages)

    logs = (
        query.order_by(ActivityLog.timestamp.desc())
        .offset((page - 1) * ADMIN_PAGE_SIZE)
        .limit(ADMIN_PAGE_SIZE)
        .all()
    )

    return templates.TemplateResponse(
        "admin/activity_log.html",
        _admin_ctx(
            request,
            active_page="activity_log",
            logs=logs,
            total=total,
            page=page,
            total_pages=total_pages,
            filter_category=category,
            filter_q=q,
            filter_date_from=date_from,
            filter_date_to=date_to,
        ),
    )


# ---------------------------------------------------------------------------
# Platform Settings
# ---------------------------------------------------------------------------

_SETTING_KEYS = [
    "platform_logo_url", "company_name", "company_address",
    "company_email", "company_phone", "company_gstin", "company_pan", "gst_rate",
]


def _load_settings(db: Session) -> dict:
    from app.config import settings as cfg
    defaults = {
        "platform_logo_url": "",
        "company_name": cfg.company_name,
        "company_address": cfg.company_address,
        "company_email": cfg.company_email,
        "company_phone": cfg.company_phone,
        "company_gstin": cfg.company_gstin,
        "company_pan": cfg.company_pan,
        "gst_rate": str(cfg.gst_rate),
    }
    rows = db.query(SiteSetting).all()
    for row in rows:
        defaults[row.key] = row.value or ""
    return defaults


class _SettingsProxy:
    def __init__(self, d: dict):
        self._d = d

    def __getattr__(self, key):
        return self._d.get(key, "")


@router.get("/settings")
def settings_page(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login?next=/admin/settings", status_code=303)
    s = _SettingsProxy(_load_settings(db))
    return templates.TemplateResponse(
        "admin/settings.html",
        _admin_ctx(request, active_page="settings", s=s),
    )


@router.post("/settings")
async def settings_update(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login?next=/admin/settings", status_code=303)
    form = await request.form()
    for key in _SETTING_KEYS:
        val = form.get(key, "").strip()
        row = db.query(SiteSetting).filter(SiteSetting.key == key).first()
        if row:
            row.value = val
        else:
            db.add(SiteSetting(key=key, value=val))
    db.commit()
    flash(request, "Settings saved.", "success")
    return RedirectResponse("/admin/settings", status_code=303)


# ---------------------------------------------------------------------------
# Events
# ---------------------------------------------------------------------------

@router.get("/events")
def events_list(request: Request, db: Session = Depends(get_db), view: str = Query("current")):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    today = date.today()
    stale = (
        db.query(Event)
        .filter(
            Event.status == "published",
            or_(
                Event.end_date < today,
                (Event.end_date.is_(None)) & (Event.start_date < today),
            ),
        )
        .all()
    )
    for ev in stale:
        ev.status = "completed"
    if stale:
        db.commit()

    if view == "completed":
        events = db.query(Event).filter(Event.status == "completed").order_by(Event.created_at.desc()).all()
    else:
        events = db.query(Event).filter(Event.status != "completed").order_by(Event.created_at.desc()).all()

    enriched = []
    for ev in events:
        session_count = len(ev.event_sessions) if ev.event_sessions else 0
        booking_count = db.query(func.count(Booking.id)).filter(
            Booking.event_id == ev.id, Booking.payment_status == "paid",
            Booking.is_shared_ticket == False,
        ).scalar() or 0
        aud = db.query(Auditorium).get(ev.auditorium_id) if ev.auditorium_id else None
        enriched.append({"event": ev, "session_count": session_count, "bookings": booking_count, "auditorium": aud})
    return templates.TemplateResponse(
        "admin/events.html",
        _admin_ctx(request, active_page="events", events=enriched, view=view),
    )


@router.post("/events/{event_id}/status")
async def event_update_status(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)
    ev = db.query(Event).get(event_id)
    if not ev:
        return JSONResponse({"error": "Event not found"}, status_code=404)
    body = await request.json()
    new_status = body.get("status", "").strip().lower()
    if new_status not in ("draft", "published", "completed", "cancelled"):
        return JSONResponse({"error": "Invalid status"}, status_code=400)
    ev.status = new_status
    db.commit()
    return JSONResponse({"ok": True, "status": new_status})


@router.get("/events/new")
def event_new_form(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    colleges = db.query(College).order_by(College.name).all()
    auditoriums = db.query(Auditorium).order_by(Auditorium.name).all()
    all_sessions = db.query(SessionModel).order_by(SessionModel.title).all()
    speakers = db.query(Speaker).order_by(Speaker.name).all()
    aud_seat_types = _auditorium_seat_types(db)
    ct_map = _custom_types_map(db)
    fb_templates = db.query(FeedbackTemplate).order_by(FeedbackTemplate.name).all()
    all_events = db.query(Event).order_by(Event.name).all()
    cert_templates = db.query(CertificateTemplate).order_by(CertificateTemplate.name).all()
    return templates.TemplateResponse(
        "admin/event_form.html",
        _admin_ctx(request, active_page="events", event=None, colleges=colleges,
                   auditoriums=auditoriums, all_sessions=all_sessions, speakers=speakers,
                   aud_seat_types=aud_seat_types, custom_types_map=ct_map,
                   fb_templates=fb_templates, all_events=all_events,
                   cert_templates=cert_templates),
    )


@router.post("/events/new")
async def event_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    aud_raw = form.get("auditorium_id", "")
    start_date_raw = form.get("start_date", "")
    end_date_raw = form.get("end_date", "")
    ev = Event(
        name=form.get("name", "").strip(),
        description=form.get("description", "").strip() or None,
        banner_url=form.get("banner_url", "").strip() or None,
        college_id=int(form.get("college_id")) if form.get("college_id") else None,
        auditorium_id=int(aud_raw) if aud_raw and aud_raw.strip().isdigit() else None,
        start_date=date.fromisoformat(start_date_raw) if start_date_raw else None,
        end_date=date.fromisoformat(end_date_raw) if end_date_raw else None,
        price=float(form.get("price", 0) or 0),
        price_vip=float(form["price_vip"]) if form.get("price_vip", "").strip() else None,
        price_accessible=float(form["price_accessible"]) if form.get("price_accessible", "").strip() else None,
        processing_fee_pct=float(form["processing_fee_pct"]) if form.get("processing_fee_pct", "").strip() else None,
        custom_prices=_extract_custom_prices(form),
        status=form.get("status", "draft"),
        cert_title=normalize_cert_scalar_for_storage(form.get("cert_title")),
        cert_subtitle=normalize_cert_scalar_for_storage(form.get("cert_subtitle")),
        cert_footer=normalize_cert_scalar_for_storage(form.get("cert_footer")),
        cert_signer_name=normalize_cert_scalar_for_storage(form.get("cert_signer_name")),
        cert_signer_designation=normalize_cert_scalar_for_storage(form.get("cert_signer_designation")),
        cert_signature_url=normalize_cert_scalar_for_storage(form.get("cert_signature_url")),
        cert_logo_url=normalize_cert_scalar_for_storage(form.get("cert_logo_url")),
        cert_bg_url=normalize_cert_scalar_for_storage(form.get("cert_bg_url")),
        cert_color_scheme=normalize_cert_scalar_for_storage(form.get("cert_color_scheme")),
        cert_style=normalize_cert_scalar_for_storage(form.get("cert_style")),
        feedback_template_id=int(form["feedback_template_id"]) if form.get("feedback_template_id", "").strip() else None,
    )
    db.add(ev)
    db.flush()

    ctid = (form.get("cert_template_id") or "").strip()
    if ctid.isdigit():
        tpl = db.query(CertificateTemplate).get(int(ctid))
        if tpl:
            _apply_certificate_template_to_event(ev, tpl)

    _save_gallery_images(db, form, "event", ev.id)

    from app.models.event_session import EventSession
    linked = _save_event_sessions(db, form, ev.id)

    _save_event_breaks(db, form, ev.id)
    _save_event_addons(db, form, ev.id)

    log_activity(db, category="admin", action="create", description=f"Created event '{ev.name}'", request=request, user_id=admin.id, target_type="event", target_id=ev.id)
    db.commit()
    msg = f"Event '{ev.name}' created"
    if linked:
        msg += f" with {linked} session(s)"
    flash(request, msg + ".", "success")
    return RedirectResponse("/admin/events", status_code=303)


@router.get("/events/{event_id}/edit")
def event_edit_form(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    if not ev:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)
    from app.models.event_session import EventSession
    colleges = db.query(College).order_by(College.name).all()
    auditoriums = db.query(Auditorium).order_by(Auditorium.name).all()
    ev_sessions = (
        db.query(EventSession).filter(EventSession.event_id == event_id)
        .order_by(EventSession.order, EventSession.start_time).all()
    )
    for es in ev_sessions:
        _ = es.session
    sessions = ev_sessions
    coupons = db.query(Coupon).filter(Coupon.event_id == event_id).order_by(Coupon.created_at.desc()).all()
    linked_ids = {es.session_id for es in ev_sessions}
    all_sessions = [s for s in db.query(SessionModel).order_by(SessionModel.title).all() if s.id not in linked_ids]
    speakers = db.query(Speaker).order_by(Speaker.name).all()
    aud_seat_types = _auditorium_seat_types(db)
    ct_map = _custom_types_map(db)
    gallery = db.query(GalleryImage).filter(
        GalleryImage.owner_type == "event", GalleryImage.owner_id == event_id
    ).order_by(GalleryImage.position).all()
    event_breaks = db.query(EventBreak).filter(EventBreak.event_id == event_id).order_by(EventBreak.order, EventBreak.start_time).all()
    event_addons = db.query(EventAddOn).filter(EventAddOn.event_id == event_id).all()
    fb_templates = db.query(FeedbackTemplate).order_by(FeedbackTemplate.name).all()
    all_events = db.query(Event).filter(Event.id != event_id).order_by(Event.name).all()
    cert_templates = db.query(CertificateTemplate).order_by(CertificateTemplate.name).all()
    return templates.TemplateResponse(
        "admin/event_form.html",
        _admin_ctx(request, active_page="events", event=ev, colleges=colleges, auditoriums=auditoriums,
                   event_sessions=sessions, event_coupons=coupons,
                   all_sessions=all_sessions, speakers=speakers,
                   aud_seat_types=aud_seat_types, custom_types_map=ct_map,
                   gallery_images=gallery, event_breaks=event_breaks,
                   event_addons=event_addons, fb_templates=fb_templates,
                   all_events=all_events, cert_templates=cert_templates),
    )


@router.post("/events/{event_id}/edit")
async def event_update(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    if not ev:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)
    form = await _form(request)
    aud_raw = form.get("auditorium_id", "")
    start_date_raw = form.get("start_date", "")
    end_date_raw = form.get("end_date", "")
    ev.name = form.get("name", "").strip()
    ev.description = form.get("description", "").strip() or None
    ev.banner_url = form.get("banner_url", "").strip() or None
    ev.college_id = int(form.get("college_id")) if form.get("college_id") else None
    ev.auditorium_id = int(aud_raw) if aud_raw and aud_raw.strip().isdigit() else ev.auditorium_id
    if start_date_raw:
        ev.start_date = date.fromisoformat(start_date_raw)
    if end_date_raw:
        ev.end_date = date.fromisoformat(end_date_raw)
    else:
        ev.end_date = None
    ev.price = float(form.get("price", ev.price or 0) or 0)
    ev.price_vip = float(form["price_vip"]) if form.get("price_vip", "").strip() else None
    ev.price_accessible = float(form["price_accessible"]) if form.get("price_accessible", "").strip() else None
    ev.processing_fee_pct = float(form["processing_fee_pct"]) if form.get("processing_fee_pct", "").strip() else None
    ev.custom_prices = _extract_custom_prices(form)
    ev.status = form.get("status", "draft")

    if "cert_title" in form:
        ev.cert_title = normalize_cert_scalar_for_storage(form.get("cert_title"))
    if "cert_subtitle" in form:
        ev.cert_subtitle = normalize_cert_scalar_for_storage(form.get("cert_subtitle"))
    if "cert_footer" in form:
        ev.cert_footer = normalize_cert_scalar_for_storage(form.get("cert_footer"))
    if "cert_signer_name" in form:
        ev.cert_signer_name = normalize_cert_scalar_for_storage(form.get("cert_signer_name"))
    if "cert_signer_designation" in form:
        ev.cert_signer_designation = normalize_cert_scalar_for_storage(form.get("cert_signer_designation"))
    if "cert_signature_url" in form:
        ev.cert_signature_url = normalize_cert_scalar_for_storage(form.get("cert_signature_url"))
    if "cert_logo_url" in form:
        ev.cert_logo_url = normalize_cert_scalar_for_storage(form.get("cert_logo_url"))
    if "cert_bg_url" in form:
        ev.cert_bg_url = normalize_cert_scalar_for_storage(form.get("cert_bg_url"))
    if "cert_color_scheme" in form:
        ev.cert_color_scheme = normalize_cert_scalar_for_storage(form.get("cert_color_scheme"))
    if "cert_style" in form:
        ev.cert_style = normalize_cert_scalar_for_storage(form.get("cert_style"))
    ev.feedback_template_id = int(form["feedback_template_id"]) if form.get("feedback_template_id", "").strip() else None

    from app.models.event_session import EventSession
    linked = _save_event_sessions(db, form, ev.id)

    _save_gallery_images(db, form, "event", ev.id)
    _save_event_breaks(db, form, ev.id)
    _save_event_addons(db, form, ev.id)

    log_activity(db, category="admin", action="update", description=f"Updated event '{ev.name}'", request=request, user_id=admin.id, target_type="event", target_id=ev.id)
    db.commit()
    msg = f"Event '{ev.name}' updated"
    if linked:
        msg += f" — {linked} session(s) linked"
    flash(request, msg + ".", "success")
    return RedirectResponse(f"/admin/events/{ev.id}/edit", status_code=303)


@router.post("/events/{event_id}/delete")
def event_delete(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    if ev:
        paid_count = db.query(func.count(Booking.id)).filter(
            Booking.event_id == ev.id, Booking.payment_status == "paid"
        ).scalar()
        if paid_count > 0:
            flash(request, f"Cannot delete — this event has {paid_count} paid booking(s). Cancel or refund them first.", "danger")
            return RedirectResponse("/admin/events", status_code=303)
        log_activity(db, category="admin", action="delete", description=f"Deleted event '{ev.name}'", request=request, user_id=admin.id, target_type="event", target_id=ev.id)
        db.delete(ev)
        db.commit()
        flash(request, f"Event '{ev.name}' deleted.", "success")
    return RedirectResponse("/admin/events", status_code=303)


# ---------------------------------------------------------------------------
# Agenda template download & import
# ---------------------------------------------------------------------------

_AGENDA_COLUMNS = [
    "Type", "Title", "Abstract", "Key Learning Outcomes", "Agenda",
    "Speaker Name", "Speaker Title", "Speaker Email", "Speaker Bio", "Speaker LinkedIn",
    "Start Time", "End Time",
]
_AGENDA_SAMPLE_ROWS = [
    ["Session", "Opening Keynote", "Welcome address and overview of today's agenda",
     "Understand the conference theme; Learn about keynote speakers",
     "Introduction (10 min); Main Presentation (25 min); Q&A (10 min)",
     "Dr. Sarah", "Professor of CS", "sarah@example.com", "Expert in distributed systems",
     "https://linkedin.com/in/sarah-example",
     "9:00 AM", "9:45 AM"],
    ["Break", "Tea Break", "", "", "", "", "", "", "", "",
     "9:45 AM", "10:00 AM"],
    ["Session", "Workshop: AI Fundamentals", "Hands-on machine learning workshop for beginners",
     "Build a basic ML model; Understand training pipelines; Evaluate model performance",
     "Theory (30 min); Lab Exercise (45 min); Wrap-up (15 min)",
     "James", "ML Engineer", "james@example.com", "",
     "https://linkedin.com/in/james-ml",
     "10:00 AM", "11:30 AM"],
    ["Break", "Lunch", "Networking lunch", "", "", "", "", "", "", "",
     "11:30 AM", "12:30 PM"],
    ["Session", "Panel Discussion", "Industry trends Q&A with leaders",
     "Gain insights into current industry trends; Hear from diverse perspectives",
     "",
     "Maria", "CTO, TechCorp", "", "Industry leader in cloud computing",
     "https://linkedin.com/in/maria-cto",
     "12:30 PM", "1:15 PM"],
]

_AGENDA_FIELDS = [
    "type", "title", "abstract", "key_learning_outcomes", "agenda",
    "speaker", "speaker_title", "speaker_email", "speaker_bio", "speaker_linkedin",
    "start_time", "end_time",
]


def _map_agenda_headers(headers: list[str]) -> dict[str, int]:
    """Map column names to field keys by fuzzy matching."""
    col_map: dict[str, int] = {}
    for i, h in enumerate(headers):
        hl = h.lower().replace("(", "").replace(")", "").replace(" ", "").replace("_", "")
        if hl == "type":
            col_map["type"] = i
        elif hl == "title":
            col_map["title"] = i
        elif hl in ("abstract", "description", "desc"):
            col_map["abstract"] = i
        elif hl in ("keylearningoutcomes", "learningoutcomes", "outcomes", "klo"):
            col_map["key_learning_outcomes"] = i
        elif hl == "agenda":
            col_map["agenda"] = i
        elif hl in ("speakername", "speaker"):
            col_map["speaker"] = i
        elif hl in ("speakertitle", "speakerrole"):
            col_map["speaker_title"] = i
        elif hl in ("speakeremail", "email"):
            col_map["speaker_email"] = i
        elif hl in ("speakerbio", "bio"):
            col_map["speaker_bio"] = i
        elif hl in ("speakerlinkedin", "linkedin", "linkedinurl"):
            col_map["speaker_linkedin"] = i
        elif hl in ("starttime", "start"):
            col_map["start_time"] = i
        elif hl in ("endtime", "end"):
            col_map["end_time"] = i
    return col_map


def _extract_agenda_row(vals: list[str], col_map: dict[str, int]) -> dict[str, str]:
    """Pull values from a row list using the column map, with safe index access."""
    def _get(key: str) -> str:
        idx = col_map.get(key)
        if idx is not None and idx < len(vals):
            return vals[idx].strip()
        return ""
    return {f: _get(f) for f in _AGENDA_FIELDS}


def _parse_agenda_items(raw: str) -> list[dict]:
    """Parse semicolon-separated agenda sub-items.

    Accepted formats per item:
      "Title (X min)"  or  "Title (Xmin)"  or  "Title - X min"  or  "Title"
    Returns list of {"title": str, "duration_minutes": int}.
    """
    import re
    items: list[dict] = []
    if not raw or not raw.strip():
        return items
    for part in raw.split(";"):
        part = part.strip()
        if not part:
            continue
        m = re.match(r'^(.+?)\s*[\(\-]\s*(\d+)\s*(?:min(?:utes?)?)?\s*\)?$', part, re.IGNORECASE)
        if m:
            items.append({"title": m.group(1).strip(), "duration_minutes": int(m.group(2))})
        else:
            items.append({"title": part, "duration_minutes": 0})
    return items


@router.get("/events/agenda-template")
def agenda_template_download(request: Request, db: Session = Depends(get_db), format: str = Query("xlsx")):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    fmt = format.lower().strip()

    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(_AGENDA_COLUMNS)
        writer.writerows(_AGENDA_SAMPLE_ROWS)
        buf.seek(0)
        return StreamingResponse(
            iter([buf.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=agenda_template.csv"},
        )

    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Agenda"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx, col_name in enumerate(_AGENDA_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, row_data in enumerate(_AGENDA_SAMPLE_ROWS, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

    col_widths = [12, 30, 35, 30, 45, 20, 20, 25, 30, 30, 20, 20]
    for i, w in enumerate(col_widths, 1):
        col_letter = chr(64 + i) if i <= 26 else chr(64 + (i - 1) // 26) + chr(65 + (i - 1) % 26)
        ws.column_dimensions[col_letter].width = w

    type_dv = DataValidation(type="list", formula1='"Session,Break"', allow_blank=False)
    type_dv.error = "Please select Session or Break"
    type_dv.errorTitle = "Invalid Type"
    type_dv.prompt = "Choose Session or Break"
    type_dv.promptTitle = "Type"
    ws.add_data_validation(type_dv)
    type_dv.add(f"A2:A200")

    instructions = wb.create_sheet("Instructions")
    notes = [
        ("Column", "Notes"),
        ("Type", "Required. Must be exactly 'Session' or 'Break'."),
        ("Title", "Required. Name of the session or break."),
        ("Abstract", "Optional. A brief overview or summary of the session."),
        ("Key Learning Outcomes", "Optional. Semicolon-separated list of what attendees will learn. "
                                  "Example: 'Build a basic ML model; Understand training pipelines'."),
        ("Agenda", "Optional. For sessions only. Semicolon-separated sub-items with duration. "
                   "Format: 'Item Title (X min); Another Item (Y min)'. "
                   "Example: 'Introduction (10 min); Main Talk (25 min); Q&A (10 min)'. "
                   "Each sub-item's speaker is automatically set to the row's Speaker Name."),
        ("Speaker Email", "The only required speaker field. If the email matches an existing speaker, "
                          "that speaker is reused regardless of name. If the email is new, a new speaker is created. "
                          "If name is left blank, the part before @ is used as the display name."),
        ("Speaker Name", "Optional. For sessions only. Used as the display name when creating a new speaker. "
                         "If email is blank, falls back to matching by name only."),
        ("Speaker Title", "Optional. Role/title for the speaker (e.g. 'Professor of CS', 'CTO'). "
                          "Used when creating a new speaker; ignored if matched by email."),
        ("Speaker Bio", "Optional. Short biography for the speaker. Used when creating a new speaker."),
        ("Speaker LinkedIn", "Optional. LinkedIn profile URL for the speaker (e.g. https://linkedin.com/in/username)."),
        ("Start Time", "Required. 12-hour format with AM/PM (e.g. 9:00 AM, 1:30 PM). The date is taken from the event."),
        ("End Time", "Required. 12-hour format with AM/PM (e.g. 10:30 AM). Duration is calculated automatically."),
    ]
    for r, (a, b) in enumerate(notes, 1):
        instructions.cell(row=r, column=1, value=a).font = Font(bold=(r == 1))
        instructions.cell(row=r, column=2, value=b)
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=agenda_template.xlsx"},
    )


# ---------------------------------------------------------------------------
# Event template download & import
# ---------------------------------------------------------------------------

_EVENT_TEMPLATE_AGENDA_COLS = [
    "Start Time", "End Time", "Type", "Title", "Abstract",
    "Key Learning Outcomes", "Speaker", "Speaker Title",
    "Speaker Email", "Speaker Bio", "Speaker LinkedIn",
]

_EVENT_TEMPLATE_AGENDA_FIELDS = [
    "start_time", "end_time", "type", "title", "abstract",
    "key_learning_outcomes", "speaker", "speaker_title",
    "speaker_email", "speaker_bio", "speaker_linkedin",
]

_EVENT_HEADER_LABELS = [
    "Event Title", "City", "State", "College", "Hall",
    "Start Date", "End Date", "Status",
]


@router.get("/events/event-template")
def event_template_download(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Event"

    label_font = Font(bold=True, size=11)
    label_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for r, label in enumerate(_EVENT_HEADER_LABELS, 1):
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = label_font
        lc.fill = label_fill
        lc.border = thin_border
        ws.cell(row=r, column=2).border = thin_border

    agenda_header_row = len(_EVENT_HEADER_LABELS) + 3
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    for ci, col_name in enumerate(_EVENT_TEMPLATE_AGENDA_COLS, 1):
        cell = ws.cell(row=agenda_header_row, column=ci, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    sample_rows = []
    for orig in _AGENDA_SAMPLE_ROWS:
        mapped = {f: orig[i] if i < len(orig) else "" for i, f in enumerate(_AGENDA_FIELDS)}
        sample_rows.append([mapped.get(f, "") for f in _EVENT_TEMPLATE_AGENDA_FIELDS])

    for ri, row_data in enumerate(sample_rows, agenda_header_row + 1):
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = thin_border

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 30
    agenda_col_widths = [15, 15, 12, 30, 35, 30, 20, 20, 25, 30, 30]
    for i, w in enumerate(agenda_col_widths):
        letter = chr(65 + i)
        if w > (ws.column_dimensions[letter].width or 0):
            ws.column_dimensions[letter].width = w

    instructions = wb.create_sheet("Instructions")
    notes = [
        ("Section", "Notes"),
        ("Event Title", "Required. The name of the event."),
        ("City", "Required. City where the event takes place. Matched by name; created if new."),
        ("State", "Required when creating a new city. State/province for the city."),
        ("College", "Required. Matched by name within the city; created if not found."),
        ("Hall", "Required. Auditorium / hall name. Matched within the college; created if not found."),
        ("Start Date", "Required. Event start date in YYYY-MM-DD format (e.g. 2026-04-07)."),
        ("End Date", "Required. Event end date in YYYY-MM-DD format."),
        ("Status", "Optional. 'Draft', 'Published', etc. Defaults to 'Draft'."),
        ("", ""),
        ("Agenda Section", "Starts below the event header. Same format as the agenda template."),
        ("Start Time", "Required. 12-hour format with AM/PM (e.g. 9:00 AM)."),
        ("End Time", "Required. 12-hour format with AM/PM."),
        ("Type", "Required. 'Session' or 'Break'."),
        ("Title", "Required. Session or break name."),
        ("Abstract", "Optional. Brief session summary."),
        ("Key Learning Outcomes", "Optional. Semicolon-separated outcomes."),
        ("Speaker", "Optional. Speaker display name."),
        ("Speaker Title", "Optional. Speaker role/title."),
        ("Speaker Email", "The primary identifier for speakers. Matched to existing speakers by email."),
        ("Speaker Bio", "Optional. Short biography."),
        ("Speaker LinkedIn", "Optional. LinkedIn profile URL."),
    ]
    for r, (a, b) in enumerate(notes, 1):
        instructions.cell(row=r, column=1, value=a).font = Font(bold=(r == 1))
        instructions.cell(row=r, column=2, value=b)
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=event_template.xlsx"},
    )


@router.get("/events/{event_id}/export-template")
def event_export_template(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    if not ev:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from app.models.event_session import EventSession

    city_name = state_name = college_name = hall_name = ""
    if ev.college_id:
        college_obj = db.query(College).get(ev.college_id)
        if college_obj:
            college_name = college_obj.name
            city_obj = db.query(City).get(college_obj.city_id) if college_obj.city_id else None
            if city_obj:
                city_name = city_obj.name
                state_name = city_obj.state or ""
    if ev.auditorium_id:
        aud_obj = db.query(Auditorium).get(ev.auditorium_id)
        if aud_obj:
            hall_name = aud_obj.name

    header_values = [
        ev.name or "",
        city_name,
        state_name,
        college_name,
        hall_name,
        ev.start_date.isoformat() if ev.start_date else "",
        ev.end_date.isoformat() if ev.end_date else "",
        (ev.status or "draft").capitalize(),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Event"

    label_font = Font(bold=True, size=11)
    label_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for r, (label, value) in enumerate(zip(_EVENT_HEADER_LABELS, header_values), 1):
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = label_font
        lc.fill = label_fill
        lc.border = thin_border
        vc = ws.cell(row=r, column=2, value=value)
        vc.border = thin_border

    agenda_header_row = len(_EVENT_HEADER_LABELS) + 3
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    for ci, col_name in enumerate(_EVENT_TEMPLATE_AGENDA_COLS, 1):
        cell = ws.cell(row=agenda_header_row, column=ci, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    ev_sessions = (
        db.query(EventSession)
        .options(joinedload(EventSession.session))
        .filter(EventSession.event_id == event_id)
        .order_by(EventSession.order, EventSession.start_time).all()
    )
    ev_breaks = (
        db.query(EventBreak).filter(EventBreak.event_id == event_id)
        .order_by(EventBreak.order, EventBreak.start_time).all()
    )

    def _fmt_time(dt):
        if not dt:
            return ""
        h = dt.hour % 12 or 12
        return f"{h}:{dt.minute:02d} {'AM' if dt.hour < 12 else 'PM'}"

    agenda_items = []
    for es in ev_sessions:
        sess = es.session
        if not sess:
            continue
        spk_id = es.speaker_id or sess.speaker_id
        speaker = db.query(Speaker).get(spk_id) if spk_id else None
        start_fmt = _fmt_time(es.start_time)
        end_time = None
        if es.start_time and sess.duration_minutes:
            end_time = es.start_time + timedelta(minutes=sess.duration_minutes)
        end_fmt = _fmt_time(end_time)
        agenda_items.append((es.order, [
            start_fmt, end_fmt, "Session",
            sess.title or "",
            sess.abstract or sess.description or "",
            sess.key_learning_outcomes or "",
            speaker.name if speaker else (es.speaker_name or sess.speaker_name or ""),
            speaker.title if speaker else "",
            speaker.email if speaker else "",
            speaker.bio if speaker else "",
            speaker.linkedin_url if speaker else "",
        ]))
    for eb in ev_breaks:
        start_fmt = _fmt_time(eb.start_time)
        end_time = None
        if eb.start_time and eb.duration_minutes:
            end_time = eb.start_time + timedelta(minutes=eb.duration_minutes)
        end_fmt = _fmt_time(end_time)
        agenda_items.append((eb.order, [
            start_fmt, end_fmt, "Break",
            eb.title or "",
            eb.description or "",
            "", "", "", "", "", "",
        ]))

    agenda_items.sort(key=lambda x: x[0])
    for ri, (_, row_data) in enumerate(agenda_items, agenda_header_row + 1):
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=ci, value=val or "")
            cell.border = thin_border

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 30
    agenda_col_widths = [15, 15, 12, 30, 35, 30, 20, 20, 25, 30, 30]
    for i, w in enumerate(agenda_col_widths):
        letter = chr(65 + i)
        if w > (ws.column_dimensions[letter].width or 0):
            ws.column_dimensions[letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe_name = (ev.name or "event").replace(" ", "_").replace("/", "_")[:40]
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={safe_name}_template.xlsx"},
    )


@router.post("/events/{event_id}/copy-from/{source_id}")
def event_copy_from(request: Request, event_id: int, source_id: int, db: Session = Depends(get_db),
                    _csrf=Depends(csrf_protection)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    src = db.query(Event).get(source_id)
    if not ev or not src:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)
    if ev.id == src.id:
        flash(request, "Cannot copy an event onto itself.", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)

    from app.models.event_session import EventSession

    # Clear existing agenda items on the target
    db.query(EventSession).filter(EventSession.event_id == event_id).delete()
    db.query(EventBreak).filter(EventBreak.event_id == event_id).delete()
    db.query(EventAddOn).filter(EventAddOn.event_id == event_id).delete()
    db.flush()

    # Copy scalar fields (keep the target's name and set status to draft)
    ev.description = src.description
    ev.banner_url = src.banner_url
    ev.college_id = src.college_id
    ev.auditorium_id = src.auditorium_id
    ev.start_date = src.start_date
    ev.end_date = src.end_date
    ev.price = src.price
    ev.price_vip = src.price_vip
    ev.price_accessible = src.price_accessible
    ev.processing_fee_pct = src.processing_fee_pct
    ev.custom_prices = src.custom_prices
    ev.status = "draft"
    ev.cert_title = src.cert_title
    ev.cert_subtitle = src.cert_subtitle
    ev.cert_footer = src.cert_footer
    ev.cert_signer_name = src.cert_signer_name
    ev.cert_signer_designation = src.cert_signer_designation
    ev.cert_logo_url = src.cert_logo_url
    ev.cert_bg_url = src.cert_bg_url
    ev.cert_signature_url = src.cert_signature_url
    ev.cert_color_scheme = src.cert_color_scheme
    ev.cert_style = src.cert_style
    ev.feedback_template_id = src.feedback_template_id

    # Copy sessions
    src_sessions = (
        db.query(EventSession).filter(EventSession.event_id == source_id)
        .order_by(EventSession.order, EventSession.start_time).all()
    )
    for es in src_sessions:
        old_sess = es.session
        if not old_sess:
            continue
        new_sess = SessionModel(
            title=old_sess.title,
            speaker_id=old_sess.speaker_id,
            speaker_name=old_sess.speaker_name,
            description=old_sess.description,
            abstract=old_sess.abstract,
            key_learning_outcomes=old_sess.key_learning_outcomes,
            banner_url=old_sess.banner_url,
            duration_minutes=old_sess.duration_minutes,
        )
        db.add(new_sess)
        db.flush()

        old_agenda = db.query(AgendaItem).filter(AgendaItem.session_id == old_sess.id).order_by(AgendaItem.order).all()
        for ai in old_agenda:
            db.add(AgendaItem(
                session_id=new_sess.id, order=ai.order,
                title=ai.title, duration_minutes=ai.duration_minutes,
                speaker_id=ai.speaker_id, speaker_name=ai.speaker_name,
            ))

        old_speakers = db.query(SessionSpeaker).filter(SessionSpeaker.session_id == old_sess.id).all()
        for ss in old_speakers:
            db.add(SessionSpeaker(session_id=new_sess.id, speaker_id=ss.speaker_id, role=ss.role))

        db.add(EventSession(
            event_id=event_id, session_id=new_sess.id, order=es.order,
            start_time=es.start_time, speaker_id=es.speaker_id, speaker_name=es.speaker_name,
        ))

    # Copy breaks
    src_breaks = db.query(EventBreak).filter(EventBreak.event_id == source_id).order_by(EventBreak.order).all()
    for eb in src_breaks:
        db.add(EventBreak(
            event_id=event_id, title=eb.title, description=eb.description,
            duration_minutes=eb.duration_minutes, start_time=eb.start_time, order=eb.order,
        ))

    # Copy add-ons
    src_addons = db.query(EventAddOn).filter(EventAddOn.event_id == source_id).all()
    for ao in src_addons:
        db.add(EventAddOn(
            event_id=event_id, title=ao.title, description=ao.description,
            price=ao.price, max_quantity=ao.max_quantity, is_active=ao.is_active,
        ))

    log_activity(
        db, category="admin", action="copy_event",
        description=f"Copied details from '{src.name}' to '{ev.name}'",
        request=request, user_id=admin.id, target_type="event", target_id=ev.id,
    )
    db.commit()

    flash(request, f"Copied all details from '{src.name}'. Status set to Draft. Review and save.", "success")
    return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)


@router.post("/events/duplicate/{source_id}")
def event_duplicate(request: Request, source_id: int, db: Session = Depends(get_db),
                    _csrf=Depends(csrf_protection)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    src = db.query(Event).get(source_id)
    if not src:
        flash(request, "Source event not found.", "danger")
        return RedirectResponse("/admin/events/new", status_code=303)

    from app.models.event_session import EventSession

    ev = Event(
        name=f"{src.name} (Copy)",
        description=src.description,
        banner_url=src.banner_url,
        college_id=src.college_id,
        auditorium_id=src.auditorium_id,
        start_date=None,
        end_date=None,
        price=src.price,
        price_vip=src.price_vip,
        price_accessible=src.price_accessible,
        processing_fee_pct=src.processing_fee_pct,
        custom_prices=src.custom_prices,
        status="draft",
        cert_title=src.cert_title,
        cert_subtitle=src.cert_subtitle,
        cert_footer=src.cert_footer,
        cert_signer_name=src.cert_signer_name,
        cert_signer_designation=src.cert_signer_designation,
        cert_logo_url=src.cert_logo_url,
        cert_bg_url=src.cert_bg_url,
        cert_signature_url=src.cert_signature_url,
        cert_color_scheme=src.cert_color_scheme,
        cert_style=src.cert_style,
        feedback_template_id=src.feedback_template_id,
    )
    db.add(ev)
    db.flush()

    src_sessions = (
        db.query(EventSession).filter(EventSession.event_id == source_id)
        .order_by(EventSession.order, EventSession.start_time).all()
    )
    for es in src_sessions:
        old_sess = es.session
        if not old_sess:
            continue
        new_sess = SessionModel(
            title=old_sess.title,
            speaker_id=old_sess.speaker_id,
            speaker_name=old_sess.speaker_name,
            description=old_sess.description,
            abstract=old_sess.abstract,
            key_learning_outcomes=old_sess.key_learning_outcomes,
            banner_url=old_sess.banner_url,
            duration_minutes=old_sess.duration_minutes,
        )
        db.add(new_sess)
        db.flush()

        old_agenda = db.query(AgendaItem).filter(AgendaItem.session_id == old_sess.id).order_by(AgendaItem.order).all()
        for ai in old_agenda:
            db.add(AgendaItem(
                session_id=new_sess.id, order=ai.order,
                title=ai.title, duration_minutes=ai.duration_minutes,
                speaker_id=ai.speaker_id, speaker_name=ai.speaker_name,
            ))

        old_speakers = db.query(SessionSpeaker).filter(SessionSpeaker.session_id == old_sess.id).all()
        for ss in old_speakers:
            db.add(SessionSpeaker(session_id=new_sess.id, speaker_id=ss.speaker_id, role=ss.role))

        db.add(EventSession(
            event_id=ev.id, session_id=new_sess.id, order=es.order,
            start_time=es.start_time, speaker_id=es.speaker_id, speaker_name=es.speaker_name,
        ))

    src_breaks = db.query(EventBreak).filter(EventBreak.event_id == source_id).order_by(EventBreak.order).all()
    for eb in src_breaks:
        db.add(EventBreak(
            event_id=ev.id, title=eb.title, description=eb.description,
            duration_minutes=eb.duration_minutes, start_time=eb.start_time, order=eb.order,
        ))

    src_addons = db.query(EventAddOn).filter(EventAddOn.event_id == source_id).all()
    for ao in src_addons:
        db.add(EventAddOn(
            event_id=ev.id, title=ao.title, description=ao.description,
            price=ao.price, max_quantity=ao.max_quantity, is_active=ao.is_active,
        ))

    log_activity(
        db, category="admin", action="duplicate_event",
        description=f"Duplicated '{src.name}' as '{ev.name}'",
        request=request, user_id=admin.id, target_type="event", target_id=ev.id,
    )
    db.commit()

    flash(request, f"Duplicated '{src.name}'. Dates left blank — please set new dates.", "success")
    return RedirectResponse(f"/admin/events/{ev.id}/edit", status_code=303)


@router.post("/events/import-event")
async def event_import_from_template(
    request: Request,
    event_file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _csrf=Depends(csrf_protection),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    content = await event_file.read()
    fname = (event_file.filename or "").lower()
    if not fname.endswith((".xlsx", ".xls")):
        flash(request, "Event import only supports Excel (.xlsx) files.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        flash(request, f"Could not read the Excel file: {exc}", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # --- Parse event header (key-value pairs) ---
    event_meta: dict[str, str] = {}
    agenda_header_idx = None

    _label_map = {
        "eventtitle": "event_title", "title": "event_title",
        "city": "city", "state": "state",
        "college": "college", "hall": "hall",
        "startdate": "start_date", "enddate": "end_date",
        "status": "status",
    }

    for ri, row in enumerate(all_rows):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if len(cells) >= 2 and cells[0]:
            key = cells[0].lower().replace(" ", "").replace("_", "")
            if key in _label_map:
                event_meta[_label_map[key]] = cells[1]
                continue

        header_lower = [c.lower().replace(" ", "").replace("_", "") if c else "" for c in cells]
        if "type" in header_lower and "title" in header_lower:
            agenda_header_idx = ri
            break

    # --- Validate event header ---
    errors: list[str] = []
    event_title = event_meta.get("event_title", "").strip()
    if not event_title:
        errors.append("Event Title is required.")
    city_name = event_meta.get("city", "").strip()
    state_name = event_meta.get("state", "").strip()
    college_name = event_meta.get("college", "").strip()
    hall_name = event_meta.get("hall", "").strip()
    if not city_name:
        errors.append("City is required.")
    if not college_name:
        errors.append("College is required.")
    if not hall_name:
        errors.append("Hall is required.")

    start_date_str = event_meta.get("start_date", "").strip()
    end_date_str = event_meta.get("end_date", "").strip()
    ev_start_date = ev_end_date = None
    if not start_date_str:
        errors.append("Start Date is required.")
    else:
        try:
            ev_start_date = date.fromisoformat(start_date_str.replace("/", "-"))
        except ValueError:
            errors.append(f"Invalid Start Date format: '{start_date_str}'. Use YYYY-MM-DD.")
    if not end_date_str:
        errors.append("End Date is required.")
    else:
        try:
            ev_end_date = date.fromisoformat(end_date_str.replace("/", "-"))
        except ValueError:
            errors.append(f"Invalid End Date format: '{end_date_str}'. Use YYYY-MM-DD.")

    ev_status = (event_meta.get("status", "") or "draft").strip().lower()

    if errors:
        flash(request, " | ".join(errors), "danger")
        return RedirectResponse("/admin/events", status_code=303)

    # --- Resolve City / College / Auditorium ---
    city_obj = db.query(City).filter(func.lower(City.name) == city_name.lower()).first()
    if not city_obj:
        if not state_name:
            flash(request, f"City '{city_name}' not found and no State provided to create it.", "danger")
            return RedirectResponse("/admin/events", status_code=303)
        city_obj = City(name=city_name, state=state_name)
        db.add(city_obj)
        db.flush()

    college_obj = (
        db.query(College)
        .filter(func.lower(College.name) == college_name.lower(), College.city_id == city_obj.id)
        .first()
    )
    if not college_obj:
        college_obj = College(name=college_name, city_id=city_obj.id)
        db.add(college_obj)
        db.flush()

    hall_obj = (
        db.query(Auditorium)
        .filter(func.lower(Auditorium.name) == hall_name.lower(), Auditorium.college_id == college_obj.id)
        .first()
    )
    if not hall_obj:
        hall_obj = Auditorium(name=hall_name, college_id=college_obj.id, location=college_name)
        db.add(hall_obj)
        db.flush()

    # --- Create the Event ---
    new_event = Event(
        name=event_title,
        college_id=college_obj.id,
        auditorium_id=hall_obj.id,
        start_date=ev_start_date,
        end_date=ev_end_date,
        status=ev_status,
    )
    db.add(new_event)
    db.flush()

    # --- Parse & import agenda rows ---
    sessions_created = 0
    breaks_created = 0
    speakers_created = 0
    agenda_errors: list[str] = []

    if agenda_header_idx is not None and agenda_header_idx + 1 < len(all_rows):
        header_cells = [str(c).strip() if c is not None else "" for c in all_rows[agenda_header_idx]]
        col_map = _map_agenda_headers(header_cells)

        agenda_data_rows: list[dict] = []
        for row in all_rows[agenda_header_idx + 1:]:
            vals = [str(c) if c is not None else "" for c in row]
            if not any(v.strip() for v in vals):
                continue
            agenda_data_rows.append(_extract_agenda_row(vals, col_map))

        for i, row in enumerate(agenda_data_rows, 1):
            errs = _validate_agenda_row(row, i, ev_start_date)
            if errs:
                agenda_errors.extend(errs)

        if not agenda_errors:
            from app.models.event_session import EventSession

            email_cache: dict[str, Speaker | None] = {}
            name_cache: dict[str, Speaker | None] = {}

            def _resolve_speaker(name_raw: str, email_raw: str, row_data: dict) -> tuple[int | None, str | None]:
                nonlocal speakers_created
                email_key = email_raw.lower() if email_raw else ""
                if email_key:
                    if email_key not in email_cache:
                        email_cache[email_key] = (
                            db.query(Speaker).filter(func.lower(Speaker.email) == email_key).first()
                        )
                    matched = email_cache[email_key]
                    if matched:
                        name_cache[matched.name.lower()] = matched
                        return matched.id, matched.name
                    display_name = name_raw or email_raw.split("@")[0]
                    new_sp = Speaker(
                        name=display_name, title=row_data.get("speaker_title") or None,
                        email=email_raw or None, bio=row_data.get("speaker_bio") or None,
                        linkedin_url=row_data.get("speaker_linkedin") or None,
                    )
                    db.add(new_sp); db.flush()
                    email_cache[email_key] = new_sp
                    name_cache[display_name.lower()] = new_sp
                    speakers_created += 1
                    return new_sp.id, new_sp.name
                if not name_raw:
                    return None, None
                name_key = name_raw.lower()
                if name_key not in name_cache:
                    name_cache[name_key] = (
                        db.query(Speaker).filter(func.lower(Speaker.name) == name_key).first()
                    )
                matched = name_cache[name_key]
                if matched:
                    return matched.id, matched.name
                new_sp = Speaker(
                    name=name_raw, title=row_data.get("speaker_title") or None,
                    bio=row_data.get("speaker_bio") or None,
                    linkedin_url=row_data.get("speaker_linkedin") or None,
                )
                db.add(new_sp); db.flush()
                name_cache[name_key] = new_sp
                speakers_created += 1
                return new_sp.id, new_sp.name

            for i, row in enumerate(agenda_data_rows, 1):
                row_type = row["type"].lower()
                title = row["title"]
                start_time = _parse_time_value(row["start_time"], ev_start_date)
                end_time = _parse_time_value(row["end_time"], ev_start_date)
                duration = int((end_time - start_time).total_seconds() // 60)
                order = i - 1

                if row_type == "session":
                    speaker_id, speaker_name = _resolve_speaker(
                        row.get("speaker", ""), row.get("speaker_email", ""), row
                    )
                    sess = SessionModel(
                        title=title, speaker_id=speaker_id,
                        speaker_name=speaker_name or "",
                        abstract=row.get("abstract") or None,
                        key_learning_outcomes=row.get("key_learning_outcomes") or None,
                        description=row.get("abstract") or None,
                        duration_minutes=duration,
                    )
                    db.add(sess); db.flush()
                    if speaker_id:
                        db.add(SessionSpeaker(session_id=sess.id, speaker_id=speaker_id, role="Guest"))
                        db.flush()
                    for ai_idx, ai in enumerate(_parse_agenda_items(row.get("agenda", ""))):
                        db.add(AgendaItem(
                            session_id=sess.id, order=ai_idx,
                            title=ai["title"], duration_minutes=ai["duration_minutes"] or 0,
                            speaker_id=speaker_id, speaker_name=speaker_name,
                        ))
                    db.add(EventSession(
                        event_id=new_event.id, session_id=sess.id, order=order,
                        start_time=start_time, speaker_id=speaker_id, speaker_name=speaker_name,
                    ))
                    sessions_created += 1
                else:
                    db.add(EventBreak(
                        event_id=new_event.id, title=title,
                        description=row.get("abstract") or None,
                        duration_minutes=duration, start_time=start_time, order=order,
                    ))
                    breaks_created += 1

    log_activity(
        db, category="admin", action="import_event",
        description=f"Imported event '{event_title}': {sessions_created} session(s), {breaks_created} break(s)"
                    + (f", {speakers_created} new speaker(s)" if speakers_created else "")
                    + (f" — agenda warnings: {'; '.join(agenda_errors)}" if agenda_errors else ""),
        request=request, user_id=admin.id, target_type="event", target_id=new_event.id,
    )
    db.commit()

    summary = f"Event '{event_title}' created with {sessions_created} session(s) and {breaks_created} break(s)."
    if speakers_created:
        summary += f" {speakers_created} new speaker(s) created."
    if agenda_errors:
        summary += f" Agenda warnings: {'; '.join(agenda_errors)}"
    flash(request, summary, "success")
    return RedirectResponse(f"/admin/events/{new_event.id}/edit?step=4", status_code=303)


@router.post("/events/{event_id}/update-from-template")
async def event_update_from_template(
    request: Request,
    event_id: int,
    event_file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _csrf=Depends(csrf_protection),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    if not ev:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    content = await event_file.read()
    fname = (event_file.filename or "").lower()
    if not fname.endswith((".xlsx", ".xls")):
        flash(request, "Event update only supports Excel (.xlsx) files.", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)

    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        flash(request, f"Could not read the Excel file: {exc}", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)

    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    event_meta: dict[str, str] = {}
    agenda_header_idx = None

    _label_map = {
        "eventtitle": "event_title", "title": "event_title",
        "city": "city", "state": "state",
        "college": "college", "hall": "hall",
        "startdate": "start_date", "enddate": "end_date",
        "status": "status",
    }

    for ri, row in enumerate(all_rows):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if len(cells) >= 2 and cells[0]:
            key = cells[0].lower().replace(" ", "").replace("_", "")
            if key in _label_map:
                event_meta[_label_map[key]] = cells[1]
                continue

        header_lower = [c.lower().replace(" ", "").replace("_", "") if c else "" for c in cells]
        if "type" in header_lower and "title" in header_lower:
            agenda_header_idx = ri
            break

    errors: list[str] = []
    event_title = event_meta.get("event_title", "").strip()
    if not event_title:
        errors.append("Event Title is required.")

    city_name = event_meta.get("city", "").strip()
    state_name = event_meta.get("state", "").strip()
    college_name = event_meta.get("college", "").strip()
    hall_name = event_meta.get("hall", "").strip()

    start_date_str = event_meta.get("start_date", "").strip()
    end_date_str = event_meta.get("end_date", "").strip()
    ev_start_date = ev_end_date = None
    if start_date_str:
        try:
            ev_start_date = date.fromisoformat(start_date_str.replace("/", "-"))
        except ValueError:
            errors.append(f"Invalid Start Date format: '{start_date_str}'. Use YYYY-MM-DD.")
    if end_date_str:
        try:
            ev_end_date = date.fromisoformat(end_date_str.replace("/", "-"))
        except ValueError:
            errors.append(f"Invalid End Date format: '{end_date_str}'. Use YYYY-MM-DD.")

    if errors:
        flash(request, " | ".join(errors), "danger")
        return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)

    ev.name = event_title

    if city_name:
        city_obj = db.query(City).filter(func.lower(City.name) == city_name.lower()).first()
        if not city_obj:
            if not state_name:
                flash(request, f"City '{city_name}' not found and no State provided to create it.", "danger")
                return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)
            city_obj = City(name=city_name, state=state_name)
            db.add(city_obj)
            db.flush()

        if college_name:
            college_obj = (
                db.query(College)
                .filter(func.lower(College.name) == college_name.lower(), College.city_id == city_obj.id)
                .first()
            )
            if not college_obj:
                college_obj = College(name=college_name, city_id=city_obj.id)
                db.add(college_obj)
                db.flush()
            ev.college_id = college_obj.id

            if hall_name:
                hall_obj = (
                    db.query(Auditorium)
                    .filter(func.lower(Auditorium.name) == hall_name.lower(), Auditorium.college_id == college_obj.id)
                    .first()
                )
                if not hall_obj:
                    hall_obj = Auditorium(name=hall_name, college_id=college_obj.id, location=college_name)
                    db.add(hall_obj)
                    db.flush()
                ev.auditorium_id = hall_obj.id

    if ev_start_date:
        ev.start_date = ev_start_date
    if ev_end_date:
        ev.end_date = ev_end_date

    ev_status = (event_meta.get("status", "") or "").strip().lower()
    if ev_status in ("draft", "published", "completed", "cancelled"):
        ev.status = ev_status

    from app.models.event_session import EventSession

    db.query(EventSession).filter(EventSession.event_id == event_id).delete()
    db.query(EventBreak).filter(EventBreak.event_id == event_id).delete()
    db.flush()

    sessions_created = 0
    breaks_created = 0
    speakers_created = 0
    agenda_errors: list[str] = []

    if agenda_header_idx is not None and agenda_header_idx + 1 < len(all_rows):
        header_cells = [str(c).strip() if c is not None else "" for c in all_rows[agenda_header_idx]]
        col_map = _map_agenda_headers(header_cells)

        agenda_data_rows: list[dict] = []
        for row in all_rows[agenda_header_idx + 1:]:
            vals = [str(c) if c is not None else "" for c in row]
            if not any(v.strip() for v in vals):
                continue
            agenda_data_rows.append(_extract_agenda_row(vals, col_map))

        for i, row in enumerate(agenda_data_rows, 1):
            errs = _validate_agenda_row(row, i, ev_start_date or ev.start_date)
            if errs:
                agenda_errors.extend(errs)

        if not agenda_errors:
            email_cache: dict[str, Speaker | None] = {}
            name_cache: dict[str, Speaker | None] = {}

            def _resolve_speaker(name_raw: str, email_raw: str, row_data: dict) -> tuple[int | None, str | None]:
                nonlocal speakers_created
                email_key = email_raw.lower() if email_raw else ""
                if email_key:
                    if email_key not in email_cache:
                        email_cache[email_key] = (
                            db.query(Speaker).filter(func.lower(Speaker.email) == email_key).first()
                        )
                    matched = email_cache[email_key]
                    if matched:
                        name_cache[matched.name.lower()] = matched
                        return matched.id, matched.name
                    display_name = name_raw or email_raw.split("@")[0]
                    new_sp = Speaker(
                        name=display_name, title=row_data.get("speaker_title") or None,
                        email=email_raw or None, bio=row_data.get("speaker_bio") or None,
                        linkedin_url=row_data.get("speaker_linkedin") or None,
                    )
                    db.add(new_sp); db.flush()
                    email_cache[email_key] = new_sp
                    name_cache[display_name.lower()] = new_sp
                    speakers_created += 1
                    return new_sp.id, new_sp.name
                if not name_raw:
                    return None, None
                name_key = name_raw.lower()
                if name_key not in name_cache:
                    name_cache[name_key] = (
                        db.query(Speaker).filter(func.lower(Speaker.name) == name_key).first()
                    )
                matched = name_cache[name_key]
                if matched:
                    return matched.id, matched.name
                new_sp = Speaker(
                    name=name_raw, title=row_data.get("speaker_title") or None,
                    bio=row_data.get("speaker_bio") or None,
                    linkedin_url=row_data.get("speaker_linkedin") or None,
                )
                db.add(new_sp); db.flush()
                name_cache[name_key] = new_sp
                speakers_created += 1
                return new_sp.id, new_sp.name

            for i, row in enumerate(agenda_data_rows, 1):
                row_type = row["type"].lower()
                title = row["title"]
                start_time = _parse_time_value(row["start_time"], ev_start_date or ev.start_date)
                end_time = _parse_time_value(row["end_time"], ev_start_date or ev.start_date)
                duration = int((end_time - start_time).total_seconds() // 60)
                order = i - 1

                if row_type == "session":
                    speaker_id, speaker_name = _resolve_speaker(
                        row.get("speaker", ""), row.get("speaker_email", ""), row
                    )
                    sess = SessionModel(
                        title=title, speaker_id=speaker_id,
                        speaker_name=speaker_name or "",
                        abstract=row.get("abstract") or None,
                        key_learning_outcomes=row.get("key_learning_outcomes") or None,
                        description=row.get("abstract") or None,
                        duration_minutes=duration,
                    )
                    db.add(sess); db.flush()
                    if speaker_id:
                        db.add(SessionSpeaker(session_id=sess.id, speaker_id=speaker_id, role="Guest"))
                        db.flush()
                    for ai_idx, ai in enumerate(_parse_agenda_items(row.get("agenda", ""))):
                        db.add(AgendaItem(
                            session_id=sess.id, order=ai_idx,
                            title=ai["title"], duration_minutes=ai["duration_minutes"] or 0,
                            speaker_id=speaker_id, speaker_name=speaker_name,
                        ))
                    db.add(EventSession(
                        event_id=event_id, session_id=sess.id, order=order,
                        start_time=start_time, speaker_id=speaker_id, speaker_name=speaker_name,
                    ))
                    sessions_created += 1
                else:
                    db.add(EventBreak(
                        event_id=event_id, title=title,
                        description=row.get("abstract") or None,
                        duration_minutes=duration, start_time=start_time, order=order,
                    ))
                    breaks_created += 1

    log_activity(
        db, category="admin", action="update_event_from_template",
        description=f"Updated event '{event_title}' from template: {sessions_created} session(s), {breaks_created} break(s)"
                    + (f", {speakers_created} new speaker(s)" if speakers_created else "")
                    + (f" — agenda warnings: {'; '.join(agenda_errors)}" if agenda_errors else ""),
        request=request, user_id=admin.id, target_type="event", target_id=event_id,
    )
    db.commit()

    summary = f"Event '{event_title}' updated with {sessions_created} session(s) and {breaks_created} break(s)."
    if speakers_created:
        summary += f" {speakers_created} new speaker(s) created."
    if agenda_errors:
        summary += f" Agenda warnings: {'; '.join(agenda_errors)}"
    flash(request, summary, "success")
    return RedirectResponse(f"/admin/events/{event_id}/edit?step=4", status_code=303)


def _parse_agenda_file(content: bytes, filename: str) -> list[dict]:
    """Parse an uploaded XLSX or CSV file into a list of row dicts."""
    rows: list[dict] = []
    fname = (filename or "").lower()
    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = _map_agenda_headers(header)
        for row in ws.iter_rows(min_row=2, values_only=True):
            vals = [str(c) if c is not None else "" for c in row]
            if not any(v.strip() for v in vals):
                continue
            rows.append(_extract_agenda_row(vals, col_map))
        wb.close()
    else:
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        csv_headers = [str(h or "") for h in (reader.fieldnames or [])]
        csv_col_map = _map_agenda_headers(csv_headers)
        for raw_row in reader:
            vals = [raw_row.get(h, "") or "" for h in csv_headers]
            norm = _extract_agenda_row(vals, csv_col_map)
            if norm["type"] or norm["title"]:
                rows.append(norm)
    return rows


import re as _re

_TIME_12H_RE = _re.compile(
    r'^(\d{1,2}):(\d{2})\s*(AM|PM|am|pm|a\.m\.|p\.m\.)$', _re.IGNORECASE
)
_TIME_24H_RE = _re.compile(r'^\d{1,2}:\d{2}$')


def _parse_time_value(raw: str, event_date: date | None = None) -> datetime | None:
    """Parse a time string in 12-hour (e.g. '9:00 AM') or 24-hour (e.g. '09:00') format,
    combined with event_date. Full datetime strings are also accepted for backward compat."""
    raw = raw.strip()
    if not raw:
        return None
    base = event_date or date.today()

    m12 = _TIME_12H_RE.match(raw)
    if m12:
        h, mins, period = int(m12.group(1)), int(m12.group(2)), m12.group(3).upper().replace(".", "")
        if period == "PM" and h != 12:
            h += 12
        elif period == "AM" and h == 12:
            h = 0
        return datetime(base.year, base.month, base.day, h, mins)

    if _TIME_24H_RE.match(raw):
        h, mins = raw.split(":")
        return datetime(base.year, base.month, base.day, int(h), int(mins))

    return datetime.fromisoformat(raw.replace(" ", "T"))


def _validate_agenda_row(row: dict, index: int, event_date: date | None = None) -> list[str]:
    """Validate a single parsed agenda row. Returns a list of error strings (empty = valid)."""
    errors: list[str] = []
    row_type = (row.get("type") or "").strip().lower()
    title = (row.get("title") or "").strip()

    if row_type not in ("session", "break"):
        errors.append(f"Invalid type '{row.get('type', '')}' (must be Session or Break)")
        return errors
    if not title:
        errors.append("Title is required")
        return errors

    start_time = end_time = None
    if row.get("start_time"):
        try:
            start_time = _parse_time_value(row["start_time"], event_date)
        except (ValueError, TypeError):
            errors.append(f"Invalid Start Time '{row['start_time']}' (use e.g. 9:00 AM)")
            return errors
    if row.get("end_time"):
        try:
            end_time = _parse_time_value(row["end_time"], event_date)
        except (ValueError, TypeError):
            errors.append(f"Invalid End Time '{row['end_time']}' (use e.g. 9:00 AM)")
            return errors

    if not start_time or not end_time:
        errors.append("Both Start Time and End Time are required")
        return errors
    if end_time <= start_time:
        errors.append("End Time must be after Start Time")

    return errors


def _row_to_values(row: dict) -> list[str]:
    """Convert a parsed row dict back to a list of cell values in _AGENDA_COLUMNS order."""
    return [row.get(f, "") for f in _AGENDA_FIELDS]


@router.post("/events/{event_id}/preview-agenda")
async def event_preview_agenda(
    request: Request,
    event_id: int,
    db: Session = Depends(get_db),
    agenda_file: UploadFile = File(...),
):
    """Parse and validate an uploaded agenda file, returning JSON for the preview modal."""
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Not authenticated"}, status_code=401)

    ev = db.query(Event).get(event_id)
    if not ev:
        return JSONResponse({"error": "Event not found"}, status_code=404)

    content = await agenda_file.read()
    try:
        rows = _parse_agenda_file(content, agenda_file.filename)
    except Exception as exc:
        return JSONResponse({"error": f"Could not read the file: {exc}"}, status_code=400)

    if not rows:
        return JSONResponse({"error": "The uploaded file has no data rows."}, status_code=400)

    result_rows = []
    for i, row in enumerate(rows, start=1):
        errs = _validate_agenda_row(row, i, event_date=ev.start_date)
        result_rows.append({
            "values": _row_to_values(row),
            "errors": errs,
            "valid": len(errs) == 0,
        })

    return JSONResponse({
        "columns": list(_AGENDA_COLUMNS),
        "rows": result_rows,
    })


@router.post("/events/{event_id}/confirm-agenda")
async def event_confirm_agenda(
    request: Request,
    event_id: int,
    db: Session = Depends(get_db),
):
    """Accept edited rows as JSON, validate, create DB records, return JSON result."""
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Not authenticated"}, status_code=401)

    ev = db.query(Event).get(event_id)
    if not ev:
        return JSONResponse({"error": "Event not found"}, status_code=404)

    body = await request.json()
    raw_rows = body.get("rows", [])
    if not raw_rows:
        return JSONResponse({"error": "No rows provided."}, status_code=400)

    rows: list[dict] = []
    for vals in raw_rows:
        row = {}
        for idx, field in enumerate(_AGENDA_FIELDS):
            row[field] = (vals[idx] if idx < len(vals) else "").strip()
        rows.append(row)

    validation_results = []
    has_errors = False
    for i, row in enumerate(rows, start=1):
        errs = _validate_agenda_row(row, i, event_date=ev.start_date)
        validation_results.append({"index": i - 1, "errors": errs, "valid": len(errs) == 0})
        if errs:
            has_errors = True

    if has_errors:
        return JSONResponse({"success": False, "rows": validation_results})

    from app.models.event_session import EventSession

    existing_order = max(
        [es.order for es in db.query(EventSession).filter(EventSession.event_id == event_id).all()]
        + [eb.order for eb in db.query(EventBreak).filter(EventBreak.event_id == event_id).all()]
        + [-1]
    ) + 1

    email_cache: dict[str, Speaker | None] = {}
    name_cache: dict[str, Speaker | None] = {}
    sessions_created = 0
    breaks_created = 0
    speakers_created = 0

    def _resolve_speaker(name_raw: str, email_raw: str, row_data: dict) -> tuple[int | None, str | None]:
        nonlocal speakers_created
        email_key = email_raw.lower() if email_raw else ""
        if email_key:
            if email_key not in email_cache:
                email_cache[email_key] = (
                    db.query(Speaker).filter(func.lower(Speaker.email) == email_key).first()
                )
            matched = email_cache[email_key]
            if matched:
                name_cache[matched.name.lower()] = matched
                return matched.id, matched.name
            display_name = name_raw or email_raw.split("@")[0]
            new_sp = Speaker(
                name=display_name,
                title=row_data.get("speaker_title") or None,
                email=email_raw or None,
                bio=row_data.get("speaker_bio") or None,
                linkedin_url=row_data.get("speaker_linkedin") or None,
            )
            db.add(new_sp)
            db.flush()
            email_cache[email_key] = new_sp
            name_cache[display_name.lower()] = new_sp
            speakers_created += 1
            return new_sp.id, new_sp.name
        if not name_raw:
            return None, None
        name_key = name_raw.lower()
        if name_key not in name_cache:
            name_cache[name_key] = (
                db.query(Speaker).filter(func.lower(Speaker.name) == name_key).first()
            )
        matched = name_cache[name_key]
        if matched:
            return matched.id, matched.name
        new_sp = Speaker(
            name=name_raw,
            title=row_data.get("speaker_title") or None,
            bio=row_data.get("speaker_bio") or None,
            linkedin_url=row_data.get("speaker_linkedin") or None,
        )
        db.add(new_sp)
        db.flush()
        name_cache[name_key] = new_sp
        speakers_created += 1
        return new_sp.id, new_sp.name

    for i, row in enumerate(rows, start=1):
        row_type = row["type"].lower()
        title = row["title"]
        start_time = _parse_time_value(row["start_time"], ev.start_date)
        end_time = _parse_time_value(row["end_time"], ev.start_date)
        duration = int((end_time - start_time).total_seconds() // 60)
        order = existing_order + i - 1

        if row_type == "session":
            speaker_id, speaker_name = _resolve_speaker(
                row.get("speaker", ""), row.get("speaker_email", ""), row
            )
            sess = SessionModel(
                title=title,
                speaker_id=speaker_id,
                speaker_name=speaker_name or "",
                abstract=row.get("abstract") or None,
                key_learning_outcomes=row.get("key_learning_outcomes") or None,
                description=row.get("abstract") or None,
                duration_minutes=duration,
            )
            db.add(sess)
            db.flush()
            if speaker_id:
                db.add(SessionSpeaker(session_id=sess.id, speaker_id=speaker_id, role="Guest"))
                db.flush()
            for ai_idx, ai in enumerate(_parse_agenda_items(row.get("agenda", ""))):
                db.add(AgendaItem(
                    session_id=sess.id, order=ai_idx,
                    title=ai["title"], duration_minutes=ai["duration_minutes"] or 0,
                    speaker_id=speaker_id, speaker_name=speaker_name,
                ))
            db.add(EventSession(
                event_id=event_id, session_id=sess.id, order=order,
                start_time=start_time, speaker_id=speaker_id, speaker_name=speaker_name,
            ))
            sessions_created += 1
        else:
            db.add(EventBreak(
                event_id=event_id, title=title,
                description=row.get("abstract") or None,
                duration_minutes=duration, start_time=start_time, order=order,
            ))
            breaks_created += 1

    log_activity(
        db, category="admin", action="import_agenda",
        description=f"Imported agenda for '{ev.name}': {sessions_created} session(s), {breaks_created} break(s)"
                    + (f", {speakers_created} new speaker(s)" if speakers_created else ""),
        request=request, user_id=admin.id, target_type="event", target_id=ev.id,
    )
    db.commit()

    parts = [f"{sessions_created} session(s)", f"{breaks_created} break(s)"]
    if speakers_created:
        parts.append(f"{speakers_created} new speaker(s)")
    return JSONResponse({
        "success": True,
        "message": "Imported " + " and ".join(parts) + ".",
        "sessions_created": sessions_created,
        "breaks_created": breaks_created,
        "speakers_created": speakers_created,
    })


@router.post("/events/{event_id}/import-agenda")
async def event_import_agenda(
    request: Request,
    event_id: int,
    db: Session = Depends(get_db),
    agenda_file: UploadFile = File(...),
):
    """Legacy form-POST fallback for non-JS browsers."""
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    if not ev:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)

    content = await agenda_file.read()
    try:
        rows = _parse_agenda_file(content, agenda_file.filename)
    except Exception as exc:
        flash(request, f"Could not read the file: {exc}", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/edit?step=4", status_code=303)

    if not rows:
        flash(request, "The uploaded file has no data rows.", "warning")
        return RedirectResponse(f"/admin/events/{event_id}/edit?step=4", status_code=303)

    from app.models.event_session import EventSession

    existing_order = max(
        [es.order for es in db.query(EventSession).filter(EventSession.event_id == event_id).all()]
        + [eb.order for eb in db.query(EventBreak).filter(EventBreak.event_id == event_id).all()]
        + [-1]
    ) + 1

    email_cache: dict[str, Speaker | None] = {}
    name_cache: dict[str, Speaker | None] = {}
    errors: list[str] = []
    sessions_created = 0
    breaks_created = 0
    speakers_created = 0

    def _resolve_speaker(name_raw: str, email_raw: str, row_data: dict) -> tuple[int | None, str | None]:
        nonlocal speakers_created
        email_key = email_raw.lower() if email_raw else ""
        if email_key:
            if email_key not in email_cache:
                email_cache[email_key] = (
                    db.query(Speaker).filter(func.lower(Speaker.email) == email_key).first()
                )
            matched = email_cache[email_key]
            if matched:
                name_cache[matched.name.lower()] = matched
                return matched.id, matched.name
            display_name = name_raw or email_raw.split("@")[0]
            new_sp = Speaker(
                name=display_name, title=row_data.get("speaker_title") or None,
                email=email_raw or None, bio=row_data.get("speaker_bio") or None,
                linkedin_url=row_data.get("speaker_linkedin") or None,
            )
            db.add(new_sp); db.flush()
            email_cache[email_key] = new_sp
            name_cache[display_name.lower()] = new_sp
            speakers_created += 1
            return new_sp.id, new_sp.name
        if not name_raw:
            return None, None
        name_key = name_raw.lower()
        if name_key not in name_cache:
            name_cache[name_key] = (
                db.query(Speaker).filter(func.lower(Speaker.name) == name_key).first()
            )
        matched = name_cache[name_key]
        if matched:
            return matched.id, matched.name
        new_sp = Speaker(name=name_raw, title=row_data.get("speaker_title") or None, bio=row_data.get("speaker_bio") or None,
                         linkedin_url=row_data.get("speaker_linkedin") or None)
        db.add(new_sp); db.flush()
        name_cache[name_key] = new_sp
        speakers_created += 1
        return new_sp.id, new_sp.name

    for i, row in enumerate(rows, start=1):
        row_type = row["type"].lower()
        title = row["title"]
        if row_type not in ("session", "break"):
            errors.append(f"Row {i}: Invalid type '{row['type']}' (must be Session or Break)")
            continue
        if not title:
            errors.append(f"Row {i}: Title is required")
            continue
        start_time = end_time = None
        if row.get("start_time"):
            try:
                start_time = _parse_time_value(row["start_time"], ev.start_date)
            except (ValueError, TypeError):
                errors.append(f"Row {i}: Invalid Start Time '{row['start_time']}'"); continue
        if row.get("end_time"):
            try:
                end_time = _parse_time_value(row["end_time"], ev.start_date)
            except (ValueError, TypeError):
                errors.append(f"Row {i}: Invalid End Time '{row['end_time']}'"); continue
        if not start_time or not end_time:
            errors.append(f"Row {i}: Both Start Time and End Time are required"); continue
        if end_time <= start_time:
            errors.append(f"Row {i}: End Time must be after Start Time"); continue
        duration = int((end_time - start_time).total_seconds() // 60)
        order = existing_order + i - 1
        if row_type == "session":
            speaker_id, speaker_name = _resolve_speaker(row.get("speaker", ""), row.get("speaker_email", ""), row)
            sess = SessionModel(title=title, speaker_id=speaker_id, speaker_name=speaker_name or "",
                                abstract=row.get("abstract") or None,
                                key_learning_outcomes=row.get("key_learning_outcomes") or None,
                                description=row.get("abstract") or None, duration_minutes=duration)
            db.add(sess); db.flush()
            if speaker_id:
                db.add(SessionSpeaker(session_id=sess.id, speaker_id=speaker_id, role="Guest")); db.flush()
            for ai_idx, ai in enumerate(_parse_agenda_items(row.get("agenda", ""))):
                db.add(AgendaItem(session_id=sess.id, order=ai_idx, title=ai["title"],
                                  duration_minutes=ai["duration_minutes"] or 0,
                                  speaker_id=speaker_id, speaker_name=speaker_name))
            db.add(EventSession(event_id=event_id, session_id=sess.id, order=order,
                                start_time=start_time, speaker_id=speaker_id, speaker_name=speaker_name))
            sessions_created += 1
        else:
            db.add(EventBreak(event_id=event_id, title=title, description=row.get("abstract") or None,
                              duration_minutes=duration, start_time=start_time, order=order))
            breaks_created += 1

    if errors:
        flash(request, f"Skipped {len(errors)} row(s): " + "; ".join(errors[:5]), "warning")

    if sessions_created or breaks_created:
        log_activity(
            db, category="admin", action="import_agenda",
            description=f"Imported agenda for '{ev.name}': {sessions_created} session(s), {breaks_created} break(s)"
                        + (f", {speakers_created} new speaker(s)" if speakers_created else ""),
            request=request, user_id=admin.id, target_type="event", target_id=ev.id,
        )
        db.commit()
        parts = [f"{sessions_created} session(s)", f"{breaks_created} break(s)"]
        if speakers_created:
            parts.append(f"{speakers_created} new speaker(s)")
        flash(request, "Imported " + " and ".join(parts) + ".", "success")
    else:
        flash(request, "No valid rows found to import.", "warning")

    return RedirectResponse(f"/admin/events/{event_id}/edit?step=4", status_code=303)


# ---------------------------------------------------------------------------
# Coupon CRUD (event-level)
# ---------------------------------------------------------------------------

@router.get("/events/{event_id}/coupons/new")
def coupon_new_form(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    if not ev:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)
    return templates.TemplateResponse(
        "admin/coupon_form.html",
        _admin_ctx(request, active_page="events", event=ev, coupon=None),
    )


@router.post("/events/{event_id}/coupons/new")
async def coupon_create(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    code = form.get("code", "").strip().upper()
    if not code:
        flash(request, "Coupon code is required.", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/coupons/new", status_code=303)
    existing = db.query(Coupon).filter(Coupon.code == code).first()
    if existing:
        flash(request, f"Coupon code '{code}' already exists.", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/coupons/new", status_code=303)
    valid_from_raw = form.get("valid_from", "")
    valid_until_raw = form.get("valid_until", "")
    coupon = Coupon(
        code=code,
        event_id=event_id,
        discount_pct=float(form["discount_pct"]) if form.get("discount_pct", "").strip() else None,
        discount_amount=float(form["discount_amount"]) if form.get("discount_amount", "").strip() else None,
        max_uses=int(form["max_uses"]) if form.get("max_uses", "").strip() else None,
        valid_from=datetime.fromisoformat(valid_from_raw) if valid_from_raw else None,
        valid_until=datetime.fromisoformat(valid_until_raw) if valid_until_raw else None,
        is_active=form.get("is_active") == "on",
    )
    db.add(coupon)
    log_activity(db, category="admin", action="create", description=f"Created coupon '{code}' for event #{event_id}", request=request, user_id=admin.id, target_type="coupon", target_id=coupon.id)
    db.commit()
    flash(request, f"Coupon '{code}' created.", "success")
    return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)


@router.get("/events/{event_id}/coupons/{coupon_id}/edit")
def coupon_edit_form(request: Request, event_id: int, coupon_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    ev = db.query(Event).get(event_id)
    coupon = db.query(Coupon).get(coupon_id)
    if not ev or not coupon:
        flash(request, "Not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)
    return templates.TemplateResponse(
        "admin/coupon_form.html",
        _admin_ctx(request, active_page="events", event=ev, coupon=coupon),
    )


@router.post("/events/{event_id}/coupons/{coupon_id}/edit")
async def coupon_update(request: Request, event_id: int, coupon_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    coupon = db.query(Coupon).get(coupon_id)
    if not coupon:
        flash(request, "Coupon not found.", "danger")
        return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)
    form = await _form(request)
    code = form.get("code", "").strip().upper()
    if code and code != coupon.code:
        dup = db.query(Coupon).filter(Coupon.code == code, Coupon.id != coupon_id).first()
        if dup:
            flash(request, f"Coupon code '{code}' already exists.", "danger")
            return RedirectResponse(f"/admin/events/{event_id}/coupons/{coupon_id}/edit", status_code=303)
        coupon.code = code
    coupon.discount_pct = float(form["discount_pct"]) if form.get("discount_pct", "").strip() else None
    coupon.discount_amount = float(form["discount_amount"]) if form.get("discount_amount", "").strip() else None
    coupon.max_uses = int(form["max_uses"]) if form.get("max_uses", "").strip() else None
    valid_from_raw = form.get("valid_from", "")
    valid_until_raw = form.get("valid_until", "")
    coupon.valid_from = datetime.fromisoformat(valid_from_raw) if valid_from_raw else None
    coupon.valid_until = datetime.fromisoformat(valid_until_raw) if valid_until_raw else None
    coupon.is_active = form.get("is_active") == "on"
    log_activity(db, category="admin", action="update", description=f"Updated coupon '{coupon.code}'", request=request, user_id=admin.id, target_type="coupon", target_id=coupon.id)
    db.commit()
    flash(request, f"Coupon '{coupon.code}' updated.", "success")
    return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)


@router.post("/events/{event_id}/coupons/{coupon_id}/delete")
def coupon_delete(request: Request, event_id: int, coupon_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    coupon = db.query(Coupon).get(coupon_id)
    if coupon:
        log_activity(db, category="admin", action="delete", description=f"Deleted coupon '{coupon.code}'", request=request, user_id=admin.id, target_type="coupon", target_id=coupon.id)
        db.delete(coupon)
        db.commit()
        flash(request, "Coupon deleted.", "success")
    return RedirectResponse(f"/admin/events/{event_id}/edit", status_code=303)


# ---------------------------------------------------------------------------
# Feedback Management
# ---------------------------------------------------------------------------

@router.get("/feedback")
def feedback_list(
    request: Request,
    db: Session = Depends(get_db),
    rating_filter: str = Query("", alias="rating"),
    featured_filter: str = Query("", alias="featured"),
    page: int = Query(1, ge=1),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    query = db.query(Feedback)
    if rating_filter:
        try:
            query = query.filter(Feedback.rating == int(rating_filter))
        except ValueError:
            pass
    if featured_filter == "yes":
        query = query.filter(Feedback.is_featured == True)
    elif featured_filter == "no":
        query = query.filter(Feedback.is_featured == False)

    total_count = query.count()
    total_pages = max(1, (total_count + ADMIN_PAGE_SIZE - 1) // ADMIN_PAGE_SIZE)
    page = min(page, total_pages)

    feedback_rows = (
        query.order_by(Feedback.created_at.desc())
        .offset((page - 1) * ADMIN_PAGE_SIZE)
        .limit(ADMIN_PAGE_SIZE)
        .all()
    )
    enriched = []
    for fb in feedback_rows:
        user = db.query(User).get(fb.user_id)
        event = db.query(Event).get(fb.event_id) if fb.event_id else None
        enriched.append({"feedback": fb, "user": user, "event": event})

    return templates.TemplateResponse(
        "admin/feedback.html",
        _admin_ctx(
            request,
            active_page="feedback",
            feedback_items=enriched,
            rating_filter=rating_filter,
            featured_filter=featured_filter,
            page=page,
            total_pages=total_pages,
            total_count=total_count,
        ),
    )


@router.post("/feedback/{feedback_id}/toggle-featured")
def feedback_toggle_featured(request: Request, feedback_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    fb = db.query(Feedback).get(feedback_id)
    if fb:
        fb.is_featured = not fb.is_featured
        status = "featured" if fb.is_featured else "unfeatured"
        log_activity(db, category="admin", action="update", description=f"Toggled feedback #{fb.id} to {status}", request=request, user_id=admin.id, target_type="feedback", target_id=fb.id)
        db.commit()
        flash(request, f"Feedback #{fb.id} is now {status}.", "success")
    return RedirectResponse("/admin/feedback", status_code=303)


# ─── Admin Poll Management ───
# Session-level poll list/create routes removed — managed via Event Management Hub


@router.post("/polls/{poll_id}/toggle")
async def admin_toggle_poll(request: Request, poll_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"ok": False}, status_code=403)
    poll = db.query(Poll).get(poll_id)
    if not poll:
        return JSONResponse({"ok": False, "error": "Poll not found."}, status_code=404)
    if not poll.is_active:
        db.query(Poll).filter(
            Poll.session_id == poll.session_id, Poll.event_id == poll.event_id, Poll.is_active == True
        ).update({Poll.is_active: False})
        poll.is_active = True
    else:
        poll.is_active = False
    db.commit()
    from app.services.poll_events import publish
    from app.services.polls import poll_results as _poll_results, notify_event_attendees_of_poll as _notify_event_attendees_of_poll, notify_event_attendees_poll_closed as _notify_event_attendees_poll_closed
    if poll.is_active:
        results = _poll_results(db, poll)
        await publish(poll.session_id, poll.event_id, results)
        _notify_event_attendees_of_poll(db, poll, results)
    else:
        await publish(poll.session_id, poll.event_id, {"poll_id": poll.id, "is_active": False, "closed": True})
        _notify_event_attendees_poll_closed(db, poll)
    return JSONResponse({"ok": True, "is_active": poll.is_active})


@router.post("/polls/{poll_id}/close")
async def admin_close_poll(request: Request, poll_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"ok": False}, status_code=403)
    poll = db.query(Poll).get(poll_id)
    if not poll:
        return JSONResponse({"ok": False, "error": "Poll not found."}, status_code=404)
    poll.is_active = False
    poll.closed_at = now_ist()
    db.commit()
    from app.services.poll_events import publish
    from app.services.polls import notify_event_attendees_poll_closed as _notify_event_attendees_poll_closed
    await publish(poll.session_id, poll.event_id, {"poll_id": poll.id, "is_active": False, "closed": True})
    _notify_event_attendees_poll_closed(db, poll)
    return JSONResponse({"ok": True})


@router.post("/polls/{poll_id}/delete")
def admin_delete_poll(request: Request, poll_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    poll = db.query(Poll).get(poll_id)
    if not poll:
        flash(request, "Poll not found.", "danger")
        return RedirectResponse("/admin/events", status_code=303)
    sid = poll.session_id
    eid = poll.event_id
    db.delete(poll)
    db.commit()
    flash(request, "Poll deleted.", "success")
    return RedirectResponse(f"/admin/sessions/{sid}/polls?event_id={eid}", status_code=303)


# ─── Certificate Templates (library) ───


def _certificate_template_designer_bootstrap(tpl: CertificateTemplate) -> dict:
    from app.services.certificate import (
        _raw_cert_style_dict,
        default_freeform_cert_style_dict,
        should_render_certificate_as_freeform,
        _parse_cert_style,
    )

    raw = _raw_cert_style_dict(tpl)
    if should_render_certificate_as_freeform(raw):
        initial_style = raw
    else:
        initial_style = default_freeform_cert_style_dict()
        sty = _parse_cert_style(tpl)
        initial_style["border_style"] = sty.get("border_style", initial_style["border_style"])
        initial_style["border_width"] = float(sty.get("border_width", initial_style["border_width"]))
        initial_style["bg_size"] = sty.get("bg_size", initial_style["bg_size"])
        initial_style["bg_offset_x"] = float(sty.get("bg_offset_x", 0))
        initial_style["bg_offset_y"] = float(sty.get("bg_offset_y", 0))
        for k, sk in (
            ("border_color_primary", "border_color_primary"),
            ("border_color_secondary", "border_color_secondary"),
            ("border_color_tertiary", "border_color_tertiary"),
        ):
            if sty.get(sk):
                initial_style[k] = sty[sk]

    if not str(initial_style.get("background_image_url") or "").strip():
        bg = getattr(tpl, "cert_bg_url", None) or None
        if bg:
            initial_style = dict(initial_style)
            initial_style["background_image_url"] = bg

    tid = tpl.id
    _save = f"/admin/certificate-templates/{tid}/certificate/designer"
    _leg = f"/admin/certificate-templates/{tid}/certificate/legacy-to-freeform"
    return {
        "eventId": None,
        "templateId": tid,
        "isPersistedFreeform": should_render_certificate_as_freeform(raw),
        "initialStyle": initial_style,
        "saveUrl": _save,
        "legacyConvertUrl": _leg,
        "designerSaveUrl": _save,
        "designerLegacyUrl": _leg,
        "aiGenerateUrl": "/admin/certificate-ai/generate-template",
    }


@router.get("/certificate-templates/{template_id}/apply-payload")
def certificate_template_apply_payload(
    request: Request, template_id: int, db: Session = Depends(get_db)
):
    """JSON for admin event form (new event): merge template into certificate fields client-side."""
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)
    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        return JSONResponse({"error": "Not found"}, status_code=404)
    return JSONResponse(
        {
            "cert_style": tpl.cert_style or "",
            "cert_title": tpl.cert_title,
            "cert_subtitle": tpl.cert_subtitle,
            "cert_footer": tpl.cert_footer,
            "cert_signer_name": tpl.cert_signer_name,
            "cert_signer_designation": tpl.cert_signer_designation,
            "cert_signature_url": tpl.cert_signature_url,
            "cert_logo_url": tpl.cert_logo_url,
            "cert_bg_url": tpl.cert_bg_url,
            "cert_color_scheme": tpl.cert_color_scheme,
        }
    )


@router.get("/certificate-templates")
def certificate_template_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    cert_templates = (
        db.query(CertificateTemplate)
        .options(joinedload(CertificateTemplate.events))
        .order_by(CertificateTemplate.name)
        .all()
    )
    return templates.TemplateResponse(
        "admin/certificate_template_list.html",
        _admin_ctx(request, active_page="cert-templates", cert_templates=cert_templates),
    )


@router.get("/certificate-templates/new")
def certificate_template_new(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    return templates.TemplateResponse(
        "admin/certificate_template_form.html",
        _admin_ctx(request, active_page="cert-templates", tpl=None),
    )


@router.post("/certificate-templates/new")
async def certificate_template_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    name = (form.get("name") or "").strip()
    if not name:
        flash(request, "Template name is required.", "danger")
        return RedirectResponse("/admin/certificate-templates/new", status_code=303)
    tpl = CertificateTemplate(
        name=name,
        description=(form.get("description") or "").strip() or None,
        created_by=admin.id,
    )
    db.add(tpl)
    db.flush()
    log_activity(
        db,
        category="admin",
        action="create",
        description=f"Created certificate template '{tpl.name}'",
        request=request,
        user_id=admin.id,
        target_type="certificate_template",
        target_id=tpl.id,
    )
    db.commit()
    flash(request, f"Template “{tpl.name}” created — open the designer to build the layout.", "success")
    return RedirectResponse(f"/admin/certificate-templates/{tpl.id}/designer", status_code=303)


@router.get("/certificate-templates/{template_id}/edit")
def certificate_template_edit(request: Request, template_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/certificate-templates", status_code=303)
    return templates.TemplateResponse(
        "admin/certificate_template_form.html",
        _admin_ctx(request, active_page="cert-templates", tpl=tpl),
    )


@router.post("/certificate-templates/{template_id}/edit")
async def certificate_template_update(request: Request, template_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/certificate-templates", status_code=303)
    form = await _form(request)
    tpl.name = (form.get("name") or "").strip() or tpl.name
    tpl.description = (form.get("description") or "").strip() or None
    tpl.cert_title = normalize_cert_scalar_for_storage(form.get("cert_title"))
    tpl.cert_subtitle = normalize_cert_scalar_for_storage(form.get("cert_subtitle"))
    tpl.cert_footer = normalize_cert_scalar_for_storage(form.get("cert_footer"))
    tpl.cert_signer_name = normalize_cert_scalar_for_storage(form.get("cert_signer_name"))
    tpl.cert_signer_designation = normalize_cert_scalar_for_storage(form.get("cert_signer_designation"))
    tpl.cert_signature_url = normalize_cert_scalar_for_storage(form.get("cert_signature_url"))
    tpl.cert_logo_url = normalize_cert_scalar_for_storage(form.get("cert_logo_url"))
    tpl.cert_bg_url = normalize_cert_scalar_for_storage(form.get("cert_bg_url"))
    tpl.cert_color_scheme = normalize_cert_scalar_for_storage(form.get("cert_color_scheme"))
    tpl.cert_style = normalize_cert_scalar_for_storage(form.get("cert_style"))
    log_activity(
        db,
        category="admin",
        action="update",
        description=f"Updated certificate template '{tpl.name}'",
        request=request,
        user_id=admin.id,
        target_type="certificate_template",
        target_id=tpl.id,
    )
    db.commit()
    flash(request, f"Template “{tpl.name}” saved.", "success")
    return RedirectResponse("/admin/certificate-templates", status_code=303)


@router.post("/certificate-templates/{template_id}/delete")
def certificate_template_delete(request: Request, template_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/certificate-templates", status_code=303)
    name = tpl.name
    db.query(Event).filter(Event.cert_template_id == tpl.id).update({Event.cert_template_id: None})
    log_activity(
        db,
        category="admin",
        action="delete",
        description=f"Deleted certificate template '{name}'",
        request=request,
        user_id=admin.id,
        target_type="certificate_template",
        target_id=template_id,
    )
    db.delete(tpl)
    db.commit()
    flash(request, "Certificate template deleted.", "success")
    return RedirectResponse("/admin/certificate-templates", status_code=303)


@router.get("/certificate-templates/{template_id}/designer")
def certificate_template_designer_page(
    request: Request, template_id: int, db: Session = Depends(get_db)
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/certificate-templates", status_code=303)
    designer_bootstrap = _certificate_template_designer_bootstrap(tpl)
    return templates.TemplateResponse(
        "admin/certificate_designer.html",
        _admin_ctx(
            request,
            active_page="cert-templates",
            cert_template=tpl,
            event=None,
            designer_bootstrap=designer_bootstrap,
        ),
    )


@router.post("/certificate-templates/{template_id}/certificate/designer")
async def certificate_template_designer_save(
    request: Request, template_id: int, db: Session = Depends(get_db)
):
    ct = (request.headers.get("content-type") or "").lower()
    admin = _require_admin(request, db)
    if not admin:
        if "application/json" in ct:
            return JSONResponse({"error": "Unauthorized"}, status_code=401)
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        if "application/json" in ct:
            return JSONResponse({"error": "Not found"}, status_code=404)
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/certificate-templates", status_code=303)

    if "application/json" in ct:
        body = await request.json()
        payload = body.get("cert_style")
        if isinstance(payload, dict):
            tpl.cert_style = json.dumps(payload)
        elif payload is None or payload == "":
            tpl.cert_style = None
        else:
            tpl.cert_style = str(payload).strip() or None
        log_activity(
            db,
            category="admin",
            action="update",
            description=f"Updated certificate visual layout for template '{tpl.name}'",
            request=request,
            user_id=admin.id,
            target_type="certificate_template",
            target_id=tpl.id,
        )
        db.commit()
        return JSONResponse({"ok": True})
    form = await request.form()
    tpl.cert_style = (form.get("cert_style") or "").strip() or None
    log_activity(
        db,
        category="admin",
        action="update",
        description=f"Updated certificate visual layout for template '{tpl.name}'",
        request=request,
        user_id=admin.id,
        target_type="certificate_template",
        target_id=tpl.id,
    )
    db.commit()
    flash(request, "Certificate layout saved.", "success")
    return RedirectResponse(f"/admin/certificate-templates/{template_id}/designer", status_code=303)


@router.get("/certificate-templates/{template_id}/certificate/legacy-to-freeform")
def certificate_template_legacy_to_freeform(
    request: Request, template_id: int, db: Session = Depends(get_db)
):
    from app.services.certificate import legacy_cert_style_to_freeform_dict

    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)
    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        return JSONResponse({"error": "Not found"}, status_code=404)
    return JSONResponse(legacy_cert_style_to_freeform_dict(tpl))


@router.get("/certificate-templates/{template_id}/certificate/preview")
def certificate_template_certificate_preview(
    request: Request, template_id: int, db: Session = Depends(get_db)
):
    import io as _io
    from types import SimpleNamespace
    from app.services.certificate import generate_certificate_pdf

    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)

    tpl = db.query(CertificateTemplate).get(template_id)
    if not tpl:
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/certificate-templates", status_code=303)

    dummy_booking = SimpleNamespace(
        booking_ref="PREVIEW",
        qr_code_data="CERT-PREVIEW-SAMPLE",
    )
    dummy_user = SimpleNamespace(
        full_name="Sample Attendee",
        username="sample_attendee",
    )
    dummy_event = SimpleNamespace(
        name=tpl.name or "Certificate Template",
        start_date=date.today(),
    )

    pdf_bytes = generate_certificate_pdf(dummy_booking, dummy_user, tpl, dummy_event, None)

    return StreamingResponse(
        _io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={
            "Content-Disposition": 'inline; filename="certificate-template-preview.pdf"',
            "Cache-Control": "no-store",
        },
    )


# ─── Feedback Templates ───


@router.get("/feedback-templates")
def feedback_template_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpls = db.query(FeedbackTemplate).order_by(FeedbackTemplate.created_at.desc()).all()
    return templates.TemplateResponse(
        "admin/feedback_template_list.html",
        _admin_ctx(request, active_page="feedback-templates", templates_list=tpls),
    )


@router.get("/feedback-templates/new")
def feedback_template_new(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    return templates.TemplateResponse(
        "admin/feedback_template_form.html",
        _admin_ctx(request, active_page="feedback-templates", template=None),
    )


@router.post("/feedback-templates/new")
async def feedback_template_create(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    tpl = FeedbackTemplate(
        name=form.get("name", "").strip(),
        description=form.get("description", "").strip() or None,
        created_by=admin.id,
        session_ratings_enabled="session_ratings_enabled" in form,
        session_ratings_required="session_ratings_required" in form,
    )
    db.add(tpl)
    db.flush()

    for idx_str in form.getlist("q_idx"):
        idx = int(idx_str)
        text = form.get(f"q_text_{idx}", "").strip()
        if not text:
            continue
        q_type = form.get(f"q_type_{idx}", "text")
        options = None
        if q_type == "multiple_choice":
            raw = form.get(f"q_options_{idx}", "").strip()
            options = [o.strip() for o in raw.splitlines() if o.strip()] if raw else None
        page = int(form.get(f"q_page_{idx}", "1") or "1")
        tq = TemplateQuestion(
            template_id=tpl.id,
            order=idx,
            page=page,
            question_text=text,
            question_type=q_type,
            options_json=options,
            is_required=f"q_required_{idx}" in form,
        )
        db.add(tq)

    log_activity(db, category="admin", action="create", description=f"Created feedback template '{tpl.name}'", request=request, user_id=admin.id, target_type="feedback_template", target_id=tpl.id)
    db.commit()
    flash(request, f"Template '{tpl.name}' created.", "success")
    return RedirectResponse("/admin/feedback-templates", status_code=303)


@router.get("/feedback-templates/{template_id}/edit")
def feedback_template_edit(request: Request, template_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(FeedbackTemplate).get(template_id)
    if not tpl:
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/feedback-templates", status_code=303)
    return templates.TemplateResponse(
        "admin/feedback_template_form.html",
        _admin_ctx(request, active_page="feedback-templates", template=tpl),
    )


@router.post("/feedback-templates/{template_id}/edit")
async def feedback_template_update(request: Request, template_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(FeedbackTemplate).get(template_id)
    if not tpl:
        flash(request, "Template not found.", "danger")
        return RedirectResponse("/admin/feedback-templates", status_code=303)

    form = await _form(request)
    tpl.name = form.get("name", "").strip()
    tpl.description = form.get("description", "").strip() or None
    tpl.session_ratings_enabled = "session_ratings_enabled" in form
    tpl.session_ratings_required = "session_ratings_required" in form

    db.query(TemplateQuestion).filter(TemplateQuestion.template_id == tpl.id).delete()
    for idx_str in form.getlist("q_idx"):
        idx = int(idx_str)
        text = form.get(f"q_text_{idx}", "").strip()
        if not text:
            continue
        q_type = form.get(f"q_type_{idx}", "text")
        options = None
        if q_type == "multiple_choice":
            raw = form.get(f"q_options_{idx}", "").strip()
            options = [o.strip() for o in raw.splitlines() if o.strip()] if raw else None
        page = int(form.get(f"q_page_{idx}", "1") or "1")
        tq = TemplateQuestion(
            template_id=tpl.id,
            order=idx,
            page=page,
            question_text=text,
            question_type=q_type,
            options_json=options,
            is_required=f"q_required_{idx}" in form,
        )
        db.add(tq)

    log_activity(db, category="admin", action="update", description=f"Updated feedback template '{tpl.name}'", request=request, user_id=admin.id, target_type="feedback_template", target_id=tpl.id)
    db.commit()
    flash(request, f"Template '{tpl.name}' updated.", "success")
    return RedirectResponse("/admin/feedback-templates", status_code=303)


@router.post("/feedback-templates/{template_id}/delete")
def feedback_template_delete(request: Request, template_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    tpl = db.query(FeedbackTemplate).get(template_id)
    if tpl:
        db.query(Event).filter(Event.feedback_template_id == tpl.id).update({Event.feedback_template_id: None})
        log_activity(db, category="admin", action="delete", description=f"Deleted feedback template '{tpl.name}'", request=request, user_id=admin.id, target_type="feedback_template", target_id=tpl.id)
        db.delete(tpl)
        db.commit()
        flash(request, "Template deleted.", "success")
    return RedirectResponse("/admin/feedback-templates", status_code=303)


# ─── Newsletter Campaigns ───


@router.get("/newsletters")
def newsletter_list(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    campaigns = db.query(Newsletter).order_by(Newsletter.created_at.desc()).all()
    subscriber_count = db.query(func.count(NewsletterSubscriber.id)).scalar() or 0
    return templates.TemplateResponse(
        "admin/newsletters.html",
        _admin_ctx(request, active_page="newsletters", campaigns=campaigns, subscriber_count=subscriber_count),
    )


@router.get("/newsletters/new")
def newsletter_new_form(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    subscriber_count = db.query(func.count(NewsletterSubscriber.id)).scalar() or 0
    return templates.TemplateResponse(
        "admin/newsletter_compose.html",
        _admin_ctx(request, active_page="newsletters", newsletter=None, subscriber_count=subscriber_count),
    )


@router.post("/newsletters/new")
async def newsletter_new_save(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    form = await _form(request)
    nl = Newsletter(
        subject=form.get("subject", "").strip(),
        body_html=form.get("body_html", ""),
        status="draft",
    )
    db.add(nl)
    db.flush()
    log_activity(db, category="admin", action="create", description=f"Created newsletter draft '{nl.subject}'", request=request, user_id=admin.id, target_type="newsletter", target_id=nl.id)
    db.commit()
    flash(request, "Newsletter draft saved.", "success")
    return RedirectResponse(f"/admin/newsletters/{nl.id}/edit", status_code=303)


@router.get("/newsletters/{newsletter_id}/edit")
def newsletter_edit_form(request: Request, newsletter_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    nl = db.query(Newsletter).get(newsletter_id)
    if not nl:
        flash(request, "Newsletter not found.", "danger")
        return RedirectResponse("/admin/newsletters", status_code=303)
    subscriber_count = db.query(func.count(NewsletterSubscriber.id)).scalar() or 0
    return templates.TemplateResponse(
        "admin/newsletter_compose.html",
        _admin_ctx(request, active_page="newsletters", newsletter=nl, subscriber_count=subscriber_count),
    )


@router.post("/newsletters/{newsletter_id}/edit")
async def newsletter_edit_save(request: Request, newsletter_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    nl = db.query(Newsletter).get(newsletter_id)
    if not nl:
        flash(request, "Newsletter not found.", "danger")
        return RedirectResponse("/admin/newsletters", status_code=303)
    if nl.status != "draft":
        flash(request, "Only draft newsletters can be edited.", "warning")
        return RedirectResponse("/admin/newsletters", status_code=303)
    form = await _form(request)
    nl.subject = form.get("subject", "").strip()
    nl.body_html = form.get("body_html", "")
    log_activity(db, category="admin", action="update", description=f"Updated newsletter draft '{nl.subject}'", request=request, user_id=admin.id, target_type="newsletter", target_id=nl.id)
    db.commit()
    flash(request, "Newsletter draft updated.", "success")
    return RedirectResponse(f"/admin/newsletters/{nl.id}/edit", status_code=303)


@router.post("/newsletters/{newsletter_id}/delete")
def newsletter_delete(request: Request, newsletter_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    nl = db.query(Newsletter).get(newsletter_id)
    if not nl:
        flash(request, "Newsletter not found.", "danger")
        return RedirectResponse("/admin/newsletters", status_code=303)
    if nl.status != "draft":
        flash(request, "Only draft newsletters can be deleted.", "warning")
        return RedirectResponse("/admin/newsletters", status_code=303)
    log_activity(db, category="admin", action="delete", description=f"Deleted newsletter draft '{nl.subject}'", request=request, user_id=admin.id, target_type="newsletter", target_id=nl.id)
    db.delete(nl)
    db.commit()
    flash(request, "Newsletter draft deleted.", "success")
    return RedirectResponse("/admin/newsletters", status_code=303)


@router.post("/newsletters/{newsletter_id}/send")
def newsletter_send(request: Request, newsletter_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    nl = db.query(Newsletter).get(newsletter_id)
    if not nl:
        flash(request, "Newsletter not found.", "danger")
        return RedirectResponse("/admin/newsletters", status_code=303)
    if nl.status != "draft":
        flash(request, "This newsletter has already been sent or is currently sending.", "warning")
        return RedirectResponse("/admin/newsletters", status_code=303)
    subscriber_count = db.query(func.count(NewsletterSubscriber.id)).scalar() or 0
    if subscriber_count == 0:
        flash(request, "No subscribers to send to.", "warning")
        return RedirectResponse(f"/admin/newsletters/{nl.id}/edit", status_code=303)
    nl.status = "sending"
    nl.total_recipients = subscriber_count
    nl.sent_count = 0
    nl.failed_count = 0
    db.commit()

    from app.services.email import send_newsletter_campaign
    send_newsletter_campaign(nl.id)

    log_activity(db, category="admin", action="send", description=f"Started sending newsletter '{nl.subject}' to {subscriber_count} subscribers", request=request, user_id=admin.id, target_type="newsletter", target_id=nl.id)
    db.commit()
    flash(request, f"Newsletter is being sent to {subscriber_count} subscriber(s).", "success")
    return RedirectResponse("/admin/newsletters", status_code=303)


@router.get("/newsletters/{newsletter_id}/status")
def newsletter_status(request: Request, newsletter_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    nl = db.query(Newsletter).get(newsletter_id)
    if not nl:
        return JSONResponse({"error": "not found"}, status_code=404)
    return JSONResponse({
        "status": nl.status,
        "total_recipients": nl.total_recipients,
        "sent_count": nl.sent_count,
        "failed_count": nl.failed_count,
    })


@router.post("/newsletters/preview")
async def newsletter_preview(request: Request, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "unauthorized"}, status_code=401)
    form = await _form(request)
    body_html = form.get("body_html", "")
    from app.services.email import wrap_newsletter_html
    full_html = wrap_newsletter_html(body_html, "#")
    from fastapi.responses import HTMLResponse
    return HTMLResponse(content=full_html)


@router.post("/upload-image")
async def admin_upload_image(request: Request, db: Session = Depends(get_db)):
    from app.models.uploaded_image import UploadedImage

    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "unauthorized"}, status_code=401)

    form = await request.form()
    file = form.get("image")
    if not file or not hasattr(file, "filename"):
        return JSONResponse({"error": "No file uploaded"}, status_code=400)

    allowed = {"image/jpeg", "image/png", "image/gif", "image/webp"}
    if file.content_type not in allowed:
        return JSONResponse({"error": "Invalid file type. Allowed: JPEG, PNG, GIF, WebP"}, status_code=400)

    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        return JSONResponse({"error": "File too large (max 5MB)"}, status_code=400)

    img = UploadedImage(
        filename=file.filename or "upload",
        content_type=file.content_type,
        data=content,
    )
    db.add(img)
    db.commit()
    db.refresh(img)

    return JSONResponse({"url": f"/uploads/{img.id}"})


@router.post("/certificate-ai/generate-template")
async def certificate_ai_generate_template(request: Request, db: Session = Depends(get_db)):
    """Two-step OpenAI pipeline: decorative background image + vision layout → cert_style v2."""
    from app.services.certificate_ai_template import CertificateAiImageError, run_ai_certificate_template_pipeline

    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "Unauthorized"}, status_code=401)
    if not (settings.openai_api_key or "").strip():
        return JSONResponse(
            {"error": "OpenAI is not configured (set OPENAI_API_KEY)."},
            status_code=503,
        )
    body: dict = {}
    ct = (request.headers.get("content-type") or "").lower()
    if "application/json" in ct:
        try:
            body = await request.json()
            if not isinstance(body, dict):
                body = {}
        except Exception:
            body = {}
    prompt_hint = (body.get("prompt_hint") or body.get("hint") or "") if isinstance(body, dict) else ""
    if isinstance(prompt_hint, str):
        prompt_hint = prompt_hint.strip()[:2000]
    else:
        prompt_hint = ""

    try:
        result = run_ai_certificate_template_pipeline(db, settings, prompt_hint=prompt_hint)
    except CertificateAiImageError as exc:
        return JSONResponse({"error": str(exc), "step": "image"}, status_code=502)

    return JSONResponse(
        {
            "cert_style": result["cert_style"],
            "upload_url": result["upload_url"],
            "layout_fallback": result["layout_fallback"],
        }
    )


@router.post("/newsletters/upload-image")
async def newsletter_upload_image(request: Request, db: Session = Depends(get_db)):
    import os
    import uuid
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"error": "unauthorized"}, status_code=401)

    form = await request.form()
    file = form.get("image")
    if not file or not hasattr(file, "filename"):
        return JSONResponse({"error": "No file uploaded"}, status_code=400)

    allowed = {"image/jpeg", "image/png", "image/gif", "image/webp"}
    if file.content_type not in allowed:
        return JSONResponse({"error": "Invalid file type. Allowed: JPEG, PNG, GIF, WebP"}, status_code=400)

    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        return JSONResponse({"error": "File too large (max 5MB)"}, status_code=400)

    ext_map = {"image/jpeg": ".jpg", "image/png": ".png", "image/gif": ".gif", "image/webp": ".webp"}
    ext = ext_map.get(file.content_type, ".jpg")
    filename = f"{uuid.uuid4().hex}{ext}"

    upload_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "uploads", "newsletters")
    os.makedirs(upload_dir, exist_ok=True)
    filepath = os.path.join(upload_dir, filename)

    with open(filepath, "wb") as f:
        f.write(content)

    url = f"/static/uploads/newsletters/{filename}"
    return JSONResponse({"url": url})


# ---------------------------------------------------------------------------
# Event Management Hub
# ---------------------------------------------------------------------------

def _hub_ctx(request: Request, event: Event, hub_tab: str, **kwargs):
    ctx = _admin_ctx(request, active_page="event_mgmt", hub_event=event, hub_tab=hub_tab, **kwargs)
    return ctx


@router.get("/event-management")
def event_management_landing(request: Request):
    return RedirectResponse("/admin/", status_code=302)


@router.get("/event-management/{event_id}")
def event_management_overview(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    bookings_count = db.query(func.count(Booking.id)).filter(
        Booking.event_id == event_id, Booking.payment_status == "paid",
        Booking.is_shared_ticket == False,
    ).scalar() or 0
    revenue = db.query(func.sum(Booking.amount_paid)).filter(
        Booking.event_id == event_id, Booking.payment_status == "paid",
        Booking.is_shared_ticket == False,
    ).scalar() or 0
    checked_in = db.query(func.count(Booking.id)).filter(
        Booking.event_id == event_id, Booking.payment_status == "paid", Booking.checked_in == True
    ).scalar() or 0
    checkin_pct = round(checked_in / bookings_count * 100) if bookings_count else 0
    waitlist_count = db.query(func.count(Waitlist.id)).filter(Waitlist.event_id == event_id).scalar() or 0
    active_polls = db.query(func.count(Poll.id)).filter(
        Poll.event_id == event_id, Poll.is_active == True
    ).scalar() or 0

    stats = {
        "bookings": bookings_count, "revenue": revenue, "checked_in": checked_in,
        "checkin_pct": checkin_pct, "waitlist": waitlist_count, "active_polls": active_polls,
    }

    from app.models.event_session import EventSession
    ev_sessions = (
        db.query(EventSession).filter(EventSession.event_id == event_id)
        .order_by(EventSession.order, EventSession.start_time).all()
    )
    sessions_enriched = []
    for es in ev_sessions:
        s = es.session
        override_speaker = es.speaker_name or (es.speaker.name if es.speaker else None)
        speakers = [override_speaker] if override_speaker else []
        if not speakers:
            speakers = [ss.speaker.name for ss in s.session_speakers if ss.speaker] if hasattr(s, "session_speakers") else []
        poll_count = db.query(func.count(Poll.id)).filter(Poll.session_id == s.id, Poll.event_id == event_id).scalar() or 0
        sessions_enriched.append({
            "session": s, "speakers": speakers, "poll_count": poll_count,
            "start_time": es.start_time, "order": es.order,
        })

    return templates.TemplateResponse(
        "admin/event_management/overview.html",
        _hub_ctx(request, event, "overview", stats=stats, sessions=sessions_enriched),
    )


@router.get("/event-management/{event_id}/bookings")
def event_management_bookings(
    request: Request, event_id: int, db: Session = Depends(get_db),
    q: str = Query("", alias="q"),
    status_filter: str = Query("", alias="status"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    query = db.query(Booking).filter(Booking.event_id == event_id)
    if status_filter:
        query = query.filter(Booking.payment_status == status_filter)
    else:
        query = query.filter(Booking.payment_status.in_(["paid", "hold", "refunded"]))
    bookings = query.order_by(Booking.booked_at.desc()).all()

    enriched = []
    for b in bookings:
        u = db.query(User).get(b.user_id)
        seat = db.query(Seat).get(b.seat_id)
        if q:
            search = q.lower()
            match = (
                (u and (search in u.username.lower() or search in u.email.lower() or (u.full_name and search in u.full_name.lower())))
                or (b.booking_ref and search in b.booking_ref.lower())
                or (b.ticket_id and search in b.ticket_id.lower())
            )
            if not match:
                continue
        enriched.append({"booking": b, "user": u, "seat": seat})

    return templates.TemplateResponse(
        "admin/event_management/bookings.html",
        _hub_ctx(request, event, "bookings", bookings=enriched, q=q, status_filter=status_filter),
    )


@router.get("/event-management/{event_id}/checkin")
def event_management_checkin_page(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    total = db.query(func.count(Booking.id)).filter(Booking.event_id == event_id, Booking.payment_status == "paid").scalar() or 0
    ci = db.query(func.count(Booking.id)).filter(Booking.event_id == event_id, Booking.payment_status == "paid", Booking.checked_in == True).scalar() or 0
    stats = {"total": total, "checked_in": ci, "pct": round(ci / total * 100) if total else 0}

    return templates.TemplateResponse(
        "admin/event_management/checkin.html",
        _hub_ctx(request, event, "checkin", stats=stats, result=None),
    )


@router.post("/event-management/{event_id}/checkin")
async def event_management_checkin_verify(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    form = await _form(request)
    ticket_id = form.get("ticket_id", "").strip()

    if not ticket_id:
        total = db.query(func.count(Booking.id)).filter(Booking.event_id == event_id, Booking.payment_status == "paid").scalar() or 0
        ci = db.query(func.count(Booking.id)).filter(Booking.event_id == event_id, Booking.payment_status == "paid", Booking.checked_in == True).scalar() or 0
        stats = {"total": total, "checked_in": ci, "pct": round(ci / total * 100) if total else 0}
        return templates.TemplateResponse(
            "admin/event_management/checkin.html",
            _hub_ctx(request, event, "checkin", stats=stats, result={"status": "error", "msg": "Please enter a ticket ID."}),
        )

    is_group = ticket_id.startswith("GROUP-")
    result = None

    if is_group:
        group_id = ticket_id[6:]
        all_group_any = db.query(Booking).filter(Booking.booking_group == group_id).all()
        all_group = [b for b in all_group_any if b.payment_status == "paid" and b.event_id == event_id]
        refunded_count = sum(1 for b in all_group_any if b.event_id == event_id and b.payment_status in ("refunded", "cancelled"))

        if not all_group_any:
            result = {"status": "error", "msg": f"Group '{group_id}' not found."}
        elif not all_group:
            result = {"status": "error", "msg": f"No valid tickets in this group for this event."}
        else:
            now = now_ist()
            newly_checked, already_checked = [], []
            for gb in all_group:
                seat = db.query(Seat).get(gb.seat_id)
                label = seat.label if seat else gb.ticket_id
                if gb.checked_in:
                    already_checked.append(label)
                else:
                    gb.checked_in = True
                    gb.checked_in_at = now
                    newly_checked.append(label)
            db.commit()
            user = db.query(User).get(all_group[0].user_id)
            refunded_note = f" ({refunded_count} ticket(s) refunded/cancelled.)" if refunded_count else ""
            if newly_checked and not already_checked:
                msg = f"Check-in successful! {len(newly_checked)} ticket(s).{refunded_note}"
                status = "success"
            elif newly_checked:
                msg = f"Checked in {len(newly_checked)} ticket(s). {len(already_checked)} already checked in.{refunded_note}"
                status = "success"
            else:
                msg = f"Re-entry — all {len(already_checked)} ticket(s) already checked in.{refunded_note}"
                status = "reentry"
            if newly_checked:
                log_activity(db, category="admin", action="checkin", description=f"Group check-in: {len(newly_checked)} ticket(s) for '{event.name}'", request=request, user_id=admin.id, target_type="booking", target_id=all_group[0].id)
            result = {
                "status": status, "msg": msg, "is_group": True,
                "user_name": user.full_name or user.username if user else "Unknown",
                "user_email": user.email if user else "",
                "event_name": event.name,
                "newly_checked": newly_checked, "already_checked": already_checked,
            }
    else:
        booking = db.query(Booking).filter(
            Booking.ticket_id == ticket_id, Booking.payment_status == "paid", Booking.event_id == event_id
        ).first()
        if not booking:
            result = {"status": "error", "msg": f"Ticket '{ticket_id}' not found or not valid for this event."}
        elif booking.checked_in:
            user = db.query(User).get(booking.user_id)
            seat = db.query(Seat).get(booking.seat_id)
            time_str = booking.checked_in_at.strftime('%I:%M %p') if booking.checked_in_at else 'earlier'
            result = {
                "status": "reentry", "msg": f"Re-entry — ticket valid. Originally checked in at {time_str}.",
                "user_name": user.full_name or user.username if user else "Unknown",
                "user_email": user.email if user else "",
                "seat_label": seat.label if seat else "", "ticket_id": ticket_id,
            }
        else:
            booking.checked_in = True
            booking.checked_in_at = now_ist()
            user = db.query(User).get(booking.user_id)
            seat = db.query(Seat).get(booking.seat_id)
            log_activity(db, category="admin", action="checkin", description=f"Checked in ticket '{ticket_id}' for '{event.name}'", request=request, user_id=admin.id, target_type="booking", target_id=booking.id)
            db.commit()
            result = {
                "status": "success", "msg": "Check-in successful!",
                "user_name": user.full_name or user.username if user else "Unknown",
                "user_email": user.email if user else "",
                "seat_label": seat.label if seat else "", "ticket_id": ticket_id,
            }

    total = db.query(func.count(Booking.id)).filter(Booking.event_id == event_id, Booking.payment_status == "paid").scalar() or 0
    ci = db.query(func.count(Booking.id)).filter(Booking.event_id == event_id, Booking.payment_status == "paid", Booking.checked_in == True).scalar() or 0
    stats = {"total": total, "checked_in": ci, "pct": round(ci / total * 100) if total else 0}

    return templates.TemplateResponse(
        "admin/event_management/checkin.html",
        _hub_ctx(request, event, "checkin", stats=stats, result=result),
    )


@router.get("/event-management/{event_id}/waitlist")
def event_management_waitlist(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    wl_entries = db.query(Waitlist).filter(Waitlist.event_id == event_id).order_by(Waitlist.joined_at.desc()).all()
    enriched = []
    for w in wl_entries:
        u = db.query(User).get(w.user_id)
        enriched.append({"entry": w, "user": u})

    return templates.TemplateResponse(
        "admin/event_management/waitlist.html",
        _hub_ctx(request, event, "waitlist", entries=enriched),
    )


def _enrich_poll(p):
    """Build display-ready dict for a single Poll object."""
    enriched: dict = {"poll": p}
    if p.poll_type in ("multiple_choice", "yes_no"):
        total = sum(len(o.votes) for o in p.options)
        opts = []
        for o in p.options:
            count = len(o.votes)
            opts.append({"option": o, "votes": count, "pct": round(count / total * 100, 1) if total else 0})
        enriched.update(total_votes=total, options=opts)
    elif p.poll_type == "rating":
        votes = [v for v in p.votes if v.rating_value is not None]
        total = len(votes)
        avg = round(sum(v.rating_value for v in votes) / total, 1) if total else 0
        dist = {s: 0 for s in range(1, 6)}
        for v in votes:
            dist[v.rating_value] = dist.get(v.rating_value, 0) + 1
        enriched.update(total_votes=total, average=avg, distribution=dist)
    elif p.poll_type == "text":
        votes = [v for v in p.votes if v.text_answer]
        enriched.update(total_votes=len(votes), text_responses=[v.text_answer for v in votes])
    else:
        enriched.update(total_votes=0, options=[])
    return enriched


@router.get("/event-management/{event_id}/polls")
def event_management_polls(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    from app.models.event_session import EventSession
    polls = db.query(Poll).filter(Poll.event_id == event_id).order_by(Poll.created_at.desc()).all()

    by_session: dict[int, list] = {}
    for p in polls:
        by_session.setdefault(p.session_id, []).append(_enrich_poll(p))

    ev_sessions = (
        db.query(EventSession).filter(EventSession.event_id == event_id)
        .order_by(EventSession.order, EventSession.start_time).all()
    )
    sessions = [es.session for es in ev_sessions]
    grouped = []
    for s in sessions:
        session_polls = by_session.pop(s.id, [])
        grouped.append({"session": s, "polls": session_polls})
    for sid, remaining in by_session.items():
        orphan = db.query(SessionModel).get(sid)
        grouped.append({"session": orphan, "polls": remaining})

    return templates.TemplateResponse(
        "admin/event_management/polls.html",
        _hub_ctx(request, event, "polls", grouped=grouped, sessions=sessions, total_polls=len(polls)),
    )


@router.post("/event-management/{event_id}/polls")
async def event_management_create_poll(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"ok": False}, status_code=403)
    event = db.query(Event).get(event_id)
    if not event:
        return JSONResponse({"ok": False, "error": "Event not found."}, status_code=404)
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "error": "Invalid request."}, status_code=400)

    session_id = body.get("session_id")
    if not session_id:
        return JSONResponse({"ok": False, "error": "Select a session."}, status_code=400)
    session_obj = db.query(SessionModel).get(int(session_id))
    if not session_obj:
        return JSONResponse({"ok": False, "error": "Session not found."}, status_code=404)

    question = (body.get("question") or "").strip()
    poll_type = body.get("poll_type", "multiple_choice")
    if poll_type not in ("multiple_choice", "yes_no", "rating", "text"):
        poll_type = "multiple_choice"
    options_list = body.get("options", [])
    if not question:
        return JSONResponse({"ok": False, "error": "Provide a question."}, status_code=400)
    if poll_type == "multiple_choice" and len(options_list) < 2:
        return JSONResponse({"ok": False, "error": "Provide at least 2 options."}, status_code=400)

    poll = Poll(
        session_id=int(session_id), event_id=event_id, question=question, poll_type=poll_type,
        allow_multiple=bool(body.get("allow_multiple")), created_by=admin.id,
    )
    db.add(poll)
    db.flush()
    if poll_type == "yes_no":
        db.add(PollOption(poll_id=poll.id, option_text="Yes", order=0))
        db.add(PollOption(poll_id=poll.id, option_text="No", order=1))
    elif poll_type == "multiple_choice":
        for i, opt_text in enumerate(options_list):
            opt_text = (opt_text or "").strip()
            if opt_text:
                db.add(PollOption(poll_id=poll.id, option_text=opt_text, order=i))
    log_activity(db, category="admin", action="create", description=f"Created poll for session '{session_obj.title}' in event '{event.name}'", request=request, user_id=admin.id, target_type="poll", target_id=poll.id)
    db.commit()
    return JSONResponse({"ok": True, "poll_id": poll.id})


# ─── Event-Level Recordings (Hub) ───

@router.get("/event-management/{event_id}/recordings")
def event_management_recordings(request: Request, event_id: int, db: Session = Depends(get_db)):
    from app.services.polls import build_embed_url as _build_embed_url
    from app.models.event_session import EventSession
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    ev_sessions = (
        db.query(EventSession).filter(EventSession.event_id == event_id)
        .order_by(EventSession.order, EventSession.start_time).all()
    )
    sessions = [es.session for es in ev_sessions]
    session_ids = [es.session_id for es in ev_sessions]

    recordings = (
        db.query(SessionRecording)
        .filter(SessionRecording.session_id.in_(session_ids))
        .filter((SessionRecording.event_id == event_id) | (SessionRecording.event_id.is_(None)))
        .order_by(SessionRecording.order).all()
    ) if session_ids else []

    enriched = [{"rec": r, "embed_url": _build_embed_url(r.url), "session": r.session} for r in recordings]
    return templates.TemplateResponse(
        "admin/event_management/recordings.html",
        _hub_ctx(request, event, "recordings", recordings=enriched, sessions=sessions),
    )


@router.post("/event-management/{event_id}/recordings")
async def event_management_recording_add(request: Request, event_id: int, db: Session = Depends(get_db)):
    from app.models.event_session import EventSession
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)
    form = await request.form()
    url = form.get("url", "").strip()
    err = _validate_recording_url(url)
    if err:
        flash(request, err, "danger")
        return RedirectResponse(f"/admin/event-management/{event_id}/recordings", status_code=303)
    session_id_raw = form.get("session_id", "").strip()
    if not session_id_raw or not session_id_raw.isdigit():
        flash(request, "Select a session.", "danger")
        return RedirectResponse(f"/admin/event-management/{event_id}/recordings", status_code=303)
    session_id = int(session_id_raw)
    title = form.get("title", "").strip() or None
    is_public = "is_public" in form
    max_order = db.query(func.coalesce(func.max(SessionRecording.order), -1)).filter(
        SessionRecording.session_id == session_id, SessionRecording.event_id == event_id
    ).scalar()
    rec = SessionRecording(
        session_id=session_id, event_id=event_id,
        url=url, title=title, order=max_order + 1, is_public=is_public,
    )
    db.add(rec)
    db.commit()
    flash(request, "Recording added.", "success")
    return RedirectResponse(f"/admin/event-management/{event_id}/recordings", status_code=303)


@router.post("/event-management/{event_id}/recordings/{rec_id}/update")
async def event_management_recording_update(request: Request, event_id: int, rec_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    rec = db.query(SessionRecording).filter(SessionRecording.id == rec_id).first()
    if not rec:
        flash(request, "Recording not found.", "danger")
        return RedirectResponse(f"/admin/event-management/{event_id}/recordings", status_code=303)
    form = await request.form()
    rec.title = form.get("title", "").strip() or None
    db.commit()
    flash(request, "Recording title updated.", "success")
    return RedirectResponse(f"/admin/event-management/{event_id}/recordings", status_code=303)


@router.post("/event-management/{event_id}/recordings/{rec_id}/toggle")
def event_management_recording_toggle(request: Request, event_id: int, rec_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    rec = db.query(SessionRecording).filter(SessionRecording.id == rec_id).first()
    if rec:
        rec.is_public = not rec.is_public
        db.commit()
    return RedirectResponse(f"/admin/event-management/{event_id}/recordings", status_code=303)


@router.post("/event-management/{event_id}/recordings/{rec_id}/delete")
def event_management_recording_delete(request: Request, event_id: int, rec_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    rec = db.query(SessionRecording).filter(SessionRecording.id == rec_id).first()
    if rec:
        db.delete(rec)
        db.commit()
        flash(request, "Recording removed.", "success")
    return RedirectResponse(f"/admin/event-management/{event_id}/recordings", status_code=303)


@router.get("/event-management/{event_id}/feedback")
def event_management_feedback(
    request: Request, event_id: int, db: Session = Depends(get_db),
    rating_filter: str = Query("", alias="rating"),
    featured_filter: str = Query("", alias="featured"),
):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    query = db.query(Feedback).filter(Feedback.event_id == event_id)
    if rating_filter:
        try:
            query = query.filter(Feedback.rating == int(rating_filter))
        except ValueError:
            pass
    if featured_filter == "yes":
        query = query.filter(Feedback.is_featured == True)
    elif featured_filter == "no":
        query = query.filter(Feedback.is_featured == False)

    feedback_items = query.order_by(Feedback.created_at.desc()).all()
    enriched = []
    for fb in feedback_items:
        user = db.query(User).get(fb.user_id)
        enriched.append({"feedback": fb, "user": user})

    return templates.TemplateResponse(
        "admin/event_management/feedback.html",
        _hub_ctx(request, event, "feedback", feedback_items=enriched,
                 rating_filter=rating_filter, featured_filter=featured_filter),
    )


@router.get("/event-management/{event_id}/alerts")
def event_management_alerts(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        return RedirectResponse("/admin/event-management", status_code=303)

    alerts = db.query(EventAlert).filter(EventAlert.event_id == event_id).order_by(EventAlert.created_at.desc()).all()
    enriched = []
    for a in alerts:
        admin_user = db.query(User).get(a.admin_id)
        enriched.append({"alert": a, "admin": admin_user})

    return templates.TemplateResponse(
        "admin/event_management/alerts.html",
        _hub_ctx(request, event, "alerts", alerts=enriched),
    )


@router.post("/event-management/{event_id}/alerts/send")
async def event_management_send_alert(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return JSONResponse({"ok": False, "error": "Unauthorized"}, status_code=403)
    event = db.query(Event).get(event_id)
    if not event:
        return JSONResponse({"ok": False, "error": "Event not found"}, status_code=404)

    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "error": "Invalid request"}, status_code=400)

    message = (body.get("message") or "").strip()
    alert_type = body.get("alert_type", "info")
    if alert_type not in ("info", "warning", "urgent"):
        alert_type = "info"
    if not message:
        return JSONResponse({"ok": False, "error": "Message is required"}, status_code=400)

    alert = EventAlert(event_id=event_id, admin_id=admin.id, message=message, alert_type=alert_type)
    db.add(alert)
    db.commit()

    user_ids = set(
        r[0] for r in
        db.query(Booking.user_id).filter(
            Booking.event_id == event_id, Booking.payment_status == "paid"
        ).distinct().all()
    )
    user_ids.add(admin.id)

    from app.services.poll_events import publish_to_users
    publish_to_users(list(user_ids), {
        "type": "alert",
        "message": message,
        "alert_type": alert_type,
        "event_name": event.name,
        "event_id": event_id,
    })

    attendee_count = len(user_ids) - 1
    log_activity(db, category="admin", action="alert", description=f"Sent alert to {attendee_count} attendee(s) of '{event.name}'", request=request, user_id=admin.id, target_type="event", target_id=event_id)

    return JSONResponse({"ok": True, "recipients": attendee_count})


@router.get("/event-management/{event_id}/report")
def event_management_report(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/event-management", status_code=303)

    from app.models.event_session import EventSession
    from app.models.coupon import Coupon
    from app.models.event_addon import EventAddOn, BookingAddOn
    from app.models.feedback_template import FeedbackResponse
    from app.models.session_feedback import SessionFeedback

    paid_q = db.query(Booking).filter(
        Booking.event_id == event_id,
        Booking.payment_status == "paid",
        Booking.is_shared_ticket == False,
    )
    total_bookings = paid_q.count()
    total_revenue = float(db.query(func.coalesce(func.sum(Booking.amount_paid), 0)).filter(
        Booking.event_id == event_id, Booking.payment_status == "paid",
        Booking.is_shared_ticket == False,
    ).scalar() or 0)
    checked_in = paid_q.filter(Booking.checked_in == True).count()
    checkin_pct = round(checked_in / total_bookings * 100) if total_bookings else 0
    waitlist_count = db.query(Waitlist).filter(Waitlist.event_id == event_id).count()
    poll_count = db.query(Poll).filter(Poll.event_id == event_id).count()

    fb_count_new = db.query(FeedbackResponse).filter(FeedbackResponse.event_id == event_id).count()
    fb_count_legacy = db.query(Feedback).filter(Feedback.event_id == event_id).count()
    feedback_count = fb_count_new + fb_count_legacy
    session_count = db.query(EventSession).filter(EventSession.event_id == event_id).count()

    metrics = {
        "bookings": total_bookings, "revenue": total_revenue,
        "checked_in": checked_in, "checkin_pct": checkin_pct,
        "waitlist": waitlist_count, "polls": poll_count,
        "feedback": feedback_count, "sessions": session_count,
    }

    seat_rev = (
        db.query(Seat.seat_type, func.count(Booking.id), func.sum(Booking.amount_paid))
        .join(Booking, Booking.seat_id == Seat.id)
        .filter(Booking.event_id == event_id, Booking.payment_status == "paid",
                Booking.is_shared_ticket == False)
        .group_by(Seat.seat_type).all()
    )
    seat_revenue = [
        {"type": (st or "standard").replace("_", " ").title(), "count": cnt, "revenue": float(rev or 0)}
        for st, cnt, rev in seat_rev
    ]

    status_counts = dict(
        db.query(Booking.payment_status, func.count(Booking.id))
        .filter(Booking.event_id == event_id, Booking.is_shared_ticket == False)
        .group_by(Booking.payment_status).all()
    )

    day_col = cast(Booking.booked_at, Date).label("d")
    daily_rows = (
        db.query(day_col, func.count(Booking.id))
        .filter(Booking.event_id == event_id, Booking.payment_status == "paid",
                Booking.is_shared_ticket == False)
        .group_by("d").order_by("d").all()
    )
    daily_trend = [
        {"date": d.strftime("%b %d") if hasattr(d, "strftime") else str(d)[:10], "count": cnt}
        for d, cnt in daily_rows
    ]

    coupons = db.query(Coupon).filter(Coupon.event_id == event_id).all()
    coupon_data = []
    for c in coupons:
        disc = f"{c.discount_pct}%" if c.discount_pct else f"Rs.{float(c.discount_amount or 0):,.0f}"
        coupon_data.append({"code": c.code, "discount": disc, "used": c.used_count, "max": c.max_uses or 0})

    addons = db.query(EventAddOn).filter(EventAddOn.event_id == event_id).all()
    addon_data = []
    for a in addons:
        qty = int(db.query(func.coalesce(func.sum(BookingAddOn.quantity), 0)).filter(BookingAddOn.addon_id == a.id).scalar())
        addon_data.append({"title": a.title, "price": float(a.price or 0), "qty": qty, "revenue": float(a.price or 0) * qty})

    booked_user_ids = [
        r[0] for r in
        db.query(Booking.user_id).filter(
            Booking.event_id == event_id, Booking.payment_status == "paid",
            Booking.is_shared_ticket == False
        ).distinct().all()
    ]
    demographics = {"specializations": [], "disciplines": [], "years": []}
    if booked_user_ids:
        users = db.query(User).filter(User.id.in_(booked_user_ids)).all()
        from collections import Counter
        spec_counts = Counter(u.domain for u in users if u.domain).most_common(12)
        demographics["specializations"] = [{"name": n, "count": c} for n, c in spec_counts]
        disc_counts = Counter(u.discipline for u in users if u.discipline).most_common(8)
        demographics["disciplines"] = [{"name": n, "count": c} for n, c in disc_counts]
        year_counts = Counter(u.year_of_study for u in users if u.year_of_study)
        demographics["years"] = [{"year": yr, "count": year_counts[yr]} for yr in sorted(year_counts.keys())]

    fb_responses = db.query(FeedbackResponse).filter(FeedbackResponse.event_id == event_id).all()
    fb_legacy = db.query(Feedback).filter(Feedback.event_id == event_id).all()
    all_ratings = [fr.overall_rating for fr in fb_responses if fr.overall_rating]
    all_ratings += [fl.rating for fl in fb_legacy if fl.rating]
    avg_rating = round(sum(all_ratings) / len(all_ratings), 1) if all_ratings else 0
    rating_dist = {s: 0 for s in range(1, 6)}
    for r in all_ratings:
        if 1 <= r <= 5:
            rating_dist[int(r)] = rating_dist.get(int(r), 0) + 1

    session_fb = (
        db.query(SessionModel.title, func.avg(SessionFeedback.rating), func.count(SessionFeedback.id))
        .join(SessionFeedback, SessionFeedback.session_id == SessionModel.id)
        .filter(SessionFeedback.event_id == event_id, SessionFeedback.rating.isnot(None))
        .group_by(SessionModel.id, SessionModel.title).all()
    )
    session_ratings = [{"title": t, "avg": round(float(a), 1), "count": c} for t, a, c in session_fb]

    comments = [fr.comment.strip() for fr in fb_responses if fr.comment and fr.comment.strip()]
    comments += [fl.comment.strip() for fl in fb_legacy if fl.comment and fl.comment.strip()]

    feedback_data = {
        "avg_rating": avg_rating, "total_ratings": len(all_ratings),
        "distribution": rating_dist, "session_ratings": session_ratings,
        "comments": comments[:15],
    }

    polls = db.query(Poll).filter(Poll.event_id == event_id).order_by(Poll.created_at).all()
    poll_summary = []
    for p in polls:
        total_votes = db.query(PollVote).filter(PollVote.poll_id == p.id).count()
        sess_title = p.session.title if p.session else "Unknown"
        entry = {"question": p.question, "type": p.poll_type, "votes": total_votes,
                 "active": p.is_active, "closed": p.closed_at is not None,
                 "session_title": sess_title}
        if p.poll_type in ("multiple_choice", "yes_no"):
            opts = []
            for o in p.options:
                vc = db.query(PollVote).filter(PollVote.option_id == o.id).count()
                opts.append({"text": o.option_text, "votes": vc, "pct": round(vc / total_votes * 100, 1) if total_votes else 0})
            entry["options"] = opts
        elif p.poll_type == "rating":
            avg = db.query(func.avg(PollVote.rating_value)).filter(
                PollVote.poll_id == p.id, PollVote.rating_value.isnot(None)
            ).scalar()
            entry["average"] = round(float(avg), 1) if avg else 0
        poll_summary.append(entry)

    return templates.TemplateResponse(
        "admin/event_management/report.html",
        _hub_ctx(request, event, "report",
                 metrics=metrics, seat_revenue=seat_revenue,
                 status_counts=status_counts, daily_trend=daily_trend,
                 coupon_data=coupon_data, addon_data=addon_data,
                 demographics=demographics, feedback_data=feedback_data,
                 poll_summary=poll_summary),
    )


@router.get("/event-management/{event_id}/report/pdf")
def event_management_report_pdf(request: Request, event_id: int, db: Session = Depends(get_db)):
    admin = _require_admin(request, db)
    if not admin:
        return RedirectResponse("/auth/login", status_code=303)
    event = db.query(Event).get(event_id)
    if not event:
        flash(request, "Event not found.", "danger")
        return RedirectResponse("/admin/event-management", status_code=303)
    from app.services.event_report import generate_event_report_pdf
    pdf_bytes = generate_event_report_pdf(db, event_id)
    safe_name = "".join(c if c.isalnum() or c in " _-" else "_" for c in event.name)
    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="report-{safe_name}.pdf"'},
    )